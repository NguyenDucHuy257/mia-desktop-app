"""Desktop-only progress and performance adapter for source Excel rendering.

The vendored exporter remains authoritative for workbook structure, values,
styles, merges, and atomic saving.  This adapter only observes its existing
work units and replaces the expensive two-measurement fit pass with an
equivalent one-measurement implementation.
"""

from __future__ import annotations

import logging
import math
import sys
import time
from array import array
from copy import copy
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable


VENDOR_ROOT = Path(__file__).resolve().parent / "vendor" / "mia_crawl_service"
if str(VENDOR_ROOT) not in sys.path:
    sys.path.insert(0, str(VENDOR_ROOT))

import app.exporters.invoice_detail_excel_exporter as source_module
from app.exporters.invoice_detail_excel_exporter import InvoiceDetailExcelExporter
from openpyxl.utils import get_column_letter


logger = logging.getLogger("mia.excel_export")
ProgressCallback = Callable[[str, int, int], None]


class _ProgressRowBuilder:
    def __init__(self, delegate, owner: "ProgressiveInvoiceDetailExcelExporter") -> None:
        self.delegate = delegate
        self.owner = owner

    def build_rows(self, detail_payload, detail_record):
        started = time.perf_counter()
        try:
            rows = self.delegate.build_rows(detail_payload, detail_record)
            self.owner.generated_row_count += len(rows)
            return rows
        finally:
            self.owner.row_builder_seconds += time.perf_counter() - started
            self.owner.processed_invoice_count += 1
            emit_every = max(1, self.owner.invoice_total // 200)
            if (
                self.owner.processed_invoice_count == self.owner.invoice_total
                or self.owner.processed_invoice_count - self.owner._last_emitted_invoice >= emit_every
            ):
                self.owner._last_emitted_invoice = self.owner.processed_invoice_count
                self.owner._emit(
                    "build_rows",
                    self.owner.processed_invoice_count,
                    self.owner.invoice_total,
                )


class ProgressiveInvoiceDetailExcelExporter(InvoiceDetailExcelExporter):
    """Expose real source work units without changing source workbook output."""

    def __init__(self, template_path, row_builder=None, *, progress: ProgressCallback) -> None:
        super().__init__(template_path, row_builder=row_builder)
        self.progress = progress
        self.invoice_total = 0
        self.processed_invoice_count = 0
        self.generated_row_count = 0
        self.column_count = 0
        self.json_read_seconds = 0.0
        self.row_builder_seconds = 0.0
        self.worksheet_write_seconds = 0.0
        self.style_copy_seconds = 0.0
        self.fit_seconds = 0.0
        self.save_seconds = 0.0
        self.template_load_seconds = 0.0
        self.merge_seconds = 0.0
        self._write_style_copies = 0
        self._write_started_at: float | None = None
        self._raw_paths: set[Path] = set()
        self._last_emitted_row = -1
        self._last_emitted_invoice = 0

    def export(self, detail_records, output_path, from_date, to_date):
        self.invoice_total = len(detail_records)
        self._raw_paths = {
            Path(str(record.get("raw_detail_path") or ""))
            for record in detail_records
            if record.get("raw_detail_path")
        }
        original_builder = self.row_builder
        original_copy = source_module.copy
        original_load_workbook = source_module.load_workbook
        original_read_text = Path.read_text
        self.row_builder = _ProgressRowBuilder(original_builder, self)

        def measured_load_workbook(*args, **kwargs):
            started = time.perf_counter()
            try:
                return original_load_workbook(*args, **kwargs)
            finally:
                self.template_load_seconds += time.perf_counter() - started
                self._emit("load_template", 1, 1)

        def measured_read_text(path: Path, *args, **kwargs):
            started = time.perf_counter()
            try:
                return original_read_text(path, *args, **kwargs)
            finally:
                if path in self._raw_paths:
                    self.json_read_seconds += time.perf_counter() - started

        def measured_copy(value):
            started = time.perf_counter()
            try:
                return original_copy(value)
            finally:
                self.style_copy_seconds += time.perf_counter() - started
                if self.processed_invoice_count == self.invoice_total and self.column_count:
                    if self._write_started_at is None:
                        self._write_started_at = time.perf_counter()
                    self._write_style_copies += 1
                    completed_rows = min(
                        self.generated_row_count,
                        self._write_style_copies // self.column_count,
                    )
                    emit_every = max(1, self.generated_row_count // 200)
                    if (
                        completed_rows == self.generated_row_count
                        or completed_rows - self._last_emitted_row >= emit_every
                    ):
                        self._last_emitted_row = completed_rows
                        self._emit("write_rows", completed_rows, self.generated_row_count)

        source_module.load_workbook = measured_load_workbook
        source_module.copy = measured_copy
        Path.read_text = measured_read_text
        total_started = time.perf_counter()
        self._emit("load_template", 0, 1)
        try:
            result = super().export(detail_records, output_path, from_date, to_date)
            total_seconds = time.perf_counter() - total_started
            logger.info(
                "excel_export_summary scope=details invoice_count=%s row_count=%s "
                "column_count=%s json_read_ms=%.1f row_builder_ms=%.1f "
                "worksheet_write_ms=%.1f style_copy_ms=%.1f merge_ms=%.1f "
                "fit_columns_ms=%.1f workbook_save_ms=%.1f template_load_ms=%.1f "
                "total_ms=%.1f",
                result.get("invoice_count", 0),
                result.get("row_count", 0),
                self.column_count,
                self.json_read_seconds * 1000,
                self.row_builder_seconds * 1000,
                self.worksheet_write_seconds * 1000,
                self.style_copy_seconds * 1000,
                self.merge_seconds * 1000,
                self.fit_seconds * 1000,
                self.save_seconds * 1000,
                self.template_load_seconds * 1000,
                total_seconds * 1000,
            )
            return result
        finally:
            Path.read_text = original_read_text
            source_module.copy = original_copy
            source_module.load_workbook = original_load_workbook
            self.row_builder = original_builder

    def _prepare_prototype(self, worksheet, header_row, prototype_row, column_keys):
        self.column_count = len(column_keys)
        return super()._prepare_prototype(
            worksheet, header_row, prototype_row, column_keys
        )

    def _merge_invoice_serial_numbers(self, worksheet, first_data_row, invoice_groups):
        if self._write_started_at is not None:
            self.worksheet_write_seconds += time.perf_counter() - self._write_started_at
            self._write_started_at = None
        self._emit("write_rows", self.generated_row_count, self.generated_row_count)
        started = time.perf_counter()
        try:
            return super()._merge_invoice_serial_numbers(
                worksheet, first_data_row, invoice_groups
            )
        finally:
            self.merge_seconds += time.perf_counter() - started

    def _fit_cells(
        self,
        worksheet: Any,
        *,
        header_row: int,
        first_data_row: int,
        last_data_row: int,
        column_keys: list[str],
    ) -> None:
        """Preserve source sizing while measuring each non-empty cell once."""
        started = time.perf_counter()
        effective_last_row = max(header_row, last_data_row)
        row_count = max(0, effective_last_row - first_data_row + 1)
        column_count = len(column_keys)
        measured_widths = array("d", [math.nan]) * (row_count * column_count)
        fitted_widths: list[float] = []

        @lru_cache(maxsize=65_536)
        def display_width(text: str) -> float:
            return source_module._display_width(text)

        self._emit("format", 0, column_count)
        for column, _key in enumerate(column_keys, start=1):
            data_widths: list[float] = []
            for row_offset, row in enumerate(
                range(first_data_row, effective_last_row + 1)
            ):
                value = worksheet.cell(row, column).value
                if source_module._is_empty_excel_value(value):
                    continue
                width = display_width(str(value))
                measured_widths[
                    (row_offset * column_count) + (column - 1)
                ] = width
                data_widths.append(width)

            dimension = worksheet.column_dimensions[get_column_letter(column)]
            if not data_widths:
                dimension.width = 0
                dimension.hidden = True
                fitted_widths.append(0)
            else:
                width = min(max(source_module._robust_column_width(data_widths), 3), 255)
                dimension.width = width
                dimension.hidden = False
                fitted_widths.append(width)
            self._emit("format", column, column_count)

        for row_offset, row in enumerate(range(first_data_row, effective_last_row + 1)):
            required_lines = 1
            for column, width in enumerate(fitted_widths, start=1):
                content_width = measured_widths[
                    (row_offset * column_count) + (column - 1)
                ]
                if math.isnan(content_width):
                    continue
                usable_width = max(width, 1)
                if content_width <= usable_width + 0.01:
                    continue
                cell = worksheet.cell(row, column)
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.shrink_to_fit = False
                cell.alignment = alignment
                required_lines = max(
                    required_lines,
                    math.ceil(content_width / usable_width),
                )
            worksheet.row_dimensions[row].height = min(20 * required_lines, 409)

        worksheet.sheet_view.zoomScale = 90
        worksheet.sheet_view.zoomScaleNormal = 90
        worksheet.page_setup.orientation = "landscape"
        self.fit_seconds += time.perf_counter() - started

    def _save_atomically(self, workbook, output_path):
        self._emit("save", 0, 1)
        started = time.perf_counter()
        try:
            result = super()._save_atomically(workbook, output_path)
        finally:
            self.save_seconds += time.perf_counter() - started
        self._emit("save", 1, 1)
        return result

    def _emit(self, phase: str, processed: int, total: int) -> None:
        try:
            self.progress(phase, max(0, int(processed)), max(0, int(total)))
        except Exception:
            logger.exception("excel_export_progress_callback_failed phase=%s", phase)


__all__ = ["ProgressiveInvoiceDetailExcelExporter"]
