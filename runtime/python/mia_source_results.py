"""Source-native result presentation and Excel export for the desktop host.

Result tables deliberately use the same prepared workbook schemas as the
vendored source exporters. Column order and Vietnamese titles therefore come
from source XLSX templates rather than a desktop label map. Raw/path transport
fields never cross the local JSON-RPC presentation boundary.
"""

from __future__ import annotations

import base64
import hashlib
import json
import logging
import os
import re
import sqlite3
import tempfile
import time
from collections import OrderedDict
from contextlib import closing
from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable


logger = logging.getLogger("mia.excel_export")
ExportProgressCallback = Callable[[dict[str, Any]], None]


_RESULT_COUNT_CACHE: OrderedDict[tuple[Any, ...], int] = OrderedDict()
_RESULT_COUNT_CACHE_LIMIT = 64
_RESULT_ANALYSIS_CACHE: OrderedDict[tuple[Any, ...], dict[str, Any]] = OrderedDict()
_RESULT_ANALYSIS_CACHE_LIMIT = 32
_RECONCILIATION_CACHE: OrderedDict[tuple[Any, ...], dict[str, Any]] = OrderedDict()
_RECONCILIATION_CACHE_LIMIT = 16
_RECONCILIATION_ALGORITHM_VERSION = 2
_EXCLUSION_CACHE: OrderedDict[tuple[Any, ...], frozenset[str]] = OrderedDict()
_EXCLUSION_CACHE_LIMIT = 16
_LOOKUP_CACHE: OrderedDict[tuple[Any, ...], tuple[str, str]] = OrderedDict()
_LOOKUP_CACHE_LIMIT = 16_384

MONETARY_RESULT_FIELDS = frozenset({
    "tgtcthue", "tgtthue", "ttcktmai", "tgtphi", "tgtttbso",
    "stckhau", "thtien", "tthue",
})
INVOICE_LEVEL_MONETARY_FIELDS = frozenset({
    "tgtcthue", "tgtthue", "ttcktmai", "tgtphi", "tgtttbso",
})
RECONCILIATION_MONEY_FIELDS = (
    ("tgtcthue", "thtien", "Tổng tiền trước thuế"),
    ("tgtthue", "tthue", "Tiền thuế"),
    ("ttcktmai", "ttcktmai", "Chiết khấu"),
    ("tgtphi", "tgtphi", "Phí"),
    ("tgtttbso", "tgtttbso", "Tổng thanh toán"),
)
PERCENT_RESULT_FIELDS = frozenset({"tsuat"})
NUMBER_RESULT_FIELDS = MONETARY_RESULT_FIELDS | frozenset({"tgia", "dgia", "sluong"})
CONDITIONAL_NUMBER_RESULT_FIELDS = frozenset({"stt", "shdon"})


BLOCKED_RESULT_FIELDS = {
    "password", "token", "authorization", "proxy", "proxy_url",
    "internal_session_id", "session_hash", "raw_json_path",
    "raw_detail_path", "xml_path", "database_path", "db_path",
    "detail_path",
}

ELECTRONIC_TEMPLATE_KEYS = (
    "stt", "khmshdon", "khhdon", "shdon", "tdlap",
    "nbmst", "nbten", "nmmst", "nmten", "nmdchi",
    "tgtcthue", "tgtthue", "ttcktmai", "tgtphi", "tgtttbso",
    "dvtte", "tgia", "tthai", "kqcht",
)
CASH_PURCHASE_TEMPLATE_KEYS = (
    "stt", "khmshdon", "khhdon", "shdon", "tdlap",
    "nbmst", "nbten", "nbdchi", "nmmst", "nmten", "nmcmnd",
    "tgtcthue", "tgtthue", "ttcktmai", "tgtttbso", "tthai", "kqcht",
)
CASH_SOLD_TEMPLATE_KEYS = (
    "stt", "khmshdon", "khhdon", "shdon", "tdlap",
    "nbmst", "nbten", "nmmst", "nmten", "nmdchi", "nmcmnd",
    "tgtcthue", "tgtthue", "ttcktmai", "tgtttbso", "tthai", "kqcht",
)

# openpyxl rejects these XML control characters when assigning a cell. Real
# portal/company text occasionally contains them, so sanitize only at the Excel
# presentation boundary and leave the source-owned SQLite/raw payload untouched.
_EXCEL_ILLEGAL_CHARACTER_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


def _excel_safe_value(value: Any) -> Any:
    if isinstance(value, str):
        text = _EXCEL_ILLEGAL_CHARACTER_RE.sub("", value)
        return "'" + text if text.startswith("=") else text
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(value, ensure_ascii=False, default=str)
        text = _EXCEL_ILLEGAL_CHARACTER_RE.sub("", text)
        return "'" + text if text.startswith("=") else text
    return value


def _excel_safe_record(value: dict[str, Any]) -> dict[str, Any]:
    return {key: _excel_safe_value(field) for key, field in value.items()}


class _ExcelSafeDetailRowBuilder:
    """Delegate source row construction, sanitizing only final Excel cells."""

    def __init__(self) -> None:
        from app.parsers.invoice_detail_excel_row_builder import InvoiceDetailExcelRowBuilder

        self._delegate = InvoiceDetailExcelRowBuilder()

    def build_rows(
        self, payload: dict[str, Any], record: dict[str, Any]
    ) -> list[dict[str, Any]]:
        from mia_invoice_lookup import resolve_invoice_lookup

        lookup = resolve_invoice_lookup(
            payload,
            record,
            available_xml=record.get("xml_path"),
        )
        return [
            _excel_safe_record({
                **row,
                "url": lookup.url or "",
                "mk": lookup.code or "",
            })
            for row in self._delegate.build_rows(payload, record)
        ]


class _LookupEnrichedResultReader:
    """Enrich source result pages from persisted raw detail without recrawling."""

    def __init__(self, database_path: Path) -> None:
        from app.external_api.results import JobResultReader

        self.database_path = Path(database_path)
        self.delegate = JobResultReader(database_path)
        self._revision = _database_revision(self.database_path)

    def overview_page(self, job, *, limit: int, cursor: str | None) -> dict[str, Any]:
        return self.delegate.overview_page(job, limit=limit, cursor=cursor)

    def detail_page(self, job, *, limit: int, cursor: str | None) -> dict[str, Any]:
        self._revision = _database_revision(self.database_path)
        page = self.delegate.detail_page(job, limit=limit, cursor=cursor)
        items = page.get("items") or []
        for item in items:
            if not isinstance(item, dict):
                continue
            url, code = self._lookup(item, job)
            item["url"] = url
            item["mk"] = code
        return page

    def _lookup(self, item: dict[str, Any], job) -> tuple[str, str]:
        identity = tuple(str(item.get(name) or "") for name in (
            "nbmst", "khhdon", "shdon", "khmshdon",
        ))
        cache_key = (
            str(self.database_path.resolve()), self._revision,
            job.company_tax_code, tuple(job.parameters.get("directions") or ()),
            tuple(job.parameters.get("query_types") or ()), *identity,
        )
        cached = _LOOKUP_CACHE.get(cache_key)
        if cached is not None:
            _LOOKUP_CACHE.move_to_end(cache_key)
            return cached

        from mia_invoice_lookup import is_safe_lookup_url, normalize_lookup_key

        existing_url = str(item.get("url") or "").strip()
        existing_code = str(item.get("mk") or "").strip()
        if normalize_lookup_key(existing_code).startswith("khongtimthaymatracu"):
            existing_code = ""
        value = (existing_url if is_safe_lookup_url(existing_url) else "", existing_code)
        if self.database_path.is_file():
            directions = tuple(job.parameters.get("directions") or ())
            query_types = tuple(job.parameters.get("query_types") or ())
            if directions and query_types:
                direction_slots = ",".join("?" for _ in directions)
                query_slots = ",".join("?" for _ in query_types)
                try:
                    with closing(sqlite3.connect(self.database_path)) as connection:
                        connection.row_factory = sqlite3.Row
                        row = connection.execute(f"""
                            SELECT raw_detail_path, nbmst, khhdon, shdon, khmshdon,
                                   direction, query_type
                            FROM invoice_detail_items
                            WHERE company_tax_code=?
                              AND direction IN ({direction_slots})
                              AND query_type IN ({query_slots})
                              AND nbmst=? AND khhdon=? AND shdon=? AND khmshdon=?
                              AND TRIM(raw_detail_path)<>''
                              AND (error_message IS NULL OR TRIM(error_message)='')
                            ORDER BY id DESC LIMIT 1
                        """, (
                            job.company_tax_code, *directions, *query_types, *identity,
                        )).fetchone()
                    if row is not None:
                        raw_path = Path(str(row["raw_detail_path"] or ""))
                        payload = json.loads(raw_path.read_text(encoding="utf-8"))
                        if isinstance(payload, dict):
                            from mia_invoice_lookup import resolve_invoice_lookup

                            lookup = resolve_invoice_lookup(payload, dict(row))
                            if lookup.status != "unresolved":
                                value = (lookup.url or "", lookup.code or "")
                except (OSError, UnicodeError, ValueError, TypeError, sqlite3.Error):
                    value = ("", "")
        _LOOKUP_CACHE[cache_key] = value
        _LOOKUP_CACHE.move_to_end(cache_key)
        while len(_LOOKUP_CACHE) > _LOOKUP_CACHE_LIMIT:
            _LOOKUP_CACHE.popitem(last=False)
        return value


class _FilteredExcelSafeDetailRowBuilder(_ExcelSafeDetailRowBuilder):
    def __init__(self, *, search: str, column_filters: Any) -> None:
        super().__init__()
        self.search = search
        self.column_filters = column_filters

    def build_rows(self, payload: dict[str, Any], record: dict[str, Any]) -> list[dict[str, Any]]:
        rows = super().build_rows(payload, record)
        return [
            row for row in rows
            if _search_matches(row, self.search)
            and _matches_column_filters(row, self.column_filters)
        ]


def _public_field_name(name: str) -> bool:
    key = str(name).casefold()
    return (
        key not in BLOCKED_RESULT_FIELDS
        and not key.startswith("raw_")
        and not key.endswith("_path")
    )


def _safe_fields(item: dict[str, Any]) -> dict[str, Any]:
    value = dict(item)
    attributes = value.pop("attributes", {})
    if isinstance(attributes, dict):
        for key, attribute in attributes.items():
            value.setdefault(str(key), attribute)
    return {
        str(key): field_value
        for key, field_value in value.items()
        if _public_field_name(str(key))
    }


def _search_matches(fields: dict[str, Any], search: str) -> bool:
    if not search:
        return True
    return search in json.dumps(
        fields, ensure_ascii=False, default=str, separators=(",", ":")
    ).casefold()


def _identity_component(value: Any, *, numeric: bool = False) -> str:
    text = str(value if value is not None else "").strip()
    if not numeric or not re.fullmatch(r"[+-]?\d+(?:\.0+)?", text):
        return text
    try:
        return format(Decimal(text).quantize(Decimal("1")), "f")
    except (InvalidOperation, ValueError):
        return text


def _invoice_key(fields: dict[str, Any], direction: str, query_type: str | None = None) -> str:
    """Stable business identity shared by Overview and every Detail line."""
    return "|".join((
        _identity_component(direction),
        _identity_component(query_type or fields.get("query_type") or "query"),
        _identity_component(fields.get("nbmst")),
        _identity_component(fields.get("khhdon")),
        _identity_component(fields.get("shdon"), numeric=True),
        _identity_component(fields.get("khmshdon"), numeric=True),
    ))


def _filter_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str, sort_keys=True)
    return str(value)


def _as_decimal(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        text = re.sub(r"\s+", "", str(value).strip())
        negative_parentheses = text.startswith("(") and text.endswith(")")
        if negative_parentheses:
            text = text[1:-1]
        text = re.sub(r"(?i)(VND|VNĐ|USD|EUR|GBP|JPY|CNY|RMB)", "", text)
        text = text.replace("%", "").replace("₫", "").replace("đ", "").replace("Đ", "")
        text = text.replace("$", "").replace("€", "").replace("£", "").replace("¥", "")
        if negative_parentheses:
            text = f"-{text}"
    if not re.fullmatch(r"[+-]?[0-9.,]+", text):
        return None

    sign = ""
    if text[:1] in {"+", "-"}:
        sign, text = text[0], text[1:]
    dots = text.count(".")
    commas = text.count(",")
    if dots and commas:
        decimal_separator = "." if text.rfind(".") > text.rfind(",") else ","
        grouping_separator = "," if decimal_separator == "." else "."
        text = text.replace(grouping_separator, "")
        text = text.replace(decimal_separator, ".")
    elif dots or commas:
        separator = "." if dots else ","
        parts = text.split(separator)
        grouping = len(parts) > 2 and all(len(part) == 3 for part in parts[1:])
        if len(parts) == 2 and len(parts[1]) == 3:
            grouping = True
        text = "".join(parts) if grouping else "".join(parts[:-1]) + "." + parts[-1]
    text = sign + text
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _column_filter_matches(value: Any, rule: Any) -> bool:
    if not isinstance(rule, dict):
        return True
    text = _filter_text(value)
    normalized = text.casefold()
    values = rule.get("values")
    if isinstance(values, list):
        allowed = {_filter_text(item).casefold() for item in values}
        if normalized not in allowed:
            return False
    search = str(rule.get("search") or "").strip().casefold()
    if search and search not in normalized:
        return False
    operator = str(rule.get("operator") or "")
    operand = rule.get("value")
    if not operator:
        return True
    if operator in {"contains", "not_contains", "starts_with", "ends_with", "equals", "not_equals"}:
        expected = _filter_text(operand).casefold()
        checks = {
            "contains": expected in normalized,
            "not_contains": expected not in normalized,
            "starts_with": normalized.startswith(expected),
            "ends_with": normalized.endswith(expected),
            "equals": normalized == expected,
            "not_equals": normalized != expected,
        }
        return checks[operator]
    actual_number = _as_decimal(value)
    expected_number = _as_decimal(operand)
    if actual_number is None or expected_number is None:
        return False
    if operator == "gt": return actual_number > expected_number
    if operator == "gte": return actual_number >= expected_number
    if operator == "lt": return actual_number < expected_number
    if operator == "lte": return actual_number <= expected_number
    if operator == "number_equals": return actual_number == expected_number
    if operator == "between":
        upper = _as_decimal(rule.get("value_to"))
        return upper is not None and expected_number <= actual_number <= upper
    return True


def _matches_column_filters(fields: dict[str, Any], filters: Any) -> bool:
    if not isinstance(filters, dict):
        return True
    return all(_column_filter_matches(fields.get(str(key)), rule) for key, rule in filters.items())


def _encode_page_cursor(direction_index: int, source_cursor: str | None) -> str:
    payload = json.dumps([1, direction_index, source_cursor], separators=(",", ":"))
    return base64.urlsafe_b64encode(payload.encode("utf-8")).rstrip(b"=").decode("ascii")


def _decode_page_cursor(value: str | None) -> tuple[int, str | None]:
    if not value:
        return 0, None
    try:
        padded = value + "=" * (-len(value) % 4)
        decoded = json.loads(base64.urlsafe_b64decode(padded).decode("utf-8"))
    except Exception as error:
        raise ValueError("invalid_result_cursor") from error
    if (
        not isinstance(decoded, list)
        or len(decoded) != 3
        or decoded[0] != 1
        or not isinstance(decoded[1], int)
        or decoded[1] < 0
        or (decoded[2] is not None and not isinstance(decoded[2], str))
    ):
        raise ValueError("invalid_result_cursor")
    return decoded[1], decoded[2]


def _encode_sorted_cursor(offset: int) -> str:
    payload = json.dumps([2, offset], separators=(",", ":"))
    return base64.urlsafe_b64encode(payload.encode("utf-8")).rstrip(b"=").decode("ascii")


def _decode_sorted_cursor(value: str | None) -> int:
    if not value:
        return 0
    try:
        padded = value + "=" * (-len(value) % 4)
        decoded = json.loads(base64.urlsafe_b64decode(padded).decode("utf-8"))
    except Exception as error:
        raise ValueError("invalid_result_cursor") from error
    if (
        not isinstance(decoded, list)
        or len(decoded) != 2
        or decoded[0] != 2
        or not isinstance(decoded[1], int)
        or decoded[1] < 0
    ):
        raise ValueError("invalid_result_cursor")
    return decoded[1]


def _source_template_dir() -> Path:
    return (
        Path(__file__).resolve().parent
        / "vendor" / "mia_crawl_service" / "resources" / "templates"
    )


def _result_context(backend, query: dict[str, Any]):
    base_job = backend._result_job(str(query["connection_id"]))
    if base_job is None:
        return None

    date_from = str(query.get("date_from") or base_job.parameters["date_from"])
    date_to = str(query.get("date_to") or base_job.parameters["date_to"])
    if date.fromisoformat(date_from) > date.fromisoformat(date_to):
        raise ValueError("invalid_result_range")

    selected_direction = query.get("direction")
    if selected_direction not in (None, "purchase", "sold"):
        raise ValueError("invalid_result_direction")
    base_directions = list(
        base_job.parameters.get("directions") or ("purchase", "sold")
    )
    directions = (
        [str(selected_direction)]
        if selected_direction
        else [
            direction for direction in ("purchase", "sold")
            if direction in set(base_directions)
        ]
    )
    if not directions:
        directions = base_directions

    selected_query_type = query.get("query_type")
    if selected_query_type not in (None, "query", "sco-query"):
        raise ValueError("invalid_result_query_type")
    base_query_types = list(
        base_job.parameters.get("query_types") or ("query", "sco-query")
    )
    query_types = (
        [str(selected_query_type)] if selected_query_type else base_query_types
    )

    database_path = (
        backend.data_root / base_job.company_tax_code / "db" / "invoices.sqlite3"
    )
    return {
        "base_job": base_job,
        "date_from": date_from,
        "date_to": date_to,
        "directions": directions,
        "query_types": query_types,
        "reader": _LookupEnrichedResultReader(database_path),
        "database_path": database_path,
    }


def _job_for(context, direction: str, *, query_types: list[str] | None = None):
    base_job = context["base_job"]
    return replace(
        base_job,
        parameters={
            **base_job.parameters,
            "date_from": context["date_from"],
            "date_to": context["date_to"],
            "directions": [direction],
            "query_types": list(query_types or context["query_types"]),
        },
    )


def _overview_template_keys(category: str, direction: str) -> tuple[str, ...]:
    if category == "electronic":
        return ELECTRONIC_TEMPLATE_KEYS
    if category == "cash_register" and direction == "purchase":
        return CASH_PURCHASE_TEMPLATE_KEYS
    if category == "cash_register" and direction == "sold":
        return CASH_SOLD_TEMPLATE_KEYS
    raise ValueError("invalid_overview_template")


@lru_cache(maxsize=8)
def _overview_template_schema(
    category: str, direction: str
) -> tuple[tuple[str, str], ...]:
    from openpyxl import load_workbook
    from app.services.overview_downloader import OverviewDownloader, _find_header_row

    renderer = OverviewDownloader(
        crawler=None,
        headers_provider=lambda: {},
        template_dir=_source_template_dir(),
    )
    template_path = renderer._template_path(category, direction)
    if not template_path.is_file():
        raise FileNotFoundError(f"Missing source overview template: {template_path}")

    workbook = load_workbook(template_path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        header_row = _find_header_row(sheet)
        titles = [
            str(sheet.cell(header_row, column).value or "").strip()
            for column in range(1, sheet.max_column + 1)
        ]
    finally:
        workbook.close()

    keys = _overview_template_keys(category, direction)
    if len(titles) != len(keys):
        raise RuntimeError(
            f"Source overview template schema mismatch: {category}/{direction} "
            f"has {len(titles)} headers, expected {len(keys)}"
        )
    return tuple((key, title or key) for key, title in zip(keys, titles))


@lru_cache(maxsize=1)
def _detail_template_schema() -> tuple[tuple[str, str], ...]:
    from openpyxl import load_workbook
    from app.exporters.invoice_detail_excel_exporter import InvoiceDetailExcelExporter

    template_path = _source_template_dir() / "invoice_detail.xlsx"
    if not template_path.is_file():
        raise FileNotFoundError(f"Missing source detail template: {template_path}")

    workbook = load_workbook(template_path, read_only=False, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        header_row = InvoiceDetailExcelExporter._find_header_row(sheet)
        if header_row < 5:
            sheet.insert_rows(header_row, amount=5 - header_row)
            header_row = 5
        InvoiceDetailExcelExporter._ensure_serial_number_column(sheet, header_row)
        keys = InvoiceDetailExcelExporter._column_keys(sheet, header_row)
        titles = [
            str(sheet.cell(header_row, column).value or "").strip()
            for column in range(1, sheet.max_column + 1)
        ]
    finally:
        workbook.close()

    if len(keys) != len(titles):
        raise RuntimeError("Source detail template schema mismatch")
    return tuple((key, title or key) for key, title in zip(keys, titles))


def _result_schema(kind: str, context) -> tuple[tuple[str, str], ...]:
    if kind == "details":
        return _detail_template_schema()

    from app.config.crawl_config import query_type_to_category

    query_type = context["query_types"][0]
    context["query_types"] = [query_type]
    category = query_type_to_category(query_type)
    direction = context["directions"][0] if context["directions"] else "purchase"
    return _overview_template_schema(category, direction)


def _overview_display_value(
    fields: dict[str, Any], key: str, *, category: str
) -> Any:
    from app.services.overview_downloader import INVOICE_STATUS_LABELS, OverviewDownloader

    if key == "stt":
        return fields.get("stt")
    if key == "tdlap":
        source = fields.get("tdlap") or fields.get("nlap") or fields.get("nlap_date")
        return OverviewDownloader._format_portal_date(source)
    if key == "tthai":
        value = fields.get(key)
        return INVOICE_STATUS_LABELS.get(value, str(value or ""))
    if key == "kqcht" and category == "cash_register":
        return (
            fields.get(key)
            or "Cục Thuế đã nhận hóa đơn có mã khởi tạo từ máy tính tiền"
        )
    return fields.get(key)


def _project_fields(
    kind: str,
    fields: dict[str, Any],
    schema: tuple[tuple[str, str], ...],
    context,
) -> dict[str, Any]:
    if kind == "details":
        return {key: fields.get(key) for key, _ in schema}

    from app.config.crawl_config import query_type_to_category

    category = query_type_to_category(context["query_types"][0])
    return {
        key: _overview_display_value(fields, key, category=category)
        for key, _ in schema
    }


def _iter_matching_rows(kind: str, context, schema, search: str, column_filters: Any):
    """Stream matching source rows without materializing the dataset in React."""
    for direction in context["directions"]:
        job = _job_for(context, direction)
        cursor = None
        while True:
            page = (
                context["reader"].overview_page(job, limit=200, cursor=cursor)
                if kind == "overview"
                else context["reader"].detail_page(job, limit=200, cursor=cursor)
            )
            for item in page.get("items") or ():
                safe = _safe_fields(item)
                safe["direction"] = direction
                projected = _project_fields(kind, safe, schema, context)
                if _search_matches(safe, search) and _matches_column_filters(projected, column_filters):
                    yield safe, projected, _invoice_key(
                        safe, direction, str(safe.get("query_type") or context["query_types"][0])
                    )
            pagination = page.get("pagination") or {}
            if not pagination.get("has_more") or not pagination.get("next_cursor"):
                break
            cursor = str(pagination["next_cursor"])


def _result_item(safe: dict[str, Any], projected: dict[str, Any], direction: str, context, excluded_keys) -> dict[str, Any]:
    raw_id = safe.get("id")
    if isinstance(raw_id, int):
        row_id: int | str = raw_id
    else:
        fingerprint = json.dumps(
            safe, ensure_ascii=False, default=str, sort_keys=True
        ).encode("utf-8")
        row_id = hashlib.sha256(fingerprint).hexdigest()[:20]
    invoice_key = _invoice_key(
        safe, direction, str(safe.get("query_type") or context["query_types"][0])
    )
    return {
        "row_id": row_id,
        "direction": direction,
        "invoice_key": invoice_key,
        "excluded": invoice_key in excluded_keys,
        "fields": projected,
    }


def _read_sorted_results(kind: str, query: dict[str, Any], context, schema, excluded_keys, analysis, limit: int) -> dict[str, Any]:
    sort = query.get("sort") or {}
    column = str(sort.get("column") or "")
    direction = str(sort.get("direction") or "")
    columns = [key for key, _ in schema]
    if column not in columns or direction not in {"asc", "desc"}:
        raise ValueError("invalid_result_sort")
    search = str(query.get("search") or "").strip().casefold()
    filters = query.get("column_filters") or {}
    rows = list(_iter_matching_rows(kind, context, schema, search, filters))
    raw_populated = [row for row in rows if row[1].get(column) not in (None, "")]
    parsed = [(row, _as_decimal(row[1].get(column))) for row in raw_populated]
    numeric_ratio = sum(value is not None for _row, value in parsed) / max(1, len(parsed))
    numeric_sort = (
        column in NUMBER_RESULT_FIELDS
        or column in PERCENT_RESULT_FIELDS
        or column in CONDITIONAL_NUMBER_RESULT_FIELDS and numeric_ratio >= 0.8
    )
    if numeric_sort:
        populated = [(row, value) for row, value in parsed if value is not None]
        missing = [row for row, value in parsed if value is None]
        missing.extend(row for row in rows if row[1].get(column) in (None, ""))
    else:
        populated = [(row, None) for row in raw_populated]
        missing = [row for row in rows if row[1].get(column) in (None, "")]

    def key(entry):
        row, numeric_value = entry
        return numeric_value if numeric_sort else _filter_text(row[1].get(column)).casefold()

    populated.sort(key=key, reverse=direction == "desc")
    ordered = [row for row, _value in populated] + missing
    offset = _decode_sorted_cursor(query.get("cursor"))
    page_rows = ordered[offset:offset + limit]
    next_offset = offset + len(page_rows)
    has_more = next_offset < len(ordered)
    return {
        "items": [
            _result_item(safe, fields, safe["direction"], context, excluded_keys)
            for safe, fields, _invoice_key_value in page_rows
        ],
        "columns": columns,
        "column_labels": {key: title for key, title in schema},
        "column_types": {
            key: "percent" if key in PERCENT_RESULT_FIELDS else "number"
            for key in columns if key in NUMBER_RESULT_FIELDS or key in PERCENT_RESULT_FIELDS
        },
        "total_count": analysis["matching_row_count"],
        "aggregate": analysis,
        "pagination": {
            "limit": limit,
            "has_more": has_more,
            "next_cursor": _encode_sorted_cursor(next_offset) if has_more else None,
        },
    }


def _bounded_cache(cache: OrderedDict, key: tuple[Any, ...], factory, limit: int):
    cached = cache.get(key)
    if cached is not None:
        cache.move_to_end(key)
        return cached
    value = factory()
    cache[key] = value
    cache.move_to_end(key)
    while len(cache) > limit:
        cache.popitem(last=False)
    return value


def _query_cache_key(kind: str, context, search: str, column_filters: Any) -> tuple[Any, ...]:
    database_path = Path(context["database_path"])
    return (
        kind, str(database_path.resolve()), _database_revision(database_path),
        context["base_job"].company_tax_code, context["date_from"], context["date_to"],
        tuple(context["directions"]), tuple(context["query_types"]), search,
        json.dumps(column_filters or {}, ensure_ascii=False, sort_keys=True, default=str),
    )


def _resolve_excluded_keys(backend, exclusion: Any) -> frozenset[str]:
    if not isinstance(exclusion, dict):
        return frozenset()
    explicit = {
        str(value) for value in exclusion.get("keys") or ()
        if isinstance(value, str) and 1 <= len(value) <= 512
    }
    rules = exclusion.get("rules") or []
    if not isinstance(rules, list) or not rules:
        return frozenset(explicit)
    key = (
        "exclusion", json.dumps(rules, ensure_ascii=False, sort_keys=True, default=str),
        tuple(sorted(explicit)),
    )

    def build() -> frozenset[str]:
        output = set(explicit)
        for rule in rules:
            if not isinstance(rule, dict) or rule.get("kind") not in {"overview", "details"}:
                continue
            query = dict(rule.get("query") or {})
            context = _result_context(backend, query)
            if context is None:
                continue
            schema = _result_schema(str(rule["kind"]), context)
            search = str(query.get("search") or "").strip().casefold()
            except_keys = {str(value) for value in rule.get("except_keys") or ()}
            for _safe, _fields, invoice_key in _iter_matching_rows(
                str(rule["kind"]), context, schema, search, query.get("column_filters")
            ):
                if invoice_key not in except_keys:
                    output.add(invoice_key)
        return frozenset(output)

    return _bounded_cache(_EXCLUSION_CACHE, key, build, _EXCLUSION_CACHE_LIMIT)


def _result_analysis(backend, kind: str, context, schema, search: str, column_filters: Any, exclusion: Any):
    excluded = _resolve_excluded_keys(backend, exclusion)
    key = _query_cache_key(kind, context, search, column_filters) + (
        hashlib.sha256("\n".join(sorted(excluded)).encode("utf-8")).hexdigest(),
    )

    def build() -> dict[str, Any]:
        matching_row_count = 0
        row_count = 0
        invoice_keys: set[str] = set()
        totals = {field: Decimal("0") for field in MONETARY_RESULT_FIELDS}
        for _safe, fields, invoice_key in _iter_matching_rows(
            kind, context, schema, search, column_filters
        ):
            matching_row_count += 1
            if invoice_key in excluded:
                continue
            row_count += 1
            first_invoice_row = invoice_key not in invoice_keys
            invoice_keys.add(invoice_key)
            for field in MONETARY_RESULT_FIELDS:
                if kind == "details" and field in INVOICE_LEVEL_MONETARY_FIELDS and not first_invoice_row:
                    continue
                number = _as_decimal(fields.get(field))
                if number is not None:
                    totals[field] += number
        return {
            "matching_row_count": matching_row_count,
            "row_count": row_count,
            "invoice_count": len(invoice_keys),
            "totals": {
                field: int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
                for field, value in totals.items()
            },
        }

    return _bounded_cache(_RESULT_ANALYSIS_CACHE, key, build, _RESULT_ANALYSIS_CACHE_LIMIT), excluded


def read_results(backend, kind: str, query: dict[str, Any]) -> dict[str, Any]:
    if kind not in {"overview", "details"}:
        raise ValueError("invalid_result_kind")
    limit = int(query.get("limit", 50))
    if not 1 <= limit <= 50:
        raise ValueError("invalid_result_limit")

    context = _result_context(backend, query)
    if context is None:
        return {
            "items": [], "columns": [], "column_labels": {},
            "total_count": 0,
            "pagination": {"limit": limit, "has_more": False, "next_cursor": None},
        }

    schema = _result_schema(kind, context)
    columns = [key for key, _ in schema]
    column_labels = {key: title for key, title in schema}
    search = str(query.get("search") or "").strip().casefold()
    column_filters = query.get("column_filters") or {}
    analysis, excluded_keys = _result_analysis(
        backend, kind, context, schema, search, column_filters, query.get("exclusion")
    )
    if query.get("sort"):
        return _read_sorted_results(
            kind, query, context, schema, excluded_keys, analysis, limit
        )
    direction_index, source_cursor = _decode_page_cursor(query.get("cursor"))
    directions = context["directions"]

    if direction_index >= len(directions):
        return {
            "items": [], "columns": columns, "column_labels": column_labels,
            "total_count": analysis["matching_row_count"], "aggregate": analysis,
            "pagination": {"limit": limit, "has_more": False, "next_cursor": None},
        }

    output: list[dict[str, Any]] = []
    next_direction_index = direction_index
    next_source_cursor = source_cursor
    has_more = False

    while len(output) < limit and next_direction_index < len(directions):
        direction = directions[next_direction_index]
        job = _job_for(context, direction)
        remaining = max(1, limit - len(output))
        page = (
            context["reader"].overview_page(
                job, limit=remaining, cursor=next_source_cursor
            )
            if kind == "overview"
            else context["reader"].detail_page(
                job, limit=remaining, cursor=next_source_cursor
            )
        )

        for item in page.get("items") or ():
            safe = _safe_fields(item)
            safe["direction"] = direction
            projected = _project_fields(kind, safe, schema, context)
            if not _search_matches(safe, search) or not _matches_column_filters(projected, column_filters):
                continue
            output.append(_result_item(safe, projected, direction, context, excluded_keys))

        pagination = page.get("pagination") or {}
        page_more = bool(pagination.get("has_more"))
        page_cursor = pagination.get("next_cursor")
        if page_more and page_cursor:
            next_source_cursor = str(page_cursor)
            has_more = True
            if len(output) >= limit:
                break
            continue

        next_direction_index += 1
        next_source_cursor = None
        has_more = next_direction_index < len(directions)

    response: dict[str, Any] = {
        "items": output[:limit],
        "columns": columns,
        "column_labels": column_labels,
        "pagination": {
            "limit": limit,
            "has_more": bool(has_more),
            "next_cursor": (
                _encode_page_cursor(next_direction_index, next_source_cursor)
                if has_more else None
            ),
        },
    }

    response["total_count"] = analysis["matching_row_count"]
    response["aggregate"] = analysis
    response["column_types"] = {
        key: "percent" if key in PERCENT_RESULT_FIELDS else "number"
        for key in columns if key in NUMBER_RESULT_FIELDS or key in PERCENT_RESULT_FIELDS
    }
    return response


_RECONCILIATION_COLUMNS = (
    "stt", "reconciliation_status", "khmshdon", "khhdon", "shdon", "tdlap",
    "nbmst", "nbten", "nmmst", "nmten", "mismatch_fields",
    "overview_tgtcthue", "detail_thtien", "difference_tgtcthue",
    "overview_tgtthue", "detail_tthue", "difference_tgtthue",
    "overview_ttcktmai", "detail_ttcktmai", "difference_ttcktmai",
    "overview_tgtphi", "detail_tgtphi", "difference_tgtphi",
    "overview_tgtttbso", "detail_tgtttbso", "difference_tgtttbso",
)
_RECONCILIATION_EXPORT_COLUMNS = (
    "stt", "reconciliation_status", "reconciliation_reason", "mismatch_fields",
    "khmshdon", "khhdon", "shdon", "tdlap",
    "nbmst", "nbten", "nmmst", "nmten",
    *tuple(
        column
        for columns in (
            (
                f"overview_{overview_field}", f"detail_{detail_field}",
                f"difference_{overview_field}",
            )
            for overview_field, detail_field, _label in RECONCILIATION_MONEY_FIELDS
        )
        for column in columns
    ),
)
_RECONCILIATION_LABELS = {
    "stt": "STT",
    "reconciliation_status": "Trạng thái đối chiếu",
    "khmshdon": "Ký hiệu mẫu số",
    "khhdon": "Ký hiệu hóa đơn",
    "shdon": "Số hóa đơn",
    "tdlap": "Ngày lập",
    "nbmst": "MST người bán/người xuất hàng",
    "nbten": "Tên người bán/người xuất hàng",
    "nmmst": "MST người mua",
    "nmten": "Tên người mua",
    "mismatch_fields": "Chỉ tiêu chênh lệch",
    "reconciliation_reason": "Lý do chênh lệch",
}
for _overview_field, _detail_field, _label in RECONCILIATION_MONEY_FIELDS:
    _RECONCILIATION_LABELS[f"overview_{_overview_field}"] = f"{_label} - Tổng quan"
    _RECONCILIATION_LABELS[f"detail_{_detail_field}"] = f"{_label} - Chi tiết"
    _RECONCILIATION_LABELS[f"difference_{_overview_field}"] = f"{_label} - Chênh lệch"


def _decimal_or_zero(value: Any) -> Decimal:
    return _as_decimal(value) or Decimal("0")


def _canonical_decimal(value: Decimal) -> Decimal:
    """Remove Decimal's signed zero without changing any non-zero amount."""
    return Decimal("0") if value.is_zero() else value


def _normalize_reconciliation_money(value: Any) -> Decimal | None:
    """Normalize source totals to the VND unit displayed by the result report.

    Detail totals are sums of persisted line values and can retain microscopic
    decimal residue (for example ``-3E-9``) even when the invoice totals are
    equal to the đồng.  Reconciliation is a VND report, so comparison must be
    performed after both sides are quantized to one đồng, not merely formatted
    that way in React.
    """
    parsed = _as_decimal(value)
    if parsed is None:
        return None
    return _canonical_decimal(parsed.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _compare_money_values(overview: Any, detail: Any) -> dict[str, Any]:
    overview_value = _normalize_reconciliation_money(overview)
    detail_value = _normalize_reconciliation_money(detail)
    difference = (
        _canonical_decimal(overview_value - detail_value)
        if overview_value is not None and detail_value is not None
        else None
    )
    return {
        "overview_value": overview_value,
        "detail_value": detail_value,
        "difference": difference,
        "is_mismatch": difference is not None and difference != Decimal("0"),
    }


def _decimal_result(value: Decimal) -> int | str:
    value = _canonical_decimal(value)
    integral = value.to_integral_value()
    return int(integral) if value == integral else format(value.normalize(), "f")


def _reason_money(value: Decimal) -> str:
    safe = _decimal_result(abs(_canonical_decimal(value)))
    if isinstance(safe, int):
        return f"{safe:,}".replace(",", ".")
    return str(safe).replace(".", ",")


def _money_mismatch_reason(
    differences: list[tuple[str, Decimal, Decimal, Decimal]],
) -> str:
    explanations: list[str] = []
    for label, overview, detail, difference in differences:
        relation = "lớn hơn" if difference > 0 else "thấp hơn"
        explanations.append(
            f"{label} bên Tổng quan {relation} Chi tiết "
            f"{_reason_money(difference)} đồng"
        )
    return "; ".join(explanations) + "."


def _merge_date_ranges(ranges: list[tuple[date, date]]) -> list[tuple[date, date]]:
    merged: list[tuple[date, date]] = []
    for begin, end in sorted(ranges):
        if begin > end:
            continue
        if merged and begin <= merged[-1][1] + timedelta(days=1):
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((begin, end))
    return merged


def _intersect_date_ranges(
    left: list[tuple[date, date]], right: list[tuple[date, date]]
) -> list[tuple[date, date]]:
    output: list[tuple[date, date]] = []
    left = _merge_date_ranges(left)
    right = _merge_date_ranges(right)
    left_index = right_index = 0
    while left_index < len(left) and right_index < len(right):
        begin = max(left[left_index][0], right[right_index][0])
        end = min(left[left_index][1], right[right_index][1])
        if begin <= end:
            output.append((begin, end))
        if left[left_index][1] <= right[right_index][1]:
            left_index += 1
        else:
            right_index += 1
    return output


def _subtract_date_ranges(
    ranges: list[tuple[date, date]], blocked: list[tuple[date, date]]
) -> list[tuple[date, date]]:
    output = _merge_date_ranges(ranges)
    for blocked_from, blocked_to in _merge_date_ranges(blocked):
        next_output: list[tuple[date, date]] = []
        for begin, end in output:
            if blocked_to < begin or blocked_from > end:
                next_output.append((begin, end))
                continue
            if begin < blocked_from:
                next_output.append((begin, blocked_from - timedelta(days=1)))
            if blocked_to < end:
                next_output.append((blocked_to + timedelta(days=1), end))
        output = next_output
    return output


def _job_module_ranges(job, module_name: str) -> list[tuple[date, date]]:
    state = dict(getattr(job, "progress_state", None) or {})
    module = (state.get("modules") or {}).get(module_name) or {}
    if module.get("status") != "completed":
        return []
    output: list[tuple[date, date]] = []
    for month in module.get("months") or ():
        if not isinstance(month, dict) or month.get("status") != "completed":
            continue
        try:
            output.append((
                date.fromisoformat(str(month["from_date"])),
                date.fromisoformat(str(month["to_date"])),
            ))
        except (KeyError, TypeError, ValueError):
            continue
    return output


def _job_requested_ranges(job) -> list[tuple[date, date]]:
    parameters = dict(getattr(job, "parameters", None) or {})
    try:
        return [(
            date.fromisoformat(str(parameters["date_from"])),
            date.fromisoformat(str(parameters["date_to"])),
        )]
    except (KeyError, TypeError, ValueError):
        return []


def _reconciliation_coverage(backend, context) -> list[tuple[date, date]]:
    """Source-confirmed 2/2 coverage, excluding any active refresh range."""
    requested = (
        date.fromisoformat(context["date_from"]),
        date.fromisoformat(context["date_to"]),
    )
    account_key = str(context["base_job"].account_key)
    selected_directions = set(context["directions"])
    selected_query_types = set(context["query_types"])
    overview: list[tuple[date, date]] = []
    details: list[tuple[date, date]] = []
    active: list[tuple[date, date]] = []
    jobs = (
        backend.repository.invoice_jobs_for_account(
            account_key, owner_id=getattr(context["base_job"], "owner_id", None)
        )
        if hasattr(backend.repository, "invoice_jobs_for_account")
        else backend.repository.list_jobs_for_reconciliation()
    )
    for job in jobs:
        if str(getattr(job, "account_key", "")) != account_key:
            continue
        parameters = dict(getattr(job, "parameters", None) or {})
        if not selected_directions.intersection(parameters.get("directions") or ()):
            continue
        if not selected_query_types.intersection(parameters.get("query_types") or ()):
            continue
        status = str(getattr(job, "status", ""))
        if status in {"queued", "waiting_account", "running", "cancelling"}:
            active.extend(_job_requested_ranges(job))
            continue
        if status != "completed":
            continue
        overview.extend(_job_module_ranges(job, "overview"))
        details.extend(_job_module_ranges(job, "detail"))
    coverage = _intersect_date_ranges(overview, details)
    coverage = _intersect_date_ranges(coverage, [requested])
    return _subtract_date_ranges(coverage, active)


def _invoice_business_date(fields: dict[str, Any]) -> date | None:
    from app.utils.date_utils import normalize_business_date

    for field in ("nlap_date", "tdlap", "ntao", "nlap"):
        normalized = normalize_business_date(fields.get(field))
        if normalized:
            return date.fromisoformat(normalized)
    return None


def _inside_coverage(value: date | None, coverage: list[tuple[date, date]]) -> bool:
    return value is not None and any(begin <= value <= end for begin, end in coverage)


def _reconciliation_cache_key(backend, context) -> tuple[Any, ...]:
    database_path = Path(context["database_path"])
    control_value = getattr(backend, "control_db", None)
    control_revision = _database_revision(Path(control_value)) if control_value else ()
    return (
        "reconciliation", _RECONCILIATION_ALGORITHM_VERSION,
        str(database_path.resolve()), _database_revision(database_path),
        control_revision,
        context["base_job"].company_tax_code, context["date_from"], context["date_to"],
        tuple(context["directions"]), tuple(context["query_types"]),
    )


def _reconciliation_dataset(backend, query: dict[str, Any]) -> dict[str, Any]:
    context = _result_context(backend, query)
    if context is None:
        return {
            "items": [],
            "summary": {
                "overview_invoice_count": 0, "detail_invoice_count": 0,
                "difference": 0, "missing_detail_count": 0,
                "missing_overview_count": 0, "money_mismatch_count": 0,
                "issue_count": 0, "coverage_ranges": [],
            },
        }
    key = _reconciliation_cache_key(backend, context)

    def build() -> dict[str, Any]:
        # Reconciliation needs persisted invoice/result fields only. Bypass the
        # lookup enrichment wrapper so a full-scope comparison never opens raw
        # detail JSON merely to resolve URL/code presentation fields.
        from app.external_api.results import JobResultReader

        coverage = _reconciliation_coverage(backend, context)
        source_context = {
            **context,
            "reader": JobResultReader(context["database_path"]),
        }
        overview_by_key: dict[str, tuple[dict[str, Any], dict[str, Any]]] = {}
        detail_by_key: dict[str, dict[str, Any]] = {}
        query_types = list(source_context["query_types"])
        for query_type in query_types:
            scoped = {**source_context, "query_types": [query_type]}
            overview_schema = _result_schema("overview", scoped)
            detail_schema = _result_schema("details", scoped)
            for safe, projected, invoice_key in _iter_matching_rows(
                "overview", scoped, overview_schema, "", {}
            ):
                if not _inside_coverage(_invoice_business_date(safe), coverage):
                    continue
                safe["query_type"] = query_type
                invoice_key = _invoice_key(safe, safe["direction"], query_type)
                overview_by_key.setdefault(invoice_key, (safe, projected))
            for safe, projected, invoice_key in _iter_matching_rows(
                "details", scoped, detail_schema, "", {}
            ):
                if not _inside_coverage(_invoice_business_date(safe), coverage):
                    continue
                safe["query_type"] = query_type
                invoice_key = _invoice_key(safe, safe["direction"], query_type)
                group = detail_by_key.setdefault(invoice_key, {
                    "safe": safe, "projected": projected,
                    "line_totals": {"thtien": Decimal("0"), "tthue": Decimal("0")},
                    "invoice_totals": {},
                })
                for field in ("thtien", "tthue"):
                    group["line_totals"][field] += _decimal_or_zero(projected.get(field))
                for field in ("ttcktmai", "tgtphi", "tgtttbso"):
                    if field not in group["invoice_totals"] and projected.get(field) not in (None, ""):
                        group["invoice_totals"][field] = _decimal_or_zero(projected.get(field))

        items: list[dict[str, Any]] = []
        missing_detail_count = 0
        missing_overview_count = 0
        money_mismatch_count = 0
        for invoice_key in sorted(set(overview_by_key) | set(detail_by_key)):
            overview_entry = overview_by_key.get(invoice_key)
            detail_entry = detail_by_key.get(invoice_key)
            if overview_entry is None:
                status = "Thiếu tổng quan"
                missing_overview_count += 1
            elif detail_entry is None:
                status = "Thiếu chi tiết"
                missing_detail_count += 1
            else:
                status = ""

            overview_safe, overview_fields = overview_entry or ({}, {})
            detail_safe = detail_entry["safe"] if detail_entry else {}
            detail_fields = detail_entry["projected"] if detail_entry else {}
            fields = {
                key_name: overview_fields.get(key_name) or detail_fields.get(key_name)
                for key_name in (
                    "khmshdon", "khhdon", "shdon", "nbmst", "nbten", "nmmst", "nmten"
                )
            }
            fields["tdlap"] = (
                overview_fields.get("tdlap") or detail_fields.get("ntao")
                or overview_safe.get("nlap") or detail_safe.get("nlap")
            )
            mismatches: list[str] = []
            money_differences: list[tuple[str, Decimal, Decimal, Decimal]] = []
            for overview_field, detail_field, label in RECONCILIATION_MONEY_FIELDS:
                overview_value = (
                    _decimal_or_zero(overview_fields.get(overview_field))
                    if overview_entry is not None else None
                )
                if detail_entry is None:
                    detail_value = None
                elif detail_field in {"thtien", "tthue"}:
                    detail_value = detail_entry["line_totals"][detail_field]
                else:
                    detail_value = detail_entry["invoice_totals"].get(detail_field, Decimal("0"))
                comparison = _compare_money_values(overview_value, detail_value)
                overview_value = comparison["overview_value"]
                detail_value = comparison["detail_value"]
                difference = comparison["difference"]
                fields[f"overview_{overview_field}"] = (
                    _decimal_result(overview_value) if overview_value is not None else None
                )
                fields[f"detail_{detail_field}"] = (
                    _decimal_result(detail_value) if detail_value is not None else None
                )
                fields[f"difference_{overview_field}"] = (
                    _decimal_result(difference) if difference is not None else None
                )
                # Some source Overview templates intentionally omit a field
                # (for example ``tgtphi`` on cash-register invoices). Compare
                # only values represented by both source schemas; absence is
                # not equivalent to a financial zero supplied by the portal.
                if (
                    overview_entry is not None
                    and detail_entry is not None
                    and overview_field in overview_fields
                    and comparison["is_mismatch"]
                ):
                    mismatches.append(label)
                    money_differences.append(
                        (label, overview_value, detail_value, difference)
                    )
            if not status and mismatches:
                status = "Chênh lệch tiền"
                money_mismatch_count += 1
            if not status:
                continue
            fields["reconciliation_status"] = status
            fields["mismatch_fields"] = ", ".join(mismatches) or "—"
            if status == "Thiếu chi tiết":
                fields["reconciliation_reason"] = (
                    "Hóa đơn có trong Tổng quan nhưng không tìm thấy dữ liệu Chi tiết "
                    "sau khi phạm vi này đã đồng bộ đầy đủ Tổng quan và Chi tiết."
                )
            elif status == "Thiếu tổng quan":
                fields["reconciliation_reason"] = (
                    "Hóa đơn có trong Chi tiết nhưng không tìm thấy hóa đơn tương ứng "
                    "trong Tổng quan sau khi phạm vi này đã đồng bộ đầy đủ."
                )
            else:
                fields["reconciliation_reason"] = _money_mismatch_reason(
                    money_differences
                )
            items.append({
                "row_id": hashlib.sha256(invoice_key.encode("utf-8")).hexdigest()[:20],
                "direction": overview_safe.get("direction") or detail_safe.get("direction") or "purchase",
                "invoice_key": invoice_key,
                "excluded": False,
                "fields": fields,
            })

        overview_count = len(overview_by_key)
        detail_count = len(detail_by_key)
        return {
            "items": items,
            "summary": {
                "overview_invoice_count": overview_count,
                "detail_invoice_count": detail_count,
                "difference": overview_count - detail_count,
                "missing_detail_count": missing_detail_count,
                "missing_overview_count": missing_overview_count,
                "money_mismatch_count": money_mismatch_count,
                "issue_count": len(items),
                "coverage_ranges": [
                    {"date_from": begin.isoformat(), "date_to": end.isoformat()}
                    for begin, end in coverage
                ],
            },
        }

    return _bounded_cache(_RECONCILIATION_CACHE, key, build, _RECONCILIATION_CACHE_LIMIT)


def _filtered_reconciliation_items(dataset: dict[str, Any], query: dict[str, Any]) -> list[dict[str, Any]]:
    search = str(query.get("search") or "").strip().casefold()
    filters = query.get("column_filters") or {}
    items = [
        item for item in dataset["items"]
        if _search_matches(item["fields"], search)
        and _matches_column_filters(item["fields"], filters)
    ]
    sort = query.get("sort") or {}
    column = str(sort.get("column") or "")
    direction = str(sort.get("direction") or "")
    if column in _RECONCILIATION_COLUMNS and direction in {"asc", "desc"}:
        def sort_key(item):
            value = item["fields"].get(column)
            numeric = _as_decimal(value)
            return (value in (None, ""), numeric is None, numeric if numeric is not None else _filter_text(value).casefold())
        items.sort(key=sort_key, reverse=direction == "desc")
    return items


def read_reconciliation(backend, query: dict[str, Any]) -> dict[str, Any]:
    limit = int(query.get("limit", 50))
    if not 1 <= limit <= 50:
        raise ValueError("invalid_result_limit")
    dataset = _reconciliation_dataset(backend, query)
    items = _filtered_reconciliation_items(dataset, query)
    offset = _decode_sorted_cursor(query.get("cursor"))
    page_items = items[offset:offset + limit]
    for index, item in enumerate(page_items, start=offset + 1):
        item = {**item, "fields": {**item["fields"], "stt": index}}
        page_items[index - offset - 1] = item
    next_offset = offset + len(page_items)
    has_more = next_offset < len(items)
    money_columns = {
        column for column in _RECONCILIATION_COLUMNS
        if column.startswith(("overview_", "detail_", "difference_"))
    }
    totals = {
        column: _decimal_result(sum(
            (_decimal_or_zero(item["fields"].get(column)) for item in items),
            Decimal("0"),
        ))
        for column in money_columns
    }
    return {
        "items": page_items,
        "columns": list(_RECONCILIATION_COLUMNS),
        "column_labels": dict(_RECONCILIATION_LABELS),
        "column_types": {column: "number" for column in money_columns},
        "total_count": len(items),
        "aggregate": {
            "matching_row_count": len(items), "row_count": len(items),
            "invoice_count": len(items), "totals": totals,
        },
        "reconciliation": dataset["summary"],
        "pagination": {
            "limit": limit, "has_more": has_more,
            "next_cursor": _encode_sorted_cursor(next_offset) if has_more else None,
        },
    }


def read_result_facets(backend, query: dict[str, Any]) -> dict[str, Any]:
    kind = str(query.get("kind") or "")
    column = str(query.get("column") or "")
    if kind not in {"overview", "details", "reconciliation"} or not column:
        raise ValueError("invalid_result_facet")
    if kind == "reconciliation":
        if column not in _RECONCILIATION_COLUMNS:
            raise ValueError("invalid_result_facet")
        dataset = _reconciliation_dataset(backend, query)
        filters = dict(query.get("column_filters") or {})
        filters.pop(column, None)
        values: dict[str, Any] = {}
        limit = max(1, min(500, int(query.get("facet_limit") or 250)))
        for item in dataset["items"]:
            fields = item["fields"]
            if not _search_matches(fields, str(query.get("search") or "").strip().casefold()):
                continue
            if not _matches_column_filters(fields, filters):
                continue
            value = fields.get(column)
            values.setdefault(_filter_text(value).casefold(), value)
        ordered = sorted(values.values(), key=lambda value: _filter_text(value).casefold())
        numeric = column.startswith(("overview_", "detail_", "difference_"))
        return {
            "values": ordered[:limit], "truncated": len(ordered) > limit,
            "column_type": "number" if numeric else "text",
        }
    context = _result_context(backend, query)
    if context is None:
        return {"values": [], "truncated": False, "column_type": "text"}
    schema = _result_schema(kind, context)
    columns = {key for key, _label in schema}
    if column not in columns:
        raise ValueError("invalid_result_facet")
    filters = dict(query.get("column_filters") or {})
    filters.pop(column, None)
    search = str(query.get("search") or "").strip().casefold()
    limit = max(1, min(500, int(query.get("facet_limit") or 250)))
    values: dict[str, Any] = {}
    truncated = False
    for _safe, fields, _invoice_key_value in _iter_matching_rows(kind, context, schema, search, filters):
        value = fields.get(column)
        normalized = _filter_text(value).casefold()
        values.setdefault(normalized, value)
        if len(values) > limit:
            truncated = True
            break
    ordered = sorted(values.values(), key=lambda value: _filter_text(value).casefold())[:limit]
    column_type = "percent" if column in PERCENT_RESULT_FIELDS else (
        "number" if column in NUMBER_RESULT_FIELDS else "text"
    )
    return {"values": ordered, "truncated": truncated, "column_type": column_type}


def _all_overview_fields(
    context,
    direction: str,
    query_type: str,
    search: str,
    column_filters: Any = None,
    excluded_keys: frozenset[str] = frozenset(),
) -> list[dict[str, Any]]:
    from app.config.crawl_config import query_type_to_category

    job = _job_for(context, direction, query_types=[query_type])
    category = query_type_to_category(query_type)
    schema = _overview_template_schema(category, direction)
    projection_context = {
        **context, "query_types": [query_type], "directions": [direction]
    }
    cursor = None
    rows: list[dict[str, Any]] = []
    while True:
        page = context["reader"].overview_page(job, limit=200, cursor=cursor)
        for item in page.get("items") or ():
            fields = _safe_fields(item)
            fields["direction"] = direction
            projected = _project_fields("overview", fields, schema, projection_context)
            if (
                _search_matches(fields, search)
                and _matches_column_filters(projected, column_filters)
                and _invoice_key(fields, direction, query_type) not in excluded_keys
            ):
                fields.setdefault(
                    "tdlap", fields.get("nlap") or fields.get("nlap_date")
                )
                rows.append(fields)
        pagination = page.get("pagination") or {}
        if not pagination.get("has_more") or not pagination.get("next_cursor"):
            return rows
        cursor = str(pagination["next_cursor"])


def read_artifact_targets(backend, query: dict[str, Any]) -> list[dict[str, Any]]:
    """Return source Overview identities for one filtered artifact batch.

    This internal adapter deliberately uses the same result context and search
    matching as the Results table, but retains the canonical source date fields
    required by package persistence.  It never exposes these payloads to the
    renderer.
    """
    context = _result_context(backend, query)
    if context is None:
        return []
    search = str(query.get("search") or "").strip().casefold()
    targets: list[dict[str, Any]] = []
    for direction in context["directions"]:
        for query_type in context["query_types"]:
            for fields in _all_overview_fields(context, direction, query_type, search):
                target = {
                    "direction": direction,
                    "query_type": query_type,
                    "nbmst": str(fields.get("nbmst") or ""),
                    "khhdon": str(fields.get("khhdon") or ""),
                    "shdon": str(fields.get("shdon") or ""),
                    "khmshdon": str(fields.get("khmshdon") or ""),
                    "nlap": fields.get("nlap") or fields.get("tdlap"),
                    "nlap_date": fields.get("nlap_date"),
                    "partner_name": fields.get("nbten") if direction == "purchase" else fields.get("nmten"),
                }
                target["artifact_key"] = "|".join(str(target[name]) for name in (
                    "direction", "query_type", "nbmst", "khhdon", "shdon", "khmshdon",
                ))
                targets.append(target)
    return targets


def _available_path(destination: Path, filename: str) -> Path:
    candidate = destination / filename
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    for copy_index in range(2, 1001):
        if not candidate.exists():
            return candidate
        candidate = destination / f"{stem} ({copy_index}){suffix}"
    raise OSError("artifact_name_exhausted")


def _grouped_result_filename(
    company_tax_code: str,
    direction: str,
    scope: str,
    date_from: str,
    date_to: str,
) -> str:
    safe_tax_code = re.sub(r"[^0-9A-Za-z._-]+", "_", company_tax_code).strip("._-")
    if not safe_tax_code:
        safe_tax_code = "MIA"
    direction_label = (
        "Mua vào" if direction == "purchase" else
        "Bán ra" if direction == "sold" else "Mua vào và Bán ra"
    )
    scope_label = {
        "overview": "Tổng quan", "details": "Chi tiết",
        "reconciliation": "Đối chiếu Tổng quan và Chi tiết",
    }[scope]
    return (
        f"{safe_tax_code} - {direction_label} - {scope_label} - "
        f"{date_from}_{date_to}.xlsx"
    )


def _result_output_directory(
    destination: Path, company_tax_code: str, scope: str,
    date_from: str, date_to: str,
) -> Path:
    safe_tax_code = re.sub(
        r"[^0-9A-Za-z._-]+", "_", company_tax_code
    ).strip("._-") or "MIA"
    scope_label = {
        "overview": "Tổng quan", "details": "Chi tiết",
        "reconciliation": "Đối chiếu Tổng quan và Chi tiết",
    }[scope]
    return (
        destination / safe_tax_code
        / f"{scope_label} {date_from}_{date_to}"
    )


def _database_revision(database_path: Path) -> tuple[tuple[int, int], ...]:
    """Track SQLite and WAL writes so cached presentation counts stay fresh."""
    revision: list[tuple[int, int]] = []
    for path in (database_path, Path(f"{database_path}-wal")):
        try:
            stat = path.stat()
        except OSError:
            revision.append((0, 0))
        else:
            revision.append((stat.st_mtime_ns, stat.st_size))
    return tuple(revision)


def _cached_result_count(key: tuple[Any, ...], factory: Callable[[], int]) -> int:
    cached = _RESULT_COUNT_CACHE.get(key)
    if cached is not None:
        _RESULT_COUNT_CACHE.move_to_end(key)
        return cached
    value = factory()
    _RESULT_COUNT_CACHE[key] = value
    _RESULT_COUNT_CACHE.move_to_end(key)
    while len(_RESULT_COUNT_CACHE) > _RESULT_COUNT_CACHE_LIMIT:
        _RESULT_COUNT_CACHE.popitem(last=False)
    return value


def _normalized_detail_row_count(context, job) -> int | None:
    """Return the displayed line count when all available details are normalized.

    Mixed/legacy datasets return ``None`` so the caller can delegate row
    expansion to the authoritative source reader instead of guessing from the
    invoice count.
    """
    database_path = Path(context["database_path"])
    if not database_path.is_file():
        return None
    directions = tuple(job.parameters["directions"])
    query_types = tuple(job.parameters["query_types"])
    direction_slots = ",".join("?" for _ in directions)
    query_slots = ",".join("?" for _ in query_types)
    params = [
        job.company_tax_code,
        job.parameters["date_from"],
        job.parameters["date_to"],
        *directions,
        *query_types,
    ]
    try:
        with sqlite3.connect(database_path) as connection:
            columns = {
                str(row[1])
                for row in connection.execute("PRAGMA table_info(invoice_detail_items)")
            }
            line_table = connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' "
                "AND name='invoice_detail_lines'"
            ).fetchone()
            if "normalized_ready" not in columns or not line_table:
                return None
            normalized, legacy = connection.execute(f"""
                SELECT
                    COALESCE(SUM(CASE WHEN normalized_ready=1 THEN 1 ELSE 0 END), 0),
                    COALESCE(SUM(CASE WHEN normalized_ready<>1 THEN 1 ELSE 0 END), 0)
                FROM invoice_detail_items
                WHERE company_tax_code=? AND nlap_date BETWEEN ? AND ?
                  AND direction IN ({direction_slots})
                  AND query_type IN ({query_slots})
                  AND (error_message IS NULL OR TRIM(error_message)='')
                  AND (normalized_ready=1 OR TRIM(raw_detail_path)<>'')
            """, params).fetchone()
            if int(legacy) > 0:
                return None
            if int(normalized) == 0:
                return 0
            row = connection.execute(f"""
                SELECT COUNT(*)
                FROM invoice_detail_items d
                JOIN invoice_detail_lines l ON l.detail_item_id=d.id
                WHERE d.company_tax_code=? AND d.nlap_date BETWEEN ? AND ?
                  AND d.direction IN ({direction_slots})
                  AND d.query_type IN ({query_slots})
                  AND d.normalized_ready=1
                  AND (d.error_message IS NULL OR TRIM(d.error_message)='')
            """, params).fetchone()
            return int(row[0])
    except sqlite3.Error:
        return None


def _result_total_count(kind: str, context, search: str) -> int:
    database_path = Path(context["database_path"])
    base_job = context["base_job"]
    key = (
        kind,
        str(database_path.resolve()),
        _database_revision(database_path),
        base_job.company_tax_code,
        context["date_from"],
        context["date_to"],
        tuple(context["directions"]),
        tuple(context["query_types"]),
        search,
    )

    def count() -> int:
        total = 0
        for direction in context["directions"]:
            job = _job_for(context, direction)
            if kind == "overview" and not search:
                page = context["reader"].overview_page(job, limit=1, cursor=None)
                total += int(page.get("total_count") or 0)
                continue
            if kind == "details" and not search:
                normalized_count = _normalized_detail_row_count(context, job)
                if normalized_count is not None:
                    total += normalized_count
                    continue

            cursor = None
            while True:
                page = (
                    context["reader"].overview_page(job, limit=50, cursor=cursor)
                    if kind == "overview"
                    else context["reader"].detail_page(job, limit=50, cursor=cursor)
                )
                for item in page.get("items") or ():
                    fields = _safe_fields(item)
                    fields["direction"] = direction
                    if _search_matches(fields, search):
                        total += 1
                pagination = page.get("pagination") or {}
                if not pagination.get("has_more") or not pagination.get("next_cursor"):
                    break
                cursor = str(pagination["next_cursor"])
        return total

    return _cached_result_count(key, count)


def _combine_source_workbooks_atomically(
    staged_jobs: list[tuple[str, str, Path]],
    target: Path,
) -> None:
    """Reuse source styled-sheet copying while keeping desktop writes atomic."""
    from app.services.overview_downloader import OverviewDownloader

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    temporary.unlink(missing_ok=True)
    try:
        OverviewDownloader.combine_workbooks(staged_jobs, temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


class _ExportProgressReporter:
    """Map actual workbook work units to one monotonic request percentage."""

    _DETAIL_PHASES = {
        "load_template": (0.00, 0.05),
        "build_rows": (0.05, 0.35),
        "write_rows": (0.35, 0.75),
        "format": (0.75, 0.92),
        "save": (0.92, 1.00),
    }
    _OVERVIEW_PHASES = {
        "load_template": (0.00, 0.10),
        "write_rows": (0.10, 0.80),
        "save": (0.80, 1.00),
    }

    def __init__(self, callback: ExportProgressCallback | None) -> None:
        self.callback = callback
        self.unit_total = 0
        self.unit_index = 0
        self.last_percent = 0.0
        self.scope: str | None = None
        self.phase = "prepare"

    def planning(self, processed: int, total: int) -> None:
        fraction = self._fraction(processed, total)
        self._send("query", processed, total, fraction * 5.0)

    def set_units(self, total: int) -> None:
        self.unit_total = max(1, int(total))
        self.unit_index = 0

    def start_unit(self, scope: str) -> None:
        self.scope = scope

    def unit(self, phase: str, processed: int, total: int) -> None:
        phases = self._DETAIL_PHASES if self.scope == "details" else self._OVERVIEW_PHASES
        start, end = phases.get(phase, (0.0, 0.0))
        local = start + ((end - start) * self._fraction(processed, total))
        percent = 5.0 + (
            ((self.unit_index + local) / max(1, self.unit_total)) * 95.0
        )
        self._send(phase, processed, total, percent)

    def complete_unit(self) -> None:
        self.unit_index = min(self.unit_total, self.unit_index + 1)

    def complete(self) -> None:
        self._send("completed", 1, 1, 100.0, status="completed")

    def failed(self) -> None:
        self._send(self.phase, 0, 0, self.last_percent, status="failed")

    @staticmethod
    def _fraction(processed: int, total: int) -> float:
        if total <= 0:
            return 0.0
        return max(0.0, min(1.0, processed / total))

    def _send(
        self,
        phase: str,
        processed: int,
        total: int,
        percent: float,
        *,
        status: str = "running",
    ) -> None:
        self.phase = phase
        self.last_percent = max(self.last_percent, min(100.0, max(0.0, percent)))
        if self.callback is None:
            return
        event = {
            "status": status,
            "scope": self.scope,
            "phase": phase,
            "processed": max(0, int(processed)),
            "total": max(0, int(total)),
            "percent": round(self.last_percent, 3),
        }
        try:
            self.callback(event)
        except Exception:
            logger.exception("excel_export_progress_transport_failed phase=%s", phase)


def _write_overview_excel_from_source_template(
    rows: list[dict[str, Any]],
    *,
    direction: str,
    category: str,
    date_from: str,
    date_to: str,
    target: Path,
    progress: Callable[[str, int, int], None] | None = None,
) -> None:
    import app.services.overview_downloader as overview_module
    from app.services.overview_downloader import OverviewDownloader

    renderer = OverviewDownloader(
        crawler=None,
        headers_provider=lambda: {},
        template_dir=_source_template_dir(),
    )
    begin = date.fromisoformat(date_from)
    end = date.fromisoformat(date_to)
    safe_rows = [_excel_safe_record(row) for row in rows]
    original_load_workbook = overview_module._load_workbook
    original_clone_cell = overview_module._clone_cell
    column_count = len(_overview_template_keys(category, direction))
    timing = {"load": 0.0, "write": 0.0, "save": 0.0}
    cloned_cells = 0
    last_emitted_row = 1
    write_started_at: float | None = None

    def measured_load_workbook(*args, **kwargs):
        nonlocal write_started_at
        started = time.perf_counter()
        workbook = original_load_workbook(*args, **kwargs)
        timing["load"] += time.perf_counter() - started
        if progress is not None:
            progress("load_template", 1, 1)
            progress("write_rows", min(1, len(rows)), len(rows))
        write_started_at = time.perf_counter()
        original_save = workbook.save

        def measured_save(*save_args, **save_kwargs):
            if write_started_at is not None:
                timing["write"] += time.perf_counter() - write_started_at
            if progress is not None:
                progress("write_rows", len(rows), len(rows))
                progress("save", 0, 1)
            save_started = time.perf_counter()
            try:
                return original_save(*save_args, **save_kwargs)
            finally:
                timing["save"] += time.perf_counter() - save_started

        workbook.save = measured_save
        return workbook

    def measured_clone_cell(source, target_cell):
        nonlocal cloned_cells, last_emitted_row
        result = original_clone_cell(source, target_cell)
        cloned_cells += 1
        if progress is not None and column_count > 0:
            completed_rows = min(len(rows), 1 + (cloned_cells // column_count))
            emit_every = max(1, len(rows) // 200)
            if (
                completed_rows == len(rows)
                or completed_rows - last_emitted_row >= emit_every
            ):
                last_emitted_row = completed_rows
                progress("write_rows", completed_rows, len(rows))
        return result

    if progress is not None:
        progress("load_template", 0, 1)
    overview_module._load_workbook = measured_load_workbook
    overview_module._clone_cell = measured_clone_cell
    total_started = time.perf_counter()
    try:
        if category == "electronic":
            content = renderer._electronic_records_to_xlsx(
                safe_rows, direction, begin, end
            )
        elif category == "cash_register":
            content = renderer._cash_records_to_xlsx(
                safe_rows, direction, begin, end
            )
        else:
            raise ValueError("invalid_overview_export_category")
    finally:
        overview_module._load_workbook = original_load_workbook
        overview_module._clone_cell = original_clone_cell

    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".tmp", dir=target.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        save_started = time.perf_counter()
        temporary.write_bytes(content)
        temporary.replace(target)
        timing["save"] += time.perf_counter() - save_started
        if progress is not None:
            progress("save", 1, 1)
    finally:
        temporary.unlink(missing_ok=True)
    logger.info(
        "excel_export_summary scope=overview invoice_count=%s row_count=%s "
        "column_count=%s template_load_ms=%.1f worksheet_write_ms=%.1f "
        "workbook_save_ms=%.1f total_ms=%.1f",
        len(rows), len(rows), column_count, timing["load"] * 1000,
        timing["write"] * 1000, timing["save"] * 1000,
        (time.perf_counter() - total_started) * 1000,
    )


def _filter_detail_records(
    records: list[dict[str, Any]], search: str
) -> list[dict[str, Any]]:
    if not search:
        return records
    result: list[dict[str, Any]] = []
    for record in records:
        searchable = json.dumps(record, ensure_ascii=False, default=str).casefold()
        if search not in searchable:
            raw_path = Path(str(record.get("raw_detail_path") or ""))
            try:
                searchable = raw_path.read_text(encoding="utf-8").casefold()
            except (OSError, UnicodeError):
                searchable = ""
        if search in searchable:
            result.append(record)
    return result


def _reconciliation_export_summary(
    source_summary: dict[str, Any], rows: list[dict[str, Any]],
) -> dict[str, Any]:
    statuses = [str((item.get("fields") or {}).get("reconciliation_status") or "") for item in rows]
    difference_totals = {}
    for overview_field, _detail_field, label in RECONCILIATION_MONEY_FIELDS:
        column = f"difference_{overview_field}"
        total = sum(
            (_decimal_or_zero((item.get("fields") or {}).get(column)) for item in rows),
            Decimal("0"),
        )
        difference_totals[label] = _decimal_result(total)
    return {
        **source_summary,
        "missing_detail_count": statuses.count("Thiếu chi tiết"),
        "missing_overview_count": statuses.count("Thiếu tổng quan"),
        "money_mismatch_count": statuses.count("Chênh lệch tiền"),
        "issue_count": len(rows),
        "difference_totals": difference_totals,
    }


def _write_reconciliation_excel(
    rows: list[dict[str, Any]], target: Path,
    *, summary: dict[str, Any],
    progress: Callable[[str, int, int], None] | None = None,
) -> None:
    """Write a fixed-schema reconciliation report with an atomic replace."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if progress:
        progress("load_template", 0, 1)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Tong hop"
    worksheet = workbook.create_sheet("Bao cao doi chieu")
    headers = [
        _RECONCILIATION_LABELS[column]
        for column in _RECONCILIATION_EXPORT_COLUMNS
    ]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="EAF4EE")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="155D36")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    if progress:
        progress("load_template", 1, 1)
        progress("write_rows", 0, len(rows))
    for index, item in enumerate(rows, start=1):
        fields = item.get("fields") or {}
        worksheet.append([
            _excel_safe_value(index if column == "stt" else fields.get(column))
            for column in _RECONCILIATION_EXPORT_COLUMNS
        ])
        for column_index, column in enumerate(
            _RECONCILIATION_EXPORT_COLUMNS, start=1
        ):
            cell = worksheet.cell(index + 1, column_index)
            if column == "mismatch_fields" or (
                column.startswith("difference_") and cell.value not in (None, 0, "0")
            ):
                cell.font = Font(color="B91C1C", bold=True)
            if column == "reconciliation_reason":
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column.startswith(("overview_", "detail_", "difference_")):
                cell.number_format = '#,##0;[Red]-#,##0;0'
        if progress and (index == len(rows) or index % max(1, len(rows) // 200) == 0):
            progress("write_rows", index, len(rows))
    for column_index, column in enumerate(_RECONCILIATION_EXPORT_COLUMNS, start=1):
        width = 22
        if column == "stt":
            width = 8
        elif column in {"reconciliation_status", "mismatch_fields"}:
            width = 42
        elif column == "reconciliation_reason":
            width = 68
        elif column in {"nbten", "nmten"}:
            width = 34
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    summary_rows = [
        ("Khoảng đối chiếu", "; ".join(
            f"{value['date_from']} - {value['date_to']}"
            for value in summary.get("coverage_ranges") or ()
        )),
        ("Số hóa đơn Tổng quan", summary.get("overview_invoice_count", 0)),
        ("Số hóa đơn Chi tiết", summary.get("detail_invoice_count", 0)),
        ("Chênh lệch số lượng", summary.get("difference", 0)),
        ("Thiếu Chi tiết", summary.get("missing_detail_count", 0)),
        ("Thiếu Tổng quan", summary.get("missing_overview_count", 0)),
        ("Hóa đơn lệch tiền", summary.get("money_mismatch_count", 0)),
    ]
    summary_rows.extend(
        (f"Tổng chênh lệch - {label}", value)
        for label, value in (summary.get("difference_totals") or {}).items()
    )
    summary_sheet.append(["BÁO CÁO ĐỐI CHIẾU TỔNG QUAN & CHI TIẾT", None])
    summary_sheet.merge_cells("A1:B1")
    summary_sheet["A1"].font = Font(bold=True, color="155D36", size=14)
    summary_sheet["A1"].fill = header_fill
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    for label, value in summary_rows:
        summary_sheet.append([label, _excel_safe_value(value)])
    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 52
    summary_sheet.freeze_panes = "A2"
    if progress:
        progress(
            "format", len(_RECONCILIATION_EXPORT_COLUMNS),
            len(_RECONCILIATION_EXPORT_COLUMNS),
        )
        progress("save", 0, 1)
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    if progress:
        progress("save", 1, 1)


def _export_results_impl(
    backend,
    value: dict[str, Any],
    reporter: _ExportProgressReporter,
) -> dict[str, Any]:
    """Build Excel on demand from persisted source-owned data only."""
    from app.config.crawl_config import QUERY_TYPE_TO_CATEGORY
    from app.exporters.invoice_detail_excel_exporter import InvoiceDetailExcelExporter
    from app.repositories.invoice_detail_query_repository import InvoiceDetailQueryRepository

    destination = Path(value["destination"])
    if not destination.is_absolute():
        raise ValueError("invalid_artifact_directory")

    scopes = value.get("result_scopes") or []
    if (
        not isinstance(scopes, list)
        or not scopes
        or set(scopes) - {"overview", "details", "reconciliation"}
        or "reconciliation" in scopes and len(scopes) != 1
    ):
        raise ValueError("invalid_result_export_scope")
    connection_ids = value.get("connection_ids") or []
    if len(connection_ids) != 1:
        raise ValueError("invalid_result_export_account")

    query = {
        "connection_id": str(connection_ids[0]),
        "date_from": str(value["date_from"]),
        "date_to": str(value["date_to"]),
        "direction": value.get("direction"),
        "query_type": value.get("query_type"),
    }
    context_started = time.perf_counter()
    context = _result_context(backend, query)
    logger.info(
        "excel_export_phase scope=all phase=result_context duration_ms=%.1f",
        (time.perf_counter() - context_started) * 1000,
    )
    if context is None:
        raise ValueError("result_job_not_found")

    result_filters = value.get("result_filters") or {}
    excluded_keys = _resolve_excluded_keys(backend, value.get("exclusion"))
    destination.mkdir(parents=True, exist_ok=True)
    files: list[str] = []
    if scopes == ["reconciliation"]:
        scope_filter = (
            result_filters.get("reconciliation")
            if isinstance(result_filters, dict) else None
        )
        scope_filter = scope_filter if isinstance(scope_filter, dict) else {}
        reconciliation_query = {
            **query,
            "search": str(scope_filter.get("search", value.get("search") or "")),
            "column_filters": scope_filter.get("column_filters") or {},
            "sort": scope_filter.get("sort"),
        }
        dataset = _reconciliation_dataset(backend, reconciliation_query)
        if not dataset["summary"].get("coverage_ranges"):
            raise ValueError("result_reconciliation_coverage_missing")
        rows = _filtered_reconciliation_items(dataset, reconciliation_query)
        if not rows:
            raise ValueError("result_export_empty")
        logger.info(
            "reconciliation_export_dataset rows=%s columns=%s missing_detail=%s "
            "missing_overview=%s money_mismatch=%s coverage_ranges=%s",
            len(rows), len(_RECONCILIATION_EXPORT_COLUMNS),
            dataset["summary"]["missing_detail_count"],
            dataset["summary"]["missing_overview_count"],
            dataset["summary"]["money_mismatch_count"],
            len(dataset["summary"]["coverage_ranges"]),
        )
        reporter.planning(1, 1)
        reporter.set_units(1)
        reporter.start_unit("reconciliation")
        output_directory = _result_output_directory(
            destination, context["base_job"].company_tax_code,
            "reconciliation", context["date_from"], context["date_to"],
        )
        target = _available_path(
            output_directory,
            _grouped_result_filename(
                context["base_job"].company_tax_code,
                str(value.get("direction") or "all"), "reconciliation",
                context["date_from"], context["date_to"],
            ),
        )
        export_summary = _reconciliation_export_summary(dataset["summary"], rows)
        logger.info(
            "reconciliation_export_write output_path=%s rows=%s columns=%s",
            target, len(rows), len(_RECONCILIATION_EXPORT_COLUMNS),
        )
        _write_reconciliation_excel(
            rows, target, summary=export_summary, progress=reporter.unit,
        )
        reporter.complete_unit()
        reporter.complete()
        return {"count": 1, "files": [str(target)]}

    combinations = [
        (direction, query_type)
        for direction in context["directions"]
        for query_type in context["query_types"]
    ]
    candidate_total = len(combinations) * len(scopes)
    candidate_index = 0
    plans: list[dict[str, Any]] = []
    detail_repository = InvoiceDetailQueryRepository(context["database_path"])

    # Discover real, non-empty workbook units before rendering so mixed scope
    # progress never resets and never counts a workbook that will not exist.
    for scope in scopes:
        scope_filter = result_filters.get(scope) if isinstance(result_filters, dict) else None
        scope_filter = scope_filter if isinstance(scope_filter, dict) else {}
        scope_search = str(scope_filter.get("search", value.get("search") or "")).strip().casefold()
        scope_columns = scope_filter.get("column_filters") or {}
        for direction, query_type in combinations:
            query_started = time.perf_counter()
            if scope == "overview":
                payload = (
                    _all_overview_fields(context, direction, query_type, scope_search)
                    if not scope_columns and not excluded_keys
                    else _all_overview_fields(
                        context, direction, query_type, scope_search, scope_columns, excluded_keys
                    )
                )
            else:
                payload = detail_repository.get_detail_records_for_export(
                    context["base_job"].company_tax_code,
                    direction,
                    query_type,
                    context["date_from"],
                    context["date_to"],
                )
                if scope_search or scope_columns or excluded_keys:
                    filtered_records = []
                    row_builder = _FilteredExcelSafeDetailRowBuilder(
                        search=scope_search, column_filters=scope_columns
                    )
                    for record in payload:
                        record_key = _invoice_key(record, direction, query_type)
                        if record_key in excluded_keys:
                            continue
                        raw_path = Path(str(record.get("raw_detail_path") or ""))
                        try:
                            detail_payload = json.loads(raw_path.read_text(encoding="utf-8"))
                        except (OSError, UnicodeError, ValueError, TypeError):
                            detail_payload = {}
                        if row_builder.build_rows(detail_payload, record):
                            filtered_records.append(record)
                    payload = filtered_records
            query_seconds = time.perf_counter() - query_started
            candidate_index += 1
            reporter.planning(candidate_index, candidate_total)
            logger.info(
                "excel_export_phase scope=%s phase=query record_count=%s duration_ms=%.1f",
                scope, len(payload), query_seconds * 1000,
            )
            if payload:
                plans.append({
                    "scope": scope,
                    "direction": direction,
                    "query_type": query_type,
                    "payload": payload,
                    "search": scope_search,
                    "column_filters": scope_columns,
                })

    if not plans:
        raise ValueError("result_export_empty")
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for plan in plans:
        groups.setdefault((plan["scope"], plan["direction"]), []).append(plan)

    # One unit per rendered source sheet plus one real combine/save unit for
    # each final workbook. Query types no longer imply separate final files.
    reporter.set_units(len(plans) + len(groups))

    company_directory = destination / (
        re.sub(
            r"[^0-9A-Za-z._-]+", "_", context["base_job"].company_tax_code
        ).strip("._-") or "MIA"
    )
    company_directory.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".mia-result-sheets-", dir=company_directory) as directory:
        staging_directory = Path(directory)
        for (scope, direction), group_plans in groups.items():
            staged_jobs: list[tuple[str, str, Path]] = []
            for sheet_index, plan in enumerate(group_plans, start=1):
                query_type = plan["query_type"]
                category = str(QUERY_TYPE_TO_CATEGORY.get(query_type) or "")
                if category not in {"electronic", "cash_register"}:
                    raise ValueError("invalid_overview_export_category")
                reporter.start_unit(scope)
                staged_target = staging_directory / (
                    f"{scope}-{direction}-{category}-{sheet_index}.xlsx"
                )
                if scope == "overview":
                    _write_overview_excel_from_source_template(
                        plan["payload"],
                        direction=direction,
                        category=category,
                        date_from=context["date_from"],
                        date_to=context["date_to"],
                        target=staged_target,
                        progress=reporter.unit,
                    )
                else:
                    export_row_builder = (
                        _FilteredExcelSafeDetailRowBuilder(
                            search=plan["search"], column_filters=plan["column_filters"]
                        )
                        if plan["search"] or plan["column_filters"]
                        else _ExcelSafeDetailRowBuilder()
                    )
                    if reporter.callback is None:
                        exporter = InvoiceDetailExcelExporter(
                            _source_template_dir() / "invoice_detail.xlsx",
                            row_builder=export_row_builder,
                        )
                    else:
                        from mia_progressive_excel_exporter import (
                            ProgressiveInvoiceDetailExcelExporter,
                        )
                        exporter = ProgressiveInvoiceDetailExcelExporter(
                            _source_template_dir() / "invoice_detail.xlsx",
                            row_builder=export_row_builder,
                            progress=reporter.unit,
                        )
                    exporter.export(
                        detail_records=plan["payload"],
                        output_path=staged_target,
                        from_date=context["date_from"],
                        to_date=context["date_to"],
                    )
                combine_category = category
                if scope == "details":
                    # Source's cash_register branch intentionally reuses one
                    # Overview prototype style per column. Detail workbooks can
                    # carry row-specific styles/merges, so pass the source sheet
                    # label as the presentation category and let the same
                    # combiner take its full-cell clone path.
                    combine_category = {
                        "electronic": "Hóa đơn điện tử",
                        "cash_register": "Máy tính tiền",
                    }[category]
                staged_jobs.append((direction, combine_category, staged_target))
                reporter.complete_unit()

            reporter.start_unit(scope)
            reporter.unit("format", 0, len(staged_jobs))
            output_directory = _result_output_directory(
                destination,
                context["base_job"].company_tax_code,
                scope,
                context["date_from"],
                context["date_to"],
            )
            output_directory.mkdir(parents=True, exist_ok=True)
            target = _available_path(
                output_directory,
                _grouped_result_filename(
                    context["base_job"].company_tax_code,
                    direction,
                    scope,
                    context["date_from"],
                    context["date_to"],
                ),
            )
            _combine_source_workbooks_atomically(staged_jobs, target)
            reporter.unit("format", len(staged_jobs), len(staged_jobs))
            reporter.unit("save", 1, 1)
            reporter.complete_unit()
            files.append(str(target))

    reporter.complete()
    return {"count": len(files), "files": files}


def export_results(
    backend,
    value: dict[str, Any],
    *,
    progress_callback: ExportProgressCallback | None = None,
) -> dict[str, Any]:
    reporter = _ExportProgressReporter(progress_callback)
    started = time.perf_counter()
    try:
        return _export_results_impl(backend, value, reporter)
    except Exception:
        reporter.failed()
        raise
    finally:
        logger.info(
            "excel_export_phase scope=all phase=total duration_ms=%.1f",
            (time.perf_counter() - started) * 1000,
        )


__all__ = ["export_results", "read_results"]
