"""VAT return template exporter backed only by source-native SQLite data."""
from __future__ import annotations

import json
import errno
import os
import re
import sqlite3
import tempfile
import unicodedata
import zipfile
from copy import copy, deepcopy
from contextlib import closing
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from mia_artifact_pipeline import ArtifactInspector
from mia_export_paths import direction_export_directory
from app.utils.invoice_identity import (canonical_invoice_identity,
                                        invoice_status_is_excluded,
                                        masked_invoice_identity)

TEMPLATE_NAME = "To khai thue GTGT khau tru 01GTGT.xlsx"
SHEET_NAME = "TỜ KHAI THUẾ GTGT"
SOLD_SHEET_NAME = "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HHDV BÁN RA"
PURCHASE_SHEET_NAME = "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ CỦA HHDV MUA VÀO"
REDUCTION_SHEET_NAME = "GIẢM THUẾ GIÁ TRỊ GIA TĂNG"
QUERY_TYPES = ("query", "sco-query")
MONEY_NUMBER_FORMAT = '#,##0'
ONE_DONG = Decimal("1")
MONEY_CELLS = frozenset({
    "H11", "F14", "H14", "F16", "H16", "H17", "F19", "F20", "H20",
    "F21", "F22", "H22", "F23", "H23", "F24", "F25", "H25", "H26",
    "H28", "H29", "H30", "H32", "H33", "H34", "H35", "H36", "H38",
})
OBLIGATION_FORMULAS = {
    "H32": "=MAX(H26-H11+H28-H29-H30,0)",
    "H35": "=MAX(-(H26-H11+H28-H29-H30),0)",
    "H38": "=H35-H36",
}


@dataclass(frozen=True)
class CachedFormula:
    formula: str
    value: Decimal


def _template_path() -> Path:
    return Path(__file__).resolve().parent / "vendor" / "mia_crawl_service" / "resources" / "templates" / TEMPLATE_NAME


def _decimal(value: Any, field: str) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    if isinstance(value, Decimal):
        number = value
        return Decimal(0) if number == 0 else number.quantize(ONE_DONG, rounding=ROUND_HALF_UP)
    if isinstance(value, float):
        raise ValueError(f"vat_return_float_money:{field}")
    if isinstance(value, int):
        return Decimal(value)
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"[\s\u00a0]", "", text)
    if not text:
        return Decimal(0)
    if "," in text and "." in text:
        decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
        grouping_separator = "." if decimal_separator == "," else ","
        text = text.replace(grouping_separator, "").replace(decimal_separator, ".")
    elif "," in text:
        groups = text.split(",")
        text = "".join(groups) if len(groups) > 2 or (len(groups) == 2 and len(groups[1]) == 3) else text.replace(",", ".")
    try:
        number = Decimal(text).quantize(ONE_DONG, rounding=ROUND_HALF_UP)
        return Decimal(0) if number == 0 else number
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"vat_return_invalid_decimal:{field}") from error


def _decimal_exact(value: Any, field: str) -> Decimal:
    """Parse source money without binary floats or premature rounding."""
    if value in (None, ""):
        return Decimal(0)
    if isinstance(value, float):
        raise ValueError(f"vat_return_float_money:{field}")
    if isinstance(value, Decimal):
        return Decimal(0) if value == 0 else value
    if isinstance(value, int):
        return Decimal(value)
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"[\s\u00a0]", "", text)
    if "," in text and "." in text:
        decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
        grouping_separator = "." if decimal_separator == "," else ","
        text = text.replace(grouping_separator, "").replace(decimal_separator, ".")
    elif "," in text:
        groups = text.split(",")
        text = "".join(groups) if len(groups) > 2 or (len(groups) == 2 and len(groups[1]) == 3) else text.replace(",", ".")
    try:
        number = Decimal(text)
        return Decimal(0) if number == 0 else number
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"vat_return_invalid_decimal:{field}") from error


_MAIN_NS = "http:" "//schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http:" "//schemas.openxmlformats.org/officeDocument/2006/relationships"
ET.register_namespace("", _MAIN_NS)
ET.register_namespace("r", _REL_NS)


def _set_ooxml_cell(root: ET.Element, address: str, value: Any) -> None:
    cell = root.find(f".//{{{_MAIN_NS}}}c[@r='{address}']")
    if cell is None:
        raise RuntimeError(f"vat_return_template_cell_missing:{address}")
    for child in list(cell):
        if child.tag in {f"{{{_MAIN_NS}}}v", f"{{{_MAIN_NS}}}f", f"{{{_MAIN_NS}}}is"}:
            cell.remove(child)
    if value is None:
        cell.attrib.pop("t", None)
    elif isinstance(value, CachedFormula):
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{_MAIN_NS}}}f").text = value.formula.removeprefix("=")
        ET.SubElement(cell, f"{{{_MAIN_NS}}}v").text = format(
            Decimal(0) if value.value == 0 else value.value, "f"
        )
    elif isinstance(value, str) and value.startswith("="):
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{_MAIN_NS}}}f").text = value[1:]
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        inline = ET.SubElement(cell, f"{{{_MAIN_NS}}}is")
        ET.SubElement(inline, f"{{{_MAIN_NS}}}t").text = value
    else:
        cell.attrib.pop("t", None)
        if isinstance(value, Decimal):
            text = format(Decimal(0) if value == 0 else value, "f")
        else:
            text = str(value)
        ET.SubElement(cell, f"{{{_MAIN_NS}}}v").text = text


def _apply_number_format_styles(
    styles_root: ET.Element, sheet_root: ET.Element,
    addresses: set[str] | frozenset[str], number_format: str,
) -> None:
    num_fmts = styles_root.find(f"{{{_MAIN_NS}}}numFmts")
    if num_fmts is None:
        num_fmts = ET.Element(f"{{{_MAIN_NS}}}numFmts", {"count": "0"})
        styles_root.insert(0, num_fmts)
    existing = next((item for item in num_fmts if item.get("formatCode") == number_format), None)
    if existing is None:
        used_ids = {int(item.get("numFmtId", "0")) for item in num_fmts}
        num_fmt_id = max({163, *used_ids}) + 1
        ET.SubElement(num_fmts, f"{{{_MAIN_NS}}}numFmt", {
            "numFmtId": str(num_fmt_id), "formatCode": number_format,
        })
        num_fmts.set("count", str(len(num_fmts)))
    else:
        num_fmt_id = int(existing.get("numFmtId", "0"))

    cell_xfs = styles_root.find(f"{{{_MAIN_NS}}}cellXfs")
    if cell_xfs is None:
        raise RuntimeError("vat_return_template_styles_missing")
    style_map: dict[int, int] = {}
    for address in addresses:
        cell = sheet_root.find(f".//{{{_MAIN_NS}}}c[@r='{address}']")
        if cell is None:
            raise RuntimeError(f"vat_return_template_cell_missing:{address}")
        original_style = int(cell.get("s", "0"))
        replacement_style = style_map.get(original_style)
        if replacement_style is None:
            cloned = deepcopy(list(cell_xfs)[original_style])
            cloned.set("numFmtId", str(num_fmt_id))
            cloned.set("applyNumberFormat", "1")
            cell_xfs.append(cloned)
            replacement_style = len(cell_xfs) - 1
            style_map[original_style] = replacement_style
        cell.set("s", str(replacement_style))
    cell_xfs.set("count", str(len(cell_xfs)))


def _apply_money_styles(
    styles_root: ET.Element, sheet_root: ET.Element,
    addresses: set[str] | frozenset[str] = MONEY_CELLS,
) -> None:
    _apply_number_format_styles(styles_root, sheet_root, addresses, MONEY_NUMBER_FORMAT)


def _row_with_number(source: ET.Element, row_number: int, *, clear: bool = False) -> ET.Element:
    row = deepcopy(source)
    row.set("r", str(row_number))
    for cell in row.findall(f"{{{_MAIN_NS}}}c"):
        column = re.match(r"[A-Z]+", str(cell.get("r"))).group(0)
        cell.set("r", f"{column}{row_number}")
        if clear:
            for child in list(cell):
                if child.tag in {f"{{{_MAIN_NS}}}v", f"{{{_MAIN_NS}}}f", f"{{{_MAIN_NS}}}is"}:
                    cell.remove(child)
            cell.attrib.pop("t", None)
    return row


def _build_sold_sheet(sheet_root: ET.Element, company_name: str, tax_code: str,
                      groups: dict[str, list[dict[str, Any]]]) -> set[str]:
    source_rows = {int(row.get("r")): deepcopy(row) for row in sheet_root.findall(
        f"{{{_MAIN_NS}}}sheetData/{{{_MAIN_NS}}}row"
    )}
    sheet_data = sheet_root.find(f"{{{_MAIN_NS}}}sheetData")
    for row in list(sheet_data):
        if int(row.get("r")) > 6:
            sheet_data.remove(row)
    _set_ooxml_cell(sheet_root, "A1", company_name)
    _set_ooxml_cell(sheet_root, "A2", f"Mã số thuế: {tax_code}")
    merge_cells = sheet_root.find(f"{{{_MAIN_NS}}}mergeCells")
    for merge in list(merge_cells):
        numbers = [int(value) for value in re.findall(r"\d+", str(merge.get("ref")))]
        if numbers and max(numbers) > 6:
            merge_cells.remove(merge)
    money_addresses: set[str] = set()
    row_number = 7
    sequence = 1
    source_map = {"kct": (7, 9), "0": (10, 12), "5": (13, 15), "10": (16, 27)}
    for group_name in ("kct", "0", "5", "10"):
        header_source, total_source = source_map[group_name]
        header = _row_with_number(source_rows[header_source], row_number)
        sheet_data.append(header)
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:P{row_number}"})
        row_number += 1
        group_items = groups[group_name]
        group_base = group_tax = group_reduction = Decimal(0)
        for index, item in enumerate(group_items):
            template_row = source_rows[17 if index == 0 else 18]
            output_row = _row_with_number(template_row, row_number, clear=True)
            sheet_data.append(output_row)
            ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"G{row_number}:I{row_number}"})
            excel_date = (date.fromisoformat(item["nlap_date"]) - date(1899, 12, 30)).days
            values = {
                "A": sequence, "B": item["khmshdon"], "D": item["khhdon"],
                "E": item["shdon"], "F": excel_date, "G": item["buyer_name"],
                "J": item["buyer_tax_code"], "K": item["base"], "L": item["tax"],
                "M": "X" if item["has_reduction"] else None,
                "N": item["reduction"] if item["has_reduction"] else None,
                "O": item["note"] or None, "P": item["description"] or None,
            }
            for column, value in values.items():
                _set_ooxml_cell(output_row, f"{column}{row_number}", value)
            money_addresses.update({f"K{row_number}", f"L{row_number}", f"N{row_number}"})
            group_base += item["base"]; group_tax += item["tax"]
            group_reduction += item["reduction"] if item["has_reduction"] else Decimal(0)
            sequence += 1; row_number += 1
        total = _row_with_number(source_rows[total_source], row_number, clear=True)
        sheet_data.append(total)
        _set_ooxml_cell(total, f"A{row_number}", "Tổng")
        _set_ooxml_cell(total, f"K{row_number}", group_base)
        _set_ooxml_cell(total, f"L{row_number}", group_tax)
        _set_ooxml_cell(total, f"N{row_number}", group_reduction)
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:J{row_number}"})
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"O{row_number}:P{row_number}"})
        money_addresses.update({f"K{row_number}", f"L{row_number}", f"N{row_number}"})
        row_number += 1
    merge_cells.set("count", str(len(merge_cells)))
    sheet_root.find(f"{{{_MAIN_NS}}}dimension").set("ref", f"A1:P{row_number - 1}")
    return money_addresses


def _build_purchase_sheet(
    sheet_root: ET.Element, company_name: str, tax_code: str,
    items: list[dict[str, Any]],
) -> tuple[set[str], set[str]]:
    """Replace only the sample body while preserving the template sheet layout."""
    source_rows = {int(row.get("r")): deepcopy(row) for row in sheet_root.findall(
        f"{{{_MAIN_NS}}}sheetData/{{{_MAIN_NS}}}row"
    )}
    sheet_data = sheet_root.find(f"{{{_MAIN_NS}}}sheetData")
    if sheet_data is None:
        raise RuntimeError("vat_return_purchase_sheet_data_missing")
    for row in list(sheet_data):
        if int(row.get("r")) > 6:
            sheet_data.remove(row)
    _set_ooxml_cell(sheet_root, "A1", company_name)
    _set_ooxml_cell(sheet_root, "A2", f"Mã số thuế: {tax_code}")
    merge_cells = sheet_root.find(f"{{{_MAIN_NS}}}mergeCells")
    if merge_cells is None:
        raise RuntimeError("vat_return_purchase_merge_cells_missing")
    for merge in list(merge_cells):
        numbers = [int(value) for value in re.findall(r"\d+", str(merge.get("ref")))]
        if numbers and max(numbers) > 6:
            merge_cells.remove(merge)

    money_addresses: set[str] = set()
    date_addresses: set[str] = set()
    row_number = 7

    # The source data has no confirmed field that can safely route invoices to
    # template groups 2/3. Keep all eligible invoices in the sample's group 1.
    group_one = _row_with_number(source_rows[7], row_number)
    sheet_data.append(group_one)
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:N{row_number}"})
    row_number += 1
    total_base = total_tax = Decimal(0)
    for index, item in enumerate(items):
        if index == 0:
            template_row = source_rows[8]
        elif index == len(items) - 1:
            template_row = source_rows[30]
        else:
            template_row = source_rows[9]
        output_row = _row_with_number(template_row, row_number, clear=True)
        sheet_data.append(output_row)
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"F{row_number}:G{row_number}"})
        excel_date = (date.fromisoformat(item["nlap_date"]) - date(1899, 12, 30)).days
        values = {
            "A": index + 1,
            "B": item["khmshdon"],
            "C": item["khhdon"],
            "D": item["shdon"],
            "E": excel_date,
            "F": item["seller_name"],
            "H": item["seller_tax_code"],
            "I": item["base"],
            "J": item["tax"],
            "K": item["deductible_tax"],
            "L": item["note"] or None,
            "M": item["description"],
            "N": None,
        }
        for column, value in values.items():
            _set_ooxml_cell(output_row, f"{column}{row_number}", value)
        money_addresses.update({f"I{row_number}", f"J{row_number}", f"K{row_number}"})
        date_addresses.add(f"E{row_number}")
        total_base += item["base"]
        total_tax += item["tax"]
        row_number += 1

    total_one = _row_with_number(source_rows[31], row_number, clear=True)
    sheet_data.append(total_one)
    _set_ooxml_cell(total_one, f"A{row_number}", "Tổng")
    for column, value in (("I", total_base), ("J", total_tax), ("K", total_tax)):
        _set_ooxml_cell(total_one, f"{column}{row_number}", value)
        money_addresses.add(f"{column}{row_number}")
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:H{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"L{row_number}:N{row_number}"})
    row_number += 1

    for header_source, total_source in ((32, 34), (35, 37)):
        header = _row_with_number(source_rows[header_source], row_number)
        sheet_data.append(header)
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:N{row_number}"})
        row_number += 1
        total = _row_with_number(source_rows[total_source], row_number, clear=True)
        sheet_data.append(total)
        _set_ooxml_cell(total, f"A{row_number}", "Tổng")
        for column in ("I", "J", "K"):
            _set_ooxml_cell(total, f"{column}{row_number}", Decimal(0))
            money_addresses.add(f"{column}{row_number}")
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:H{row_number}"})
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"L{row_number}:N{row_number}"})
        row_number += 1

    merge_cells.set("count", str(len(merge_cells)))
    sheet_root.find(f"{{{_MAIN_NS}}}dimension").set("ref", f"A1:O{row_number - 1}")
    return money_addresses, date_addresses


def _build_reduction_sheet(
    sheet_root: ET.Element, company_name: str, tax_code: str,
    purchase_lines: list[dict[str, Any]], sold_lines: list[dict[str, Any]],
) -> set[str]:
    """Populate the statutory 8% purchase and sold reduction schedules."""
    source_rows = {int(row.get("r")): deepcopy(row) for row in sheet_root.findall(
        f"{{{_MAIN_NS}}}sheetData/{{{_MAIN_NS}}}row"
    )}
    sheet_data = sheet_root.find(f"{{{_MAIN_NS}}}sheetData")
    if sheet_data is None:
        raise RuntimeError("vat_return_reduction_sheet_data_missing")
    for row in list(sheet_data):
        if int(row.get("r")) > 7:
            sheet_data.remove(row)
    _set_ooxml_cell(sheet_root, "A1", company_name)
    _set_ooxml_cell(sheet_root, "A2", f"Mã số thuế: {tax_code}")
    merge_cells = sheet_root.find(f"{{{_MAIN_NS}}}mergeCells")
    if merge_cells is None:
        raise RuntimeError("vat_return_reduction_merge_cells_missing")
    for merge in list(merge_cells):
        numbers = [int(value) for value in re.findall(r"\d+", str(merge.get("ref")))]
        if numbers and max(numbers) > 7:
            merge_cells.remove(merge)

    money_addresses: set[str] = set()
    row_number = 8
    total_base = total_tax = Decimal(0)
    for index, item in enumerate(purchase_lines):
        output_row = _row_with_number(source_rows[8], row_number, clear=True)
        sheet_data.append(output_row)
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"B{row_number}:E{row_number}"})
        ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"G{row_number}:J{row_number}"})
        for column, value in (
            ("A", index + 1), ("B", item["name"]),
            ("F", item["base"]), ("G", item["tax"]),
        ):
            _set_ooxml_cell(output_row, f"{column}{row_number}", value)
        money_addresses.update({f"F{row_number}", f"G{row_number}"})
        total_base += item["base"]
        total_tax += item["tax"]
        row_number += 1

    purchase_total = _row_with_number(source_rows[23], row_number, clear=True)
    sheet_data.append(purchase_total)
    _set_ooxml_cell(purchase_total, f"A{row_number}", " Tổng")
    _set_ooxml_cell(purchase_total, f"F{row_number}", total_base)
    _set_ooxml_cell(purchase_total, f"G{row_number}", total_tax)
    money_addresses.update({f"F{row_number}", f"G{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:E{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"G{row_number}:J{row_number}"})
    row_number += 1

    for source_row in (24, 25, 26):
        sheet_data.append(_row_with_number(source_rows[source_row], row_number))
        if source_row == 24:
            ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:F{row_number}"})
        else:
            for columns in ("B:E", "G:H", "I:J", "K:L"):
                ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"{columns.split(':')[0]}{row_number}:{columns.split(':')[1]}{row_number}"})
        row_number += 1

    sold_total_base = sold_total_reduction = Decimal(0)
    for index, item in enumerate(sold_lines):
        output_row = _row_with_number(source_rows[27], row_number, clear=True)
        sheet_data.append(output_row)
        for columns in ("B:E", "G:H", "I:J", "K:L"):
            start, end = columns.split(":")
            ET.SubElement(
                merge_cells, f"{{{_MAIN_NS}}}mergeCell",
                {"ref": f"{start}{row_number}:{end}{row_number}"},
            )
        reduced_rate = item["reduced_rate"]
        reduction = item["reduction"]
        for column, value in (
            ("A", index + 1),
            ("B", item["name"]),
            ("F", item["base"]),
            ("G", item["statutory_rate"]),
            ("I", CachedFormula(f"=G{row_number}*80%", reduced_rate)),
            ("K", CachedFormula(f"=F{row_number}*(G{row_number}-I{row_number})", reduction)),
        ):
            _set_ooxml_cell(output_row, f"{column}{row_number}", value)
        money_addresses.update({f"F{row_number}", f"K{row_number}"})
        sold_total_base += item["base"]
        sold_total_reduction += reduction
        row_number += 1

    sold_total = _row_with_number(source_rows[48], row_number, clear=True)
    sheet_data.append(sold_total)
    _set_ooxml_cell(sold_total, f"A{row_number}", " Tổng")
    _set_ooxml_cell(sold_total, f"F{row_number}", sold_total_base)
    _set_ooxml_cell(sold_total, f"K{row_number}", sold_total_reduction)
    money_addresses.update({f"F{row_number}", f"K{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:E{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"G{row_number}:H{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"I{row_number}:J{row_number}"})
    sold_total_row = row_number
    row_number += 1

    part_three = _row_with_number(source_rows[49], row_number)
    sheet_data.append(part_three)
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:F{row_number}"})
    row_number += 1
    difference = _row_with_number(source_rows[50], row_number)
    reduction_difference = sold_total_reduction - total_tax
    _set_ooxml_cell(
        difference, f"C{row_number}",
        CachedFormula(f"=K{sold_total_row}-G{purchase_total.get('r')}", reduction_difference),
    )
    money_addresses.add(f"C{row_number}")
    sheet_data.append(difference)
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"A{row_number}:B{row_number}"})
    ET.SubElement(merge_cells, f"{{{_MAIN_NS}}}mergeCell", {"ref": f"C{row_number}:C{row_number + 1}"})
    row_number += 1
    sheet_data.append(_row_with_number(source_rows[51], row_number, clear=True))

    merge_cells.set("count", str(len(merge_cells)))
    sheet_root.find(f"{{{_MAIN_NS}}}dimension").set("ref", f"A1:L{row_number}")
    return money_addresses


def _write_template_workbook(template: Path, target: Path, values: dict[str, Any],
                             *, company_name: str, tax_code: str,
                             sold_groups: dict[str, list[dict[str, Any]]],
                             purchase_items: list[dict[str, Any]],
                             purchase_reduction_lines: list[dict[str, Any]],
                             sold_reduction_lines: list[dict[str, Any]]) -> None:
    with zipfile.ZipFile(template, "r") as source:
        workbook_xml = ET.fromstring(source.read("xl/workbook.xml"))
        first_sheet = workbook_xml.find(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
        if first_sheet is None or first_sheet.get("name") != SHEET_NAME:
            raise RuntimeError("vat_return_template_sheet_missing")
        relation_id = first_sheet.get(f"{{{_REL_NS}}}id")
        relationships = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        relationship = next((item for item in relationships if item.get("Id") == relation_id), None)
        if relationship is None:
            raise RuntimeError("vat_return_template_sheet_missing")
        relation_target = str(relationship.get("Target")).lstrip("/")
        sheet_path = relation_target if relation_target.startswith("xl/") else "xl/" + relation_target
        sheet_xml = ET.fromstring(source.read(sheet_path))
        sold_sheet = next((item for item in workbook_xml.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
                           if item.get("name") == SOLD_SHEET_NAME), None)
        if sold_sheet is None:
            raise RuntimeError("vat_return_sold_sheet_missing")
        sold_relation_id = sold_sheet.get(f"{{{_REL_NS}}}id")
        sold_relationship = next(item for item in relationships if item.get("Id") == sold_relation_id)
        sold_target = str(sold_relationship.get("Target")).lstrip("/")
        sold_path = sold_target if sold_target.startswith("xl/") else "xl/" + sold_target
        sold_xml = ET.fromstring(source.read(sold_path))
        purchase_sheet = next((item for item in workbook_xml.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
                               if item.get("name") == PURCHASE_SHEET_NAME), None)
        if purchase_sheet is None:
            raise RuntimeError("vat_return_purchase_sheet_missing")
        purchase_relation_id = purchase_sheet.get(f"{{{_REL_NS}}}id")
        purchase_relationship = next(item for item in relationships if item.get("Id") == purchase_relation_id)
        purchase_target = str(purchase_relationship.get("Target")).lstrip("/")
        purchase_path = purchase_target if purchase_target.startswith("xl/") else "xl/" + purchase_target
        purchase_xml = ET.fromstring(source.read(purchase_path))
        reduction_sheet = next((item for item in workbook_xml.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
                                if item.get("name") == REDUCTION_SHEET_NAME), None)
        if reduction_sheet is None:
            raise RuntimeError("vat_return_reduction_sheet_missing")
        reduction_relation_id = reduction_sheet.get(f"{{{_REL_NS}}}id")
        reduction_relationship = next(item for item in relationships if item.get("Id") == reduction_relation_id)
        reduction_target = str(reduction_relationship.get("Target")).lstrip("/")
        reduction_path = reduction_target if reduction_target.startswith("xl/") else "xl/" + reduction_target
        reduction_xml = ET.fromstring(source.read(reduction_path))
        styles_xml = ET.fromstring(source.read("xl/styles.xml"))
        for address, value in values.items():
            _set_ooxml_cell(sheet_xml, address, value)
        _apply_money_styles(styles_xml, sheet_xml)
        sold_money_cells = _build_sold_sheet(sold_xml, company_name, tax_code, sold_groups)
        _apply_money_styles(styles_xml, sold_xml, sold_money_cells)
        purchase_money_cells, purchase_date_cells = _build_purchase_sheet(
            purchase_xml, company_name, tax_code, purchase_items,
        )
        _apply_money_styles(styles_xml, purchase_xml, purchase_money_cells)
        if purchase_date_cells:
            _apply_number_format_styles(styles_xml, purchase_xml, purchase_date_cells, "dd/mm/yyyy")
        reduction_money_cells = _build_reduction_sheet(
            reduction_xml, company_name, tax_code,
            purchase_reduction_lines, sold_reduction_lines,
        )
        _apply_money_styles(styles_xml, reduction_xml, reduction_money_cells)
        defined_names = workbook_xml.find(f"{{{_MAIN_NS}}}definedNames")
        if defined_names is not None:
            last_row = sold_xml.find(f"{{{_MAIN_NS}}}dimension").get("ref").split(":")[-1][1:]
            for item in defined_names:
                if item.get("name") == "_xlnm.Print_Area" and SOLD_SHEET_NAME in str(item.text):
                    item.text = f"'{SOLD_SHEET_NAME}'!$A$1:$P${last_row}"
                if item.get("name") == "_xlnm.Print_Area" and PURCHASE_SHEET_NAME in str(item.text):
                    purchase_last = purchase_xml.find(f"{{{_MAIN_NS}}}dimension").get("ref").split(":")[-1][1:]
                    item.text = f"'{PURCHASE_SHEET_NAME}'!$A$1:$O${purchase_last}"
                if item.get("name") == "_xlnm.Print_Area" and REDUCTION_SHEET_NAME in str(item.text):
                    reduction_last = reduction_xml.find(f"{{{_MAIN_NS}}}dimension").get("ref").split(":")[-1][1:]
                    item.text = f"'{REDUCTION_SHEET_NAME}'!$A$1:$L${reduction_last}"
        calc = workbook_xml.find(f"{{{_MAIN_NS}}}calcPr")
        if calc is None:
            calc = ET.SubElement(workbook_xml, f"{{{_MAIN_NS}}}calcPr")
        calc.set("calcMode", "auto"); calc.set("fullCalcOnLoad", "1"); calc.set("forceFullCalc", "1")
        replacements = {
            sheet_path: ET.tostring(sheet_xml, encoding="utf-8", xml_declaration=True),
            sold_path: ET.tostring(sold_xml, encoding="utf-8", xml_declaration=True),
            purchase_path: ET.tostring(purchase_xml, encoding="utf-8", xml_declaration=True),
            reduction_path: ET.tostring(reduction_xml, encoding="utf-8", xml_declaration=True),
            "xl/workbook.xml": ET.tostring(workbook_xml, encoding="utf-8", xml_declaration=True),
            "xl/styles.xml": ET.tostring(styles_xml, encoding="utf-8", xml_declaration=True),
        }
        with zipfile.ZipFile(target, "w") as output:
            for info in source.infolist():
                if info.filename == "xl/calcChain.xml":
                    continue
                output.writestr(copy(info), replacements.get(info.filename, source.read(info.filename)))


def _normalized_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().strip()
    return re.sub(r"[^\w%]+", " ", text, flags=re.UNICODE).strip()


def _excluded_status(value: Any) -> bool:
    return invoice_status_is_excluded(value)


def _tax_group(value: Any) -> str | None:
    text = _normalized_text(value)
    semantic = text.replace(" ", "")
    if semantic in {"kct", "khôngchịuthuế", "khongchiuthue"}: return "kct"
    if semantic in {"kkknt", "khôngkêkhaitínhnộpthuế", "khongkekhaitinhnopthue", "khôngkêkhaikhôngtínhthuế", "khongkekhaikhongtinhthue", "khôngtínhthuế", "khongtinhthue"}: return "kkknt"
    # Keep decimal punctuation until after numeric parsing. Using
    # _normalized_text here would turn 8.0 into 80 and 8.00% into 800%.
    raw = unicodedata.normalize("NFKC", str(value or "")).strip()
    raw = re.sub(r"[\s\u00a0]", "", raw).removesuffix("%").replace(",", ".")
    try:
        number = Decimal(raw.replace(",", "."))
    except InvalidOperation:
        return None
    percent = number * 100 if abs(number) <= 1 and number != 0 else number
    if percent == 0: return "0"
    if percent == 5: return "5"
    if percent == 8: return "8"
    if percent == 10: return "10"
    return None


def _attributes(connection: sqlite3.Connection, invoice_id: int) -> dict[str, Any]:
    output = {}
    for name, raw in connection.execute(
        "SELECT field_name,value_json FROM invoice_overview_attributes WHERE invoice_item_id=?", (invoice_id,)
    ):
        try: output[str(name)] = json.loads(raw, parse_float=Decimal)
        except (TypeError, json.JSONDecodeError): output[str(name)] = raw
    return output


def _verified_purchase_totals_from_detail(
    connection: sqlite3.Connection, tax_code: str,
    overview: sqlite3.Row, fields: dict[str, Any],
) -> tuple[Decimal, Decimal, int] | None:
    """Recover absent source totals only from a complete, internally balanced detail.

    Some cash-register overview responses persist explicit JSON null for both
    invoice totals. Re-downloading Overview cannot repair those source values.
    Detail is a valid fallback only when every normalized line has both money
    values and their exact aggregate balances to the source invoice payment.
    """
    payment_raw = fields.get("tgtttbso")
    if payment_raw in (None, ""):
        return None
    if _decimal(fields.get("ttcktmai"), "ttcktmai") != 0:
        return None
    if _decimal(fields.get("tgtphi"), "tgtphi") != 0:
        return None
    candidates = connection.execute(
        """SELECT * FROM invoice_detail_items
           WHERE company_tax_code=? AND direction='purchase' AND query_type=?
             AND COALESCE(nbmst,'')=COALESCE(?, '')
             AND COALESCE(khhdon,'')=COALESCE(?, '')
             AND COALESCE(shdon,'')=COALESCE(?, '')
             AND COALESCE(khmshdon,'')=COALESCE(?, '')
             AND normalized_ready=1
             AND detail_outcome='with_lines'
             AND (error_message IS NULL OR TRIM(error_message)='')
           ORDER BY id DESC""",
        (
            tax_code, overview["query_type"], overview["nbmst"],
            overview["khhdon"], overview["shdon"], overview["khmshdon"],
        ),
    ).fetchall()
    payment = _decimal(payment_raw, "tgtttbso")
    for detail in candidates:
        lines = connection.execute(
            """SELECT CAST(thtien AS TEXT) AS thtien,
                      CAST(tthue AS TEXT) AS tthue
               FROM invoice_detail_lines
               WHERE detail_item_id=? ORDER BY line_number,id""",
            (detail["id"],),
        ).fetchall()
        if not lines or int(detail["normalized_line_count"] or 0) != len(lines):
            continue
        if any(line["thtien"] in (None, "") or line["tthue"] in (None, "") for line in lines):
            continue
        try:
            base = _decimal(sum(
                (_decimal_exact(line["thtien"], "detail_thtien") for line in lines),
                Decimal(0),
            ), "detail_base_total")
            tax = _decimal(sum(
                (_decimal_exact(line["tthue"], "detail_tthue") for line in lines),
                Decimal(0),
            ), "detail_tax_total")
        except ValueError:
            continue
        if abs((base + tax) - payment) <= Decimal("5"):
            return base, tax, int(detail["id"])
    return None


def audit_vat_detail_completeness(
    database: Path, tax_code: str, date_from: str, date_to: str,
) -> dict[str, Any]:
    """Invoice-level reconciliation used by both validation and diagnostics."""
    with closing(sqlite3.connect(database)) as connection:
        connection.row_factory = sqlite3.Row
        overview = connection.execute(
            """SELECT * FROM invoice_overview_items WHERE company_tax_code=?
               AND nlap_date BETWEEN ? AND ? AND direction IN ('purchase','sold')
               AND query_type IN ('query','sco-query')""",
            (tax_code, date_from, date_to),
        ).fetchall()
        details = connection.execute(
            """SELECT * FROM invoice_detail_items WHERE company_tax_code=?
               AND direction IN ('purchase','sold') AND nlap_date BETWEEN ? AND ?""",
            (tax_code, date_from, date_to),
        ).fetchall()
        detail_map: dict[tuple[str, ...], list[sqlite3.Row]] = {}
        for item in details:
            detail_map.setdefault(canonical_invoice_identity(tax_code, item['direction'], dict(item)), []).append(item)
        canonical: dict[tuple[str, ...], list[sqlite3.Row]] = {}
        for item in overview:
            canonical.setdefault(canonical_invoice_identity(tax_code, item['direction'], dict(item)), []).append(item)
        missing = []
        by_route: dict[str, dict[str, dict[str, int]]] = {
            direction: {query: {'overview': 0, 'with_detail': 0, 'missing': 0}
                        for query in QUERY_TYPES}
            for direction in ('purchase', 'sold')
        }
        for identity, copies in canonical.items():
            representative = copies[0]
            route = by_route[representative['direction']][representative['query_type']]
            route['overview'] += 1
            matches = detail_map.get(identity, [])
            complete = any(
                int(item['normalized_ready'] or 0) == 1
                and item['detail_outcome'] in ('with_lines', 'valid_empty')
                and not str(item['error_message'] or '').strip()
                for item in matches
            )
            if complete:
                route['with_detail'] += 1
            else:
                route['missing'] += 1
                missing.append({
                    'direction': representative['direction'],
                    'query_type': representative['query_type'],
                    'date': representative['nlap_date'],
                    'khhdon': str(representative['khhdon'] or ''),
                    'shdon': str(representative['shdon'] or ''),
                    'partner_masked': ('*' * max(0, len(str(representative['nbmst'] or '')) - 4)
                                       + str(representative['nbmst'] or '')[-4:]),
                    'identity': masked_invoice_identity(identity),
                    'reason': 'no_detail_record' if not matches else 'detail_not_finalized',
                })
        overview_ids = set(canonical)
        orphan_headers = sum(1 for identity in detail_map if identity not in overview_ids)
        orphan_lines = int(connection.execute(
            """SELECT COUNT(*) FROM invoice_detail_lines line LEFT JOIN invoice_detail_items item
               ON item.id=line.detail_item_id WHERE item.id IS NULL"""
        ).fetchone()[0])
        duplicates = sum(1 for copies in canonical.values()
                         if len({item['query_type'] for item in copies}) > 1)
        return {'by_route': by_route, 'missing': missing,
                'canonical_overview': len(canonical), 'orphan_detail_headers': orphan_headers,
                'orphan_detail_lines': orphan_lines, 'cross_query_duplicates': duplicates}


def _field(fields: dict[str, Any], row: sqlite3.Row, *names: str) -> Any:
    for name in names:
        value = fields.get(name, row[name] if name in row.keys() else None)
        if value not in (None, ""):
            return value
    return ""


def _actual_rate(base: Decimal, tax: Decimal) -> tuple[Decimal | None, str]:
    if base == 0:
        return None, "undefined"
    raw = tax / base
    for rate, label in ((Decimal("0"), "0"), (Decimal("0.05"), "5"),
                        (Decimal("0.08"), "8"), (Decimal("0.10"), "10")):
        if abs(tax - base * rate) <= Decimal("5"):
            return rate, label
    return raw, format(raw * 100, "f").rstrip("0").rstrip(".")


def _tax_matches_rate(base: Decimal, tax: Decimal, rate: Decimal) -> bool:
    """Match portal tax to a statutory rate using the agreed five-dong tolerance."""
    expected = (base * rate).quantize(ONE_DONG, rounding=ROUND_HALF_UP)
    return abs(tax - expected) <= Decimal("5")


def _masked_tax_code(value: Any) -> str:
    text = str(value or "")
    return "*" * max(0, len(text) - 4) + text[-4:]


def _invoice_number_key(value: str) -> tuple[int, str]:
    compact = re.sub(r"\D+", "", value)
    return (int(compact) if compact else 10**30, value.casefold())


def _group_sold_reduction_lines(
    lines: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Group rows that have the same displayed Part-II classification.

    The name and the two displayed rates are dimensions. Revenue is additive;
    the reduction is recalculated from the grouped revenue so the cached value
    and the Excel formula use exactly the same header-defined calculation.
    """
    grouped: dict[tuple[str, Decimal, Decimal], dict[str, Any]] = {}
    for item in lines:
        display_name = unicodedata.normalize(
            "NFKC", str(item.get("name") or "")
        ).strip()
        name_key = re.sub(r"\s+", " ", display_name).casefold()
        key = (name_key, item["statutory_rate"], item["reduced_rate"])
        target = grouped.get(key)
        if target is None:
            target = {
                **item,
                "name": re.sub(r"\s+", " ", display_name),
                "base": Decimal(0),
                "reduction": Decimal(0),
                "source_line_count": 0,
                "source_line_identities": [],
            }
            grouped[key] = target
        target["base"] += item["base"]
        target["source_line_count"] += 1
        target["source_line_identities"].append(item["line_identity"])

    output = list(grouped.values())
    for item in output:
        item["base"] = _decimal(item["base"], "sold_reduction_group_base")
        item["reduction"] = (
            item["base"] * (item["statutory_rate"] - item["reduced_rate"])
        ).quantize(ONE_DONG, rounding=ROUND_HALF_UP)
        if item["reduction"] == 0:
            item["reduction"] = Decimal(0)
    output.sort(key=lambda item: (
        item["name"].casefold(), item["statutory_rate"], item["reduced_rate"],
    ))
    return output


def build_vat_return_data(database: Path, tax_code: str, date_from: str, date_to: str) -> dict[str, Any]:
    totals = {key: Decimal(0) for key in ("purchase_base", "purchase_tax", "kct", "0", "5_base", "5_tax", "10_base", "10_tax", "kkknt")}
    unknown: set[tuple[str, str, str, str]] = set()
    grouped: dict[tuple[tuple[str, ...], str], dict[str, Any]] = {}
    trace: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    reduction_anomalies: list[dict[str, Any]] = []
    purchase_items: list[dict[str, Any]] = []
    purchase_audit: list[dict[str, Any]] = []
    invalid_purchase: list[dict[str, Any]] = []
    eligible_purchase_parents: dict[tuple[str, ...], tuple[sqlite3.Row, dict[str, Any]]] = {}
    excluded_purchase_identities: set[tuple[str, ...]] = set()
    purchase_reduction_lines: list[dict[str, Any]] = []
    purchase_reduction_audit: list[dict[str, Any]] = []
    invalid_purchase_reduction_lines: list[dict[str, Any]] = []
    sold_reduction_lines: list[dict[str, Any]] = []
    sold_reduction_audit: list[dict[str, Any]] = []
    missing_purchase_detail: list[dict[str, Any]] = []
    with closing(sqlite3.connect(database)) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """SELECT * FROM invoice_overview_items WHERE company_tax_code=?
               AND nlap_date BETWEEN ? AND ? AND direction IN ('purchase','sold')
               AND query_type IN ('query','sco-query') ORDER BY id""",
            (tax_code, date_from, date_to),
        ).fetchall()
        canonical_rows: dict[tuple[str, ...], list[sqlite3.Row]] = {}
        for row in rows:
            canonical_rows.setdefault(
                canonical_invoice_identity(tax_code, row['direction'], dict(row)), []
            ).append(row)
        detail_rows = connection.execute(
            """SELECT * FROM invoice_detail_items WHERE company_tax_code=?
               AND direction='sold' AND normalized_ready=1
               AND detail_outcome IN ('with_lines','valid_empty')
               AND (error_message IS NULL OR TRIM(error_message)='')""", (tax_code,),
        ).fetchall()
        details_by_identity: dict[tuple[str, ...], list[sqlite3.Row]] = {}
        for detail_row in detail_rows:
            details_by_identity.setdefault(
                canonical_invoice_identity(tax_code, 'sold', dict(detail_row)), []
            ).append(detail_row)
        missing_sold: list[dict[str, Any]] = []
        for identity, copies in canonical_rows.items():
            # Prefer a copy that has a matching detail source, but aggregate a
            # duplicate query/sco-query business invoice only once.
            eligible = []
            for candidate in copies:
                candidate_fields = _attributes(connection, int(candidate['id']))
                if not _excluded_status(candidate_fields.get('tthai')):
                    eligible.append((candidate, candidate_fields))
            if not eligible:
                excluded.append({"identity": masked_invoice_identity(identity), "reason": "excluded_status"})
                representative = copies[0]
                if representative["direction"] == "purchase":
                    excluded_purchase_identities.add(identity)
                    excluded_fields = _attributes(connection, int(representative["id"]))
                    purchase_audit.append({
                        "canonical_invoice_identity": masked_invoice_identity(identity),
                        "query_type": ",".join(sorted({str(item["query_type"]) for item in copies})),
                        "invoice_date": str(representative["nlap_date"] or "")[:10],
                        "model": str(representative["khmshdon"] or ""),
                        "symbol": str(representative["khhdon"] or ""),
                        "invoice_number": str(representative["shdon"] or ""),
                        "seller_name": str(excluded_fields.get("nbten") or ""),
                        "seller_tax_code_masked": _masked_tax_code(representative["nbmst"]),
                        "status_raw": str(excluded_fields.get("tthai") or ""),
                        "status_normalized": _normalized_text(excluded_fields.get("tthai")),
                        "base": None, "tax": None, "deductible_tax": None,
                        "exported": False, "reason": "excluded_status",
                    })
                continue
            row, fields = eligible[0]
            if row["direction"] == "purchase":
                invoice_date = str(_field(fields, row, "nlap_date", "nlap", "tdlap"))[:10]
                raw_number = str(_field(fields, row, "shdon"))
                raw_model = str(_field(fields, row, "khmshdon"))
                raw_symbol = str(_field(fields, row, "khhdon"))
                seller_tax_code = str(_field(fields, row, "nbmst"))
                required_missing = []
                for field_name, field_value in (
                    ("ngày lập", invoice_date), ("số hóa đơn", raw_number),
                    ("mẫu số", raw_model), ("ký hiệu", raw_symbol),
                    ("MST người bán", seller_tax_code),
                ):
                    if not field_value.strip():
                        required_missing.append(field_name)
                try:
                    date.fromisoformat(invoice_date)
                except ValueError:
                    required_missing.append("ngày lập không hợp lệ")
                totals_source = "overview"
                missing_money = [
                    field_name for field_name in ("tgtcthue", "tgtthue")
                    if fields.get(field_name) in (None, "")
                ]
                if missing_money:
                    recovered = _verified_purchase_totals_from_detail(
                        connection, tax_code, row, fields,
                    )
                    if recovered is not None:
                        recovered_base, recovered_tax, recovered_detail_id = recovered
                        fields = {
                            **fields,
                            "tgtcthue": recovered_base,
                            "tgtthue": recovered_tax,
                        }
                        totals_source = "detail_verified_fallback"
                    else:
                        recovered_detail_id = None
                for field_name in ("tgtcthue", "tgtthue"):
                    if fields.get(field_name) in (None, ""):
                        required_missing.append(field_name)
                if required_missing:
                    invalid_purchase.append({
                        "identity": masked_invoice_identity(identity),
                        "invoice_number_masked": "*" * max(0, len(raw_number) - 3) + raw_number[-3:],
                        "missing": sorted(set(required_missing)),
                    })
                    purchase_audit.append({
                        "canonical_invoice_identity": masked_invoice_identity(identity),
                        "query_type": ",".join(sorted({str(item[0]["query_type"]) for item in eligible})),
                        "invoice_date": invoice_date, "model": raw_model,
                        "symbol": raw_symbol, "invoice_number": raw_number,
                        "seller_name": str(_field(fields, row, "nbten")),
                        "seller_tax_code_masked": _masked_tax_code(seller_tax_code),
                        "status_raw": str(fields.get("tthai") or ""),
                        "status_normalized": _normalized_text(fields.get("tthai")),
                        "base": None, "tax": None, "deductible_tax": None,
                        "exported": False, "reason": "missing_required_fields",
                    })
                    continue
                base = _decimal(fields["tgtcthue"], "tgtcthue")
                tax = _decimal(fields["tgtthue"], "tgtthue")
                seller_name = str(_field(fields, row, "nbten"))
                description = (
                    f"Mua hàng của {seller_name} theo hóa đơn số {raw_number}"
                    if seller_name else f"Mua hàng theo hóa đơn số {raw_number}"
                )
                item = {
                    "identity": identity,
                    "identity_hash": masked_invoice_identity(identity),
                    "query_type": ",".join(sorted({str(item[0]["query_type"]) for item in eligible})),
                    "invoice_category": str(row["invoice_category"] or ""),
                    "khmshdon": raw_model, "khhdon": raw_symbol, "shdon": raw_number,
                    "nlap_date": invoice_date, "seller_name": seller_name,
                    "seller_tax_code": seller_tax_code,
                    "base": base, "tax": tax, "deductible_tax": tax,
                    "note": str(_field(fields, row, "ghichu", "gchu")),
                    "description": description,
                    "status_raw": str(fields.get("tthai") or ""),
                    "status_normalized": _normalized_text(fields.get("tthai")),
                    "totals_source": totals_source,
                    "totals_detail_id": recovered_detail_id if totals_source != "overview" else None,
                }
                purchase_items.append(item)
                eligible_purchase_parents[identity] = (row, fields)
                totals["purchase_base"] += base
                totals["purchase_tax"] += tax
                purchase_audit.append({
                    "canonical_invoice_identity": item["identity_hash"],
                    "query_type": item["query_type"], "invoice_date": invoice_date,
                    "model": raw_model, "symbol": raw_symbol,
                    "invoice_number": raw_number, "seller_name": seller_name,
                    "seller_tax_code_masked": _masked_tax_code(seller_tax_code),
                    "status_raw": item["status_raw"],
                    "status_normalized": item["status_normalized"],
                    "base": format(base, "f"), "tax": format(tax, "f"),
                    "deductible_tax": format(tax, "f"),
                    "exported": True,
                    "reason": "" if totals_source == "overview" else totals_source,
                    "totals_source": totals_source,
                })
                for duplicate_row, duplicate_fields in eligible[1:]:
                    purchase_audit.append({
                        "canonical_invoice_identity": item["identity_hash"],
                        "query_type": str(duplicate_row["query_type"]),
                        "invoice_date": str(duplicate_row["nlap_date"] or "")[:10],
                        "model": str(duplicate_row["khmshdon"] or ""),
                        "symbol": str(duplicate_row["khhdon"] or ""),
                        "invoice_number": str(duplicate_row["shdon"] or ""),
                        "seller_name": str(duplicate_fields.get("nbten") or ""),
                        "seller_tax_code_masked": _masked_tax_code(duplicate_row["nbmst"]),
                        "status_raw": str(duplicate_fields.get("tthai") or ""),
                        "status_normalized": _normalized_text(duplicate_fields.get("tthai")),
                        "base": None, "tax": None, "deductible_tax": None,
                        "exported": False, "reason": "duplicate_canonical",
                    })
                continue
            candidates = details_by_identity.get(identity, [])
            detail = next((item for item in candidates if item['query_type'] == row['query_type']), None)
            detail = detail or (candidates[0] if candidates else None)
            if detail is None:
                invoice_number = str(row['shdon'] or '')
                missing_sold.append({
                    'identity': masked_invoice_identity(identity),
                    'shdon_masked': ('*' * max(0, len(invoice_number) - 3) + invoice_number[-3:]),
                })
                continue
            lines = connection.execute(
                """SELECT id,line_number,ten,CAST(tsuat AS TEXT) AS tsuat,
                          CAST(thtien AS TEXT) AS thtien,CAST(tthue AS TEXT) AS tthue
                   FROM invoice_detail_lines WHERE detail_item_id=? ORDER BY line_number,id""",
                (detail["id"],),
            ).fetchall()
            groups = [_tax_group(line["tsuat"]) for line in lines]
            if not lines and detail['detail_outcome'] == 'valid_empty':
                continue
            if not lines or any(group is None for group in groups):
                unknown.add((str(row["direction"]), str(row["query_type"]), str(row["shdon"]), str(row["khhdon"])))
                continue
            for line, group in zip(lines, groups):
                base = _decimal_exact(line["thtien"], "thtien")
                tax = _decimal_exact(line["tthue"], "tthue")
                output_group = "10" if group == "8" else group
                actual_rate, actual_label = _actual_rate(base, tax)
                declared_rate = {"0": Decimal("0"), "5": Decimal("0.05"),
                                 "8": Decimal("0.08"), "10": Decimal("0.10")}.get(group)
                filing_rate = Decimal("0.10") if group in {"8", "10"} else declared_rate
                difference = None if filing_rate is None else filing_rate * base - tax
                reduced = False
                reduction_base = Decimal(0)
                reduction_anomaly = ""
                if group == "8":
                    reduced = True
                    reduction_base = base
                elif group == "10":
                    # Prefer the declared/full 10% interpretation when the
                    # five-dong bands overlap on very small invoice lines.
                    if _tax_matches_rate(base, tax, Decimal("0.10")):
                        pass
                    elif base != 0 and _tax_matches_rate(base, tax, Decimal("0.08")):
                        reduced = True
                        reduction_base = base
                    else:
                        reduction_anomaly = "declared_10_tax_is_neither_8_percent_nor_10_percent"
                elif group == "5" and not _tax_matches_rate(base, tax, Decimal("0.05")):
                    reduction_anomaly = "declared_5_tax_does_not_match_5_percent"
                trace_item = {
                    "canonical_invoice_identity": "|".join(identity),
                    "line_identity": f"{detail['id']}:{line['line_number']}",
                    "status": str(fields.get("tthai") or ""), "source_tax_rate": str(line["tsuat"]),
                    "filing_group": output_group, "base": format(base, "f"),
                    "actual_tax": format(tax, "f"),
                    "actual_rate_raw": None if base == 0 else format(tax / base, "f"),
                    "actual_rate_normalized": actual_label,
                    "difference": None if difference is None else format(difference, "f"),
                    "reduced": reduced,
                    "reduction_base": format(reduction_base, "f"),
                    "reduction_unrounded": format(reduction_base * Decimal("0.02"), "f"),
                    "reduction_anomaly": reduction_anomaly,
                    "output_group": output_group,
                    "exclusion_reason": "sheet_has_no_kkknt_group" if output_group == "kkknt" else "",
                }
                # Detail exports can contain prose/note and rounding-adjustment
                # rows carrying a nominal 0% rate but no reportable whole-VND
                # value. They are useful evidence, but they are not sale lines
                # and must not become all-zero entries in the VAT schedule.
                # Keep genuine 0% sales whenever their rounded revenue or tax
                # is non-zero.
                report_base = _decimal(base, "sold_schedule_line_base")
                report_tax = _decimal(tax, "sold_schedule_line_tax")
                if report_base == 0 and report_tax == 0:
                    trace_item["exclusion_reason"] = "zero_value_detail_line"
                    trace.append(trace_item)
                    continue
                key = (identity, str(output_group))
                target = grouped.setdefault(key, {
                    "identity": identity, "identity_hash": masked_invoice_identity(identity),
                    "group": output_group, "khmshdon": str(_field(fields, row, "khmshdon")),
                    "khhdon": str(_field(fields, row, "khhdon")),
                    "shdon": str(_field(fields, row, "shdon")),
                    "nlap_date": str(_field(fields, row, "nlap_date", "nlap", "tdlap"))[:10],
                    "buyer_name": str(_field(fields, row, "nmten", "nmtnmua")),
                    "buyer_tax_code": str(_field(fields, row, "nmmst")),
                    "note": str(_field(fields, row, "ghichu", "gchu")),
                    "description": str(_field(fields, row, "dgiai")),
                    "base_exact": Decimal(0), "tax_exact": Decimal(0),
                    "reduction_base_exact": Decimal(0),
                    "reduction": Decimal(0), "has_reduction": False,
                })
                target["base_exact"] += base
                target["tax_exact"] += tax
                target["reduction_base_exact"] += reduction_base
                target["has_reduction"] = target["has_reduction"] or reduced
                trace.append(trace_item)
                if reduced:
                    normalized_base = _decimal(base, "sold_reduction_thtien")
                    reduction = (normalized_base * Decimal("0.02")).quantize(
                        ONE_DONG, rounding=ROUND_HALF_UP,
                    )
                    sold_reduction_lines.append({
                        "parent_identity": identity,
                        "parent_identity_hash": masked_invoice_identity(identity),
                        "line_identity": trace_item["line_identity"],
                        "line_number": int(line["line_number"]),
                        "query_type": str(detail["query_type"]),
                        "nlap_date": str(_field(fields, row, "nlap_date", "nlap", "tdlap"))[:10],
                        "khhdon": str(_field(fields, row, "khhdon")),
                        "shdon": str(_field(fields, row, "shdon")),
                        "name": str(line["ten"] or "").strip(),
                        "base": normalized_base,
                        "statutory_rate": Decimal("0.10"),
                        "reduced_rate": Decimal("0.08"),
                        "reduction": reduction,
                    })
                    sold_reduction_audit.append({
                        "parent_invoice_identity": masked_invoice_identity(identity),
                        "line_identity": trace_item["line_identity"],
                        "name": str(line["ten"] or "").strip(),
                        "base": format(normalized_base, "f"),
                        "source_tax_rate": str(line["tsuat"] or ""),
                        "statutory_rate": "0.10",
                        "reduced_rate": "0.08",
                        "reduction": format(reduction, "f"),
                    })
                if reduction_anomaly:
                    reduction_anomalies.append({
                        "canonical_invoice_identity": masked_invoice_identity(identity),
                        "line_identity": trace_item["line_identity"],
                        "source_tax_rate": str(line["tsuat"]),
                        "base": format(base, "f"),
                        "actual_tax": format(tax, "f"),
                        "actual_rate_percent": actual_label,
                        "reason": reduction_anomaly,
                    })
        purchase_detail_headers = connection.execute(
            """SELECT * FROM invoice_detail_items WHERE company_tax_code=?
               AND direction='purchase' AND query_type IN ('query','sco-query')
               ORDER BY CASE query_type WHEN 'query' THEN 0 ELSE 1 END, id""",
            (tax_code,),
        ).fetchall()
        complete_purchase_details: set[tuple[str, ...]] = set()
        seen_purchase_lines: set[tuple[tuple[str, ...], int]] = set()
        for detail in purchase_detail_headers:
            identity = canonical_invoice_identity(tax_code, "purchase", dict(detail))
            lines = connection.execute(
                """SELECT id,line_number,ten,CAST(tsuat AS TEXT) AS tsuat,
                          CAST(thtien AS TEXT) AS thtien,CAST(tthue AS TEXT) AS tthue
                   FROM invoice_detail_lines WHERE detail_item_id=? ORDER BY line_number,id""",
                (detail["id"],),
            ).fetchall()
            detail_complete = (
                int(detail["normalized_ready"] or 0) == 1
                and detail["detail_outcome"] in ("with_lines", "valid_empty")
                and not str(detail["error_message"] or "").strip()
                and (detail["detail_outcome"] == "valid_empty" or bool(lines))
            )
            if detail_complete:
                complete_purchase_details.add(identity)
            parent = eligible_purchase_parents.get(identity)
            parent_status_reason = (
                "outside_period" if not (date_from <= str(detail["nlap_date"] or "")[:10] <= date_to)
                else "excluded_status" if identity in excluded_purchase_identities
                else "orphan_detail" if parent is None else ""
            )
            parent_row, parent_fields = parent if parent is not None else (None, {})
            invoice_date = str(parent_row["nlap_date"] if parent_row is not None else detail["nlap_date"] or "")[:10]
            for line in lines:
                line_number = int(line["line_number"])
                line_key = (identity, line_number)
                line_identity = f"{masked_invoice_identity(identity)}:{line_number}"
                raw_rate = line["tsuat"]
                normalized_group = _tax_group(raw_rate)
                raw_base = line["thtien"]
                raw_tax = line["tthue"]
                audit_item = {
                    "parent_invoice_identity": masked_invoice_identity(identity),
                    "line_identity": line_identity,
                    "query_type": str(detail["query_type"]),
                    "invoice_date": invoice_date,
                    "status_raw": str(parent_fields.get("tthai") or ""),
                    "status_normalized": _normalized_text(parent_fields.get("tthai")),
                    "name": str(line["ten"] or "").strip(),
                    "tax_rate_raw": "" if raw_rate is None else str(raw_rate),
                    "tax_rate_normalized": normalized_group,
                    "base": None, "tax": None,
                    "included": False, "reason": "",
                    "difference_from_8_percent": None,
                    "looks_like_8_percent": False,
                    "tax_difference_warning": False,
                }
                if parent_status_reason:
                    audit_item["reason"] = parent_status_reason
                    purchase_reduction_audit.append(audit_item)
                    continue
                if not detail_complete:
                    audit_item["reason"] = "detail_not_finalized"
                    purchase_reduction_audit.append(audit_item)
                    continue
                if line_key in seen_purchase_lines:
                    audit_item["reason"] = "duplicate_canonical_line"
                    purchase_reduction_audit.append(audit_item)
                    continue
                seen_purchase_lines.add(line_key)
                try:
                    base_exact = _decimal_exact(raw_base, "purchase_detail_thtien")
                    tax_exact = _decimal_exact(raw_tax, "purchase_detail_tthue")
                    base = _decimal(base_exact, "purchase_detail_thtien")
                    tax = _decimal(tax_exact, "purchase_detail_tthue")
                    expected_tax = (base * Decimal("0.08")).quantize(ONE_DONG, rounding=ROUND_HALF_UP)
                    difference = abs(tax - expected_tax)
                    audit_item["base"] = format(base, "f")
                    audit_item["tax"] = format(tax, "f")
                    audit_item["difference_from_8_percent"] = format(difference, "f")
                    audit_item["looks_like_8_percent"] = difference <= Decimal("5")
                    audit_item["tax_difference_warning"] = normalized_group == "8" and difference > Decimal("5")
                except ValueError:
                    if normalized_group == "8":
                        invalid_purchase_reduction_lines.append({
                            "parent_invoice_identity": masked_invoice_identity(identity),
                            "line_identity": line_identity,
                            "missing": ["thtien/tthue không hợp lệ"],
                        })
                    audit_item["reason"] = "invalid_money"
                    purchase_reduction_audit.append(audit_item)
                    continue
                if normalized_group != "8":
                    audit_item["reason"] = "tax_rate_not_8"
                    purchase_reduction_audit.append(audit_item)
                    continue
                missing_fields = []
                if not audit_item["name"]:
                    missing_fields.append("tên hàng hóa, dịch vụ")
                if raw_base in (None, ""):
                    missing_fields.append("thtien")
                if raw_tax in (None, ""):
                    missing_fields.append("tthue")
                if missing_fields:
                    invalid_purchase_reduction_lines.append({
                        "parent_invoice_identity": masked_invoice_identity(identity),
                        "line_identity": line_identity,
                        "missing": missing_fields,
                    })
                    audit_item["reason"] = "missing_required_fields"
                    purchase_reduction_audit.append(audit_item)
                    continue
                output = {
                    "parent_identity": identity,
                    "parent_identity_hash": masked_invoice_identity(identity),
                    "line_identity": line_identity,
                    "line_number": line_number,
                    "query_type": str(detail["query_type"]),
                    "nlap_date": invoice_date,
                    "khhdon": str(parent_row["khhdon"] or ""),
                    "shdon": str(parent_row["shdon"] or ""),
                    "name": audit_item["name"], "base": base, "tax": tax,
                    "difference_from_8_percent": difference,
                }
                purchase_reduction_lines.append(output)
                audit_item["included"] = True
                purchase_reduction_audit.append(audit_item)
        for identity, (parent_row, _) in eligible_purchase_parents.items():
            if identity not in complete_purchase_details:
                number = str(parent_row["shdon"] or "")
                missing_purchase_detail.append({
                    "identity": masked_invoice_identity(identity),
                    "shdon_masked": "*" * max(0, len(number) - 3) + number[-3:],
                })
    if invalid_purchase:
        raise ValueError('vat_return_purchase_invalid:' + json.dumps({
            'direction': 'purchase', 'count': len(invalid_purchase),
            'date_from': date_from, 'date_to': date_to,
            'examples': invalid_purchase[:5],
        }, ensure_ascii=False, separators=(',', ':')))
    if invalid_purchase_reduction_lines:
        raise ValueError('vat_return_purchase_reduction_invalid:' + json.dumps({
            'direction': 'purchase', 'count': len(invalid_purchase_reduction_lines),
            'date_from': date_from, 'date_to': date_to,
            'examples': invalid_purchase_reduction_lines[:5],
        }, ensure_ascii=False, separators=(',', ':')))
    if missing_sold:
        raise ValueError('vat_return_detail_missing:' + json.dumps({
            'direction': 'sold', 'count': len(missing_sold),
            'date_from': date_from, 'date_to': date_to,
            'examples': missing_sold[:3],
        }, ensure_ascii=False, separators=(',', ':')))
    if unknown:
        raise ValueError(f"vat_return_unknown_tax_rate:{len(unknown)}")
    sold_groups = {key: [] for key in ("kct", "0", "5", "10")}
    for item in grouped.values():
        item["base"] = _decimal(item.pop("base_exact"), "group_base")
        item["tax"] = _decimal(item.pop("tax_exact"), "group_tax")
        reduction_base = item.pop("reduction_base_exact")
        item["reduction"] = (
            (reduction_base * Decimal("0.02")).quantize(ONE_DONG, rounding=ROUND_HALF_UP)
            if item["has_reduction"] else Decimal(0)
        )
        if item["reduction"] == 0:
            item["reduction"] = Decimal(0)
        group = item["group"]
        if group == "kct": totals["kct"] += item["base"]
        elif group == "0": totals["0"] += item["base"]
        elif group == "5": totals["5_base"] += item["base"]; totals["5_tax"] += item["tax"]
        elif group == "10": totals["10_base"] += item["base"]; totals["10_tax"] += item["tax"]
        elif group == "kkknt": totals["kkknt"] += item["base"]
        if group in sold_groups:
            sold_groups[group].append(item)
    for items in sold_groups.values():
        items.sort(key=lambda item: (item["nlap_date"], item["khhdon"].casefold(),
                                     _invoice_number_key(item["shdon"]), item["identity"]))
    purchase_items.sort(key=lambda item: (
        item["nlap_date"], item["khhdon"].casefold(),
        _invoice_number_key(item["shdon"]), item["identity"],
    ))
    purchase_reduction_lines.sort(key=lambda item: (
        item["nlap_date"], item["khhdon"].casefold(),
        _invoice_number_key(item["shdon"]), item["line_number"], item["line_identity"],
    ))
    sold_reduction_source_line_count = len(sold_reduction_lines)
    sold_reduction_lines = _group_sold_reduction_lines(sold_reduction_lines)
    return {"totals": totals, "sold_groups": sold_groups,
            "trace": trace, "excluded": excluded,
            "reduction_anomalies": reduction_anomalies,
            "purchase_items": purchase_items,
            "purchase_audit": purchase_audit,
            "purchase_reduction_lines": purchase_reduction_lines,
            "purchase_reduction_audit": purchase_reduction_audit,
            "sold_reduction_lines": sold_reduction_lines,
            "sold_reduction_source_line_count": sold_reduction_source_line_count,
            "sold_reduction_audit": sold_reduction_audit,
            "missing_purchase_detail": missing_purchase_detail}


def aggregate_vat_return(database: Path, tax_code: str, date_from: str, date_to: str) -> dict[str, Decimal]:
    return build_vat_return_data(database, tax_code, date_from, date_to)["totals"]


def _vat_return_cell_values(totals: dict[str, Decimal], company_name: str,
                            tax_code: str) -> dict[str, Any]:
    zero = Decimal(0)
    h17 = totals["purchase_tax"]
    f20 = totals["0"] + totals["5_base"] + totals["10_base"] + totals["kkknt"]
    h20 = totals["5_tax"] + totals["10_tax"]
    f25 = totals["kct"] + f20
    h25 = h20
    h26 = h25 - h17
    s_value = h26  # [22], [37], [38], and [39a] are blank and equal zero.
    h32 = max(s_value, zero)
    h35 = max(-s_value, zero)
    h38 = h35  # [42] at H36 is blank and equal zero.
    return {
        "A1": company_name, "A2": f"Mã số thuế: {tax_code}", "H11": None,
        "F14": totals["purchase_base"], "H14": totals["purchase_tax"],
        "F16": zero, "H16": zero,
        "H17": CachedFormula("=H14", h17),
        "F19": totals["kct"],
        "F20": CachedFormula("=F21+F22+F23+F24", f20),
        "H20": CachedFormula("=H22+H23", h20),
        "F21": totals["0"], "F22": totals["5_base"], "H22": totals["5_tax"],
        "F23": totals["10_base"], "H23": totals["10_tax"], "F24": totals["kkknt"],
        "F25": CachedFormula("=F19+F20", f25),
        "H25": CachedFormula("=H20", h25),
        "H26": CachedFormula("=H25-H17", h26),
        "H28": None, "H29": None, "H30": None,
        "H32": CachedFormula(OBLIGATION_FORMULAS["H32"], h32),
        "H33": None,
        "H34": CachedFormula("=H32-H33", h32),
        "H35": CachedFormula(OBLIGATION_FORMULAS["H35"], h35),
        "H36": None,
        "H38": CachedFormula(OBLIGATION_FORMULAS["H38"], h38),
    }


def _publish_workbook(temporary: Path, target: Path) -> Path:
    """Atomically replace the deterministic target without deleting it first."""
    try:
        os.replace(temporary, target)
        return target
    except PermissionError as error:
        # The temporary workbook was already created and verified in the same
        # directory, so a replace failure against an existing writable target
        # is the Windows sharing/lock case. Keep the original file untouched.
        target_is_read_only = target.exists() and not os.access(target, os.W_OK)
        if target.exists() and not target_is_read_only:
            detail = json.dumps({"filename": target.name, "path": str(target)}, ensure_ascii=False, separators=(",", ":"))
            raise ValueError(f"vat_return_destination_file_locked:{detail}") from error
        raise ValueError("vat_return_destination_not_writable") from error


def vat_return_filename(tax_code: str, date_from: str, date_to: str) -> str:
    """Return the single deterministic filename for one taxpayer and period."""
    try:
        from_value = date.fromisoformat(date_from).strftime("%d-%m-%Y")
        to_value = date.fromisoformat(date_to).strftime("%d-%m-%Y")
    except ValueError as error:
        raise ValueError("invalid_vat_return_range") from error
    return f"To_khai_thue_GTGT_{tax_code}_{from_value}_{to_value}.xlsx"


def export_vat_return(backend, value: dict[str, Any]) -> dict[str, Any]:
    connection_ids = list(value.get("connection_ids") or ())
    if len(connection_ids) != 1: raise ValueError("invalid_vat_return_account")
    selected_destination = Path(str(value.get("destination") or ""))
    if not selected_destination.is_absolute(): raise ValueError("invalid_artifact_directory")
    date_from, date_to = str(value["date_from"]), str(value["date_to"])
    connection_id = str(connection_ids[0]); tax_code = backend.connection_tax_code(connection_id)
    destination = direction_export_directory(
        selected_destination, tax_code, ("purchase", "sold")
    )
    coverage = ArtifactInspector(backend).vat_return_coverage({"connection_ids": connection_ids, "date_from": date_from, "date_to": date_to})["accounts"][0]
    missing = []
    # The workbook now consumes purchase Detail in the 8% reduction schedule,
    # therefore every requested direction/scope must be complete before export.
    for direction in ("purchase", "sold"):
        for item in coverage[direction]["missing"]:
            missing.append({"tax_code": tax_code, "direction": direction, **item})
    if missing: raise ValueError("vat_return_coverage_missing:" + json.dumps(missing, ensure_ascii=False, separators=(",", ":")))
    # Display metadata is persisted locally when the account is connected.
    # Do not call list_connections(), whose missing-name backfill may contact
    # the portal; VAT export must be SQLite/local-metadata only.
    names = backend._load_company_names() if hasattr(backend, "_load_company_names") else {}
    company_name = str(names.get(connection_id) or "").strip()
    if not company_name: raise ValueError("vat_return_company_name_missing")
    report = build_vat_return_data(
        backend.data_root / tax_code / "db" / "invoices.sqlite3",
        tax_code, date_from, date_to,
    )
    if report["missing_purchase_detail"]:
        raise ValueError('vat_return_detail_missing:' + json.dumps({
            'direction': 'purchase', 'count': len(report["missing_purchase_detail"]),
            'date_from': date_from, 'date_to': date_to,
            'examples': report["missing_purchase_detail"][:3],
        }, ensure_ascii=False, separators=(',', ':')))
    totals = report["totals"]
    template = _template_path()
    if not template.is_file(): raise FileNotFoundError(template)
    values = _vat_return_cell_values(totals, company_name, tax_code)
    try:
        destination.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        if error.errno in {errno.EACCES, errno.EPERM, errno.EROFS}:
            raise ValueError("vat_return_destination_not_writable") from error
        raise
    target = destination / vat_return_filename(tax_code, date_from, date_to)
    try:
        descriptor, temporary_name = tempfile.mkstemp(prefix=f".{target.stem}-", suffix=".xlsx", dir=destination)
        os.close(descriptor)
    except OSError as error:
        if error.errno in {errno.EACCES, errno.EPERM, errno.EROFS}:
            raise ValueError("vat_return_destination_not_writable") from error
        raise
    temporary = Path(temporary_name)
    try:
        _write_template_workbook(
            template, temporary, values, company_name=company_name,
            tax_code=tax_code, sold_groups=report["sold_groups"],
            purchase_items=report["purchase_items"],
            purchase_reduction_lines=report["purchase_reduction_lines"],
            sold_reduction_lines=report["sold_reduction_lines"],
        )
        verification = load_workbook(temporary, data_only=False, read_only=True)
        cached_verification = load_workbook(temporary, data_only=True, read_only=True)
        original = load_workbook(template, read_only=True)
        valid = (
            verification.sheetnames == original.sheetnames
            and verification[SHEET_NAME]["H17"].value == "=H14"
            and Decimal(str(cached_verification[SHEET_NAME]["H35"].value)) == values["H35"].value
            and Decimal(str(cached_verification[SHEET_NAME]["H38"].value)) == values["H38"].value
            and verification[SOLD_SHEET_NAME]["A1"].value == company_name
            and verification[SOLD_SHEET_NAME].max_column == 16
            and verification[PURCHASE_SHEET_NAME]["A1"].value == company_name
            and verification[PURCHASE_SHEET_NAME].max_column == 15
            and verification[REDUCTION_SHEET_NAME]["A1"].value == company_name
            and verification[REDUCTION_SHEET_NAME].max_column == 12
        )
        expected_sheetnames = list(verification.sheetnames)
        original.close()
        cached_verification.close()
        if not valid:
            raise RuntimeError("vat_return_workbook_verification_failed")
        verification.close()
        target = _publish_workbook(temporary, target)
        published = load_workbook(target, data_only=False, read_only=True)
        try:
            if published.sheetnames != expected_sheetnames:
                raise RuntimeError("vat_return_workbook_verification_failed")
        finally:
            published.close()
    finally: temporary.unlink(missing_ok=True)
    return {
        "count": 1,
        "files": [str(target)],
        "audit": {
            "reduction_anomaly_count": len(report["reduction_anomalies"]),
            "reduction_anomalies": report["reduction_anomalies"],
            "purchase_invoice_count": len(report["purchase_items"]),
            "purchase_base": format(report["totals"]["purchase_base"], "f"),
            "purchase_tax": format(report["totals"]["purchase_tax"], "f"),
            "purchase_deductible_tax": format(report["totals"]["purchase_tax"], "f"),
            "purchase_rows": report["purchase_audit"],
            "purchase_reduction_line_count": len(report["purchase_reduction_lines"]),
            "purchase_reduction_base": format(sum(
                (item["base"] for item in report["purchase_reduction_lines"]), Decimal(0)
            ), "f"),
            "purchase_reduction_tax": format(sum(
                (item["tax"] for item in report["purchase_reduction_lines"]), Decimal(0)
            ), "f"),
            "purchase_reduction_rows": report["purchase_reduction_audit"],
            "sold_reduction_line_count": len(report["sold_reduction_lines"]),
            "sold_reduction_source_line_count": report["sold_reduction_source_line_count"],
            "sold_reduction_base": format(sum(
                (item["base"] for item in report["sold_reduction_lines"]), Decimal(0)
            ), "f"),
            "sold_reduction_tax": format(sum(
                (item["reduction"] for item in report["sold_reduction_lines"]), Decimal(0)
            ), "f"),
            "sold_reduction_rows": report["sold_reduction_audit"],
        },
    }
