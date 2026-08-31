"""Customer-compatible final formatting for Detail invoice workbooks.

The source exporter remains authoritative for parsing and row construction.
This module only normalizes its staged workbooks to the long-standing customer
layout: one header row, 37 fixed columns, and no content-based sizing.
"""

from __future__ import annotations

import os
import tempfile
import urllib.parse
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.exporters.invoice_detail_excel_exporter import (
    HEADER_ALIASES,
    INTERNAL_KEYS,
    InvoiceDetailExcelExporter,
    _normalize_header,
)


DETAIL_EXPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("khmshdon", "Mẫu số HD"),
    ("khhdon", "Ký hiệu hóa đơn"),
    ("shdon", "Số hóa đơn"),
    ("ntao", "Ngày lập hóa đơn"),
    ("nky", "Ngày người bán ký số"),
    ("mhdon", "MCCQT"),
    ("nky", "Ngày CQT ký số"),
    ("dvtte", "Đơn vị tiền tệ"),
    ("tgia", "Tỷ giá"),
    ("nbten", "Tên người bán"),
    ("nbmst", "MST người bán"),
    ("nbdchi", "Địa chỉ người bán"),
    ("nmten", "Tên người mua"),
    ("nmmst", "MST người mua"),
    ("nmdchi", "Địa chỉ người mua"),
    ("m_VT", "Mã VT"),
    ("ten", "Tên hàng hóa, dịch vụ"),
    ("dvtinh", "Đơn vị tính"),
    ("sluong", "Số lượng"),
    ("dgia", "Đơn giá"),
    ("stckhau", "Chiết khấu"),
    ("tsuat", "Thuế suất"),
    ("thtien", "Thành tiền chưa thuế"),
    ("tthue", "Tiền thuế"),
    ("ttcktmai", "Tổng tiền CKTM"),
    ("tgtphi", "Tổng tiền phí"),
    ("tgtttbso", "Tổng tiền thanh toán"),
    ("tthai", "Trạng thái hóa đơn"),
    ("ttxly", "Kết quả kiểm tra hóa đơn"),
    ("url", "url tra cứu hóa đơn"),
    ("mk", "Mã tra cứu"),
    ("ghichu", "Ghi chú 1"),
    ("thtttoan", "Hình thức thanh toán"),
    ("tchat", "Tính chất"),
    ("dgiai", "Ghi chú 2"),
    ("slo", "Số lô"),
    ("hdung", "Hạn dùng"),
)

# Exact A:AK widths read from the customer-supplied legacy workbook. They are
# deliberately static; long company names must never resize the export.
DETAIL_TEMPLATE_WIDTHS: tuple[float, ...] = (
    12.5703125, 21.7109375, 19.85546875, 13.0, 20.5703125,
    35.0, 44.7109375, 19.85546875, 13.0, 36.85546875,
    19.28515625, 63.85546875, 44.28515625, 28.7109375,
    51.42578125, 14.7109375, 57.28515625, 19.85546875, 13.0,
    13.0, 13.0, 13.0, 13.0, 13.0, 13.0, 13.0, 13.0,
    39.42578125, 60.7109375, 34.7109375, 23.5703125,
    21.42578125, 26.140625, 23.85546875, 36.42578125,
    18.85546875, 15.5703125,
)

_MONEY_COLUMN_INDEXES = frozenset((20, 21, 23, 24, 27))
_THIN_SIDE = Side(style="thin")
_CELL_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)


class FixedDetailExcelExporter(InvoiceDetailExcelExporter):
    """Source row generation without its runtime content-based autofit pass."""

    @staticmethod
    def _fit_cells(
        worksheet: Any,
        *,
        header_row: int,
        first_data_row: int,
        last_data_row: int,
        column_keys: list[str],
    ) -> None:
        del worksheet, header_row, first_data_row, last_data_row, column_keys


def _header_key(value: Any) -> str | None:
    normalized = _normalize_header(value)
    key = HEADER_ALIASES.get(normalized)
    if key is None and str(value or "").strip() in INTERNAL_KEYS:
        key = str(value).strip()
    return key


def _source_columns(worksheet: Any, header_row: int) -> list[int]:
    """Resolve all 37 source columns in order, including the duplicate nky."""
    positions: dict[str, list[int]] = {}
    for column in range(1, worksheet.max_column + 1):
        key = _header_key(worksheet.cell(header_row, column).value)
        if key and key != "stt":
            positions.setdefault(key, []).append(column)

    used: dict[str, int] = {}
    columns: list[int] = []
    for key, _label in DETAIL_EXPORT_COLUMNS:
        occurrence = used.get(key, 0)
        candidates = positions.get(key, [])
        if occurrence >= len(candidates):
            raise RuntimeError(f"Missing Detail export column: {key}")
        columns.append(candidates[occurrence])
        used[key] = occurrence + 1
    return columns


def _apply_header_style(cell: Any) -> None:
    cell.font = Font(size=12, bold=False)
    cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
    cell.border = _CELL_BORDER
    cell.number_format = "General"


def _apply_data_style(cell: Any, column_index: int) -> None:
    cell.font = Font(size=12, bold=False)
    cell.fill = PatternFill(fill_type=None)
    cell.alignment = Alignment(wrap_text=False)
    cell.border = _CELL_BORDER
    cell.number_format = "#,##0" if column_index in _MONEY_COLUMN_INDEXES else "General"


def combine_detail_workbooks_atomically(
    staged_jobs: list[tuple[str, str, Path]],
    target: Path,
) -> None:
    """Append staged Detail rows into one customer-format workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    for column_index, ((_key, label), width) in enumerate(
        zip(DETAIL_EXPORT_COLUMNS, DETAIL_TEMPLATE_WIDTHS, strict=True),
        start=1,
    ):
        cell = worksheet.cell(1, column_index, label)
        _apply_header_style(cell)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    output_row = 2
    try:
        for _direction, _category, staged_path in staged_jobs:
            source_book = load_workbook(staged_path, data_only=False, read_only=False)
            try:
                source_sheet = source_book.worksheets[0]
                header_row = InvoiceDetailExcelExporter._find_header_row(source_sheet)
                source_columns = _source_columns(source_sheet, header_row)
                for row_index in range(header_row + 1, source_sheet.max_row + 1):
                    values = [source_sheet.cell(row_index, column).value for column in source_columns]
                    if all(value in (None, "") for value in values):
                        continue
                    for column_index, (source_column, value) in enumerate(
                        zip(source_columns, values, strict=True), start=1
                    ):
                        source_cell = source_sheet.cell(row_index, source_column)
                        target_cell = worksheet.cell(output_row, column_index, value)
                        _apply_data_style(target_cell, column_index)
                        if column_index == 30 and value not in (None, ""):
                            link = source_cell.hyperlink.target if source_cell.hyperlink else str(value)
                            if urllib.parse.urlsplit(link).scheme in {"http", "https"}:
                                target_cell.hyperlink = link
                    output_row += 1
            finally:
                source_book.close()

        worksheet.auto_filter.ref = None
        worksheet.freeze_panes = None
        worksheet.sheet_view.zoomScale = 100
        worksheet.sheet_view.zoomScaleNormal = 100
        target.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent
        )
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()


__all__ = [
    "DETAIL_EXPORT_COLUMNS",
    "DETAIL_TEMPLATE_WIDTHS",
    "FixedDetailExcelExporter",
    "combine_detail_workbooks_atomically",
]
