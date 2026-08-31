"""Desktop detail-Excel renderer matching the customer's legacy workbook.

The source row builder remains authoritative for invoice/detail semantics. This
adapter only owns the customer-facing workbook shape: one header row, 37 fixed
columns, fixed widths, and no report metadata or automatic sizing.
"""

from __future__ import annotations

import json
import logging
import sys
import time
from copy import copy
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlsplit

VENDOR_ROOT = Path(__file__).resolve().parent / "vendor" / "mia_crawl_service"
if str(VENDOR_ROOT) not in sys.path:
    sys.path.insert(0, str(VENDOR_ROOT))

from app.exporters.invoice_detail_excel_exporter import InvoiceDetailExcelExporter
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("mia.excel_export")
ProgressCallback = Callable[[str, int, int], None]

DETAIL_EXPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("khmshdon", "Mẫu số HD"),
    ("khhdon", "Ký hiệu hóa đơn"),
    ("shdon", "Số hóa đơn"),
    ("ntao", "Ngày lập hóa đơn"),
    ("nky", "Ngày người bán ký số"),
    ("mhdon", "MCCQT"),
    ("nky", "Ngày CQT ký số"),
    ("dvtte", "Đơn vị tiền tệ"),
    ("tgia", "Tỷ giá"),
    ("nbten", "Tên người bán"),
    ("nbmst", "MST người bán"),
    ("nbdchi", "Địa chỉ người bán"),
    ("nmten", "Tên người mua"),
    ("nmmst", "MST người mua"),
    ("nmdchi", "Địa chỉ người mua"),
    ("m_VT", "Mã VT"),
    ("ten", "Tên hàng hóa, dịch vụ"),
    ("dvtinh", "Đơn vị tính"),
    ("sluong", "Số lượng"),
    ("dgia", "Đơn giá"),
    ("stckhau", "Chiết khấu"),
    ("tsuat", "Thuế suất"),
    ("thtien", "Thành tiền chưa thuế"),
    ("tthue", "Tiền thuế"),
    ("ttcktmai", "Tổng tiền CKTM"),
    ("tgtphi", "Tổng tiền phí"),
    ("tgtttbso", "Tổng tiền thanh toán"),
    ("tthai", "Trạng thái hóa đơn"),
    ("ttxly", "Kết quả kiểm tra hóa đơn"),
    ("url", "url tra cứu hóa đơn"),
    ("mk", "Mã tra cứu"),
    ("ghichu", "Ghi chú 1"),
    ("thtttoan", "Hình thức thanh toán"),
    ("tchat", "Tính chất"),
    ("dgiai", "Ghi chú 2"),
    ("slo", "Số lô"),
    ("hdung", "Hạn dùng"),
)

# Read directly from 2500709775_HDCTMuaVao 1-1-2026_4-3-2026.xlsx. These are
# static template dimensions, never recalculated from runtime cell contents.
DETAIL_TEMPLATE_WIDTHS: tuple[float, ...] = (
    12.5703125, 21.7109375, 19.85546875, 13.0, 20.5703125,
    35.0, 44.7109375, 19.85546875, 13.0, 36.85546875,
    19.28515625, 63.85546875, 44.28515625, 28.7109375,
    51.42578125, 14.7109375, 57.28515625, 19.85546875, 13.0,
    13.0, 13.0, 13.0, 13.0, 13.0, 13.0, 13.0, 13.0,
    39.42578125, 60.7109375, 34.7109375, 23.5703125,
    21.42578125, 26.140625, 23.85546875, 36.42578125,
    18.85546875, 15.5703125,
)


class CustomerDetailExcelExporter(InvoiceDetailExcelExporter):
    """Render detail rows in the exact compact legacy customer format."""

    def __init__(self, template_path, row_builder=None, *, progress=None) -> None:
        super().__init__(template_path, row_builder=row_builder)
        self.progress: ProgressCallback | None = progress

    def export(self, detail_records, output_path, from_date, to_date):
        self._validate_inputs(output_path, from_date, to_date)
        started = time.perf_counter()
        self._emit("load_template", 0, 1)
        workbook = load_workbook(self.template_path)
        self._emit("load_template", 1, 1)
        try:
            worksheet = workbook.worksheets[0]
            self._reset_sheet(worksheet)
            self._write_header(worksheet)

            rows: list[dict[str, Any]] = []
            invoice_count = 0
            skipped_count = 0
            total_records = len(detail_records)
            self._emit("build_rows", 0, total_records)
            for record_index, record in enumerate(detail_records, start=1):
                raw_path = Path(str(record.get("raw_detail_path") or ""))
                try:
                    payload = json.loads(raw_path.read_text(encoding="utf-8"))
                    if not isinstance(payload, dict):
                        raise ValueError("detail JSON root is not an object")
                    invoice_rows = self.row_builder.build_rows(payload, record)
                except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as error:
                    skipped_count += 1
                    logger.warning(
                        "Skipping unreadable invoice detail path=%s error=%s",
                        raw_path, type(error).__name__,
                    )
                    self._emit("build_rows", record_index, total_records)
                    continue
                if not invoice_rows:
                    skipped_count += 1
                else:
                    invoice_count += 1
                    rows.extend(invoice_rows)
                self._emit("build_rows", record_index, total_records)

            self._emit("write_rows", 0, len(rows))
            for row_index, row_values in enumerate(rows, start=2):
                self._write_data_row(worksheet, row_index, row_values)
                self._emit("write_rows", row_index - 1, len(rows))

            self._emit("format", 0, len(DETAIL_EXPORT_COLUMNS))
            for column_index, width in enumerate(DETAIL_TEMPLATE_WIDTHS, start=1):
                dimension = worksheet.column_dimensions[get_column_letter(column_index)]
                dimension.width = width
                dimension.hidden = False
                self._emit("format", column_index, len(DETAIL_EXPORT_COLUMNS))

            self._emit("save", 0, 1)
            self._save_atomically(workbook, output_path)
            self._emit("save", 1, 1)
        finally:
            workbook.close()

        summary = {
            "output_path": str(Path(output_path).resolve()),
            "invoice_count": invoice_count,
            "row_count": len(rows),
            "skipped_count": skipped_count,
        }
        logger.info(
            "excel_export_summary scope=details invoice_count=%s row_count=%s "
            "column_count=%s fixed_template_widths=true total_ms=%.1f",
            invoice_count, len(rows), len(DETAIL_EXPORT_COLUMNS),
            (time.perf_counter() - started) * 1000,
        )
        return summary

    @staticmethod
    def _reset_sheet(worksheet: Any) -> None:
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))
        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)
        if worksheet.max_column > len(DETAIL_EXPORT_COLUMNS):
            worksheet.delete_cols(
                len(DETAIL_EXPORT_COLUMNS) + 1,
                worksheet.max_column - len(DETAIL_EXPORT_COLUMNS),
            )
        worksheet.auto_filter.ref = None
        worksheet.freeze_panes = None

    @staticmethod
    def _write_header(worksheet: Any) -> None:
        thin = Side(style="thin", color="FF000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
        for column_index, (_key, title) in enumerate(DETAIL_EXPORT_COLUMNS, start=1):
            cell = worksheet.cell(1, column_index, title)
            cell.font = Font(size=12, bold=False)
            cell.fill = copy(fill)
            cell.border = copy(border)
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
        worksheet.row_dimensions[1].height = None

    def _write_data_row(self, worksheet: Any, row_index: int, values: dict[str, Any]) -> None:
        thin = Side(style="thin", color="FF000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for column_index, (key, _title) in enumerate(DETAIL_EXPORT_COLUMNS, start=1):
            cell = worksheet.cell(row_index, column_index)
            value = self._safe_excel_value(values.get(key, ""))
            cell.value = value
            cell.font = Font(size=12, bold=False)
            cell.border = copy(border)
            cell.alignment = Alignment(wrap_text=False)
            if key == "shdon":
                cell.number_format = "@"
            if (
                key == "url"
                and isinstance(value, str)
                and urlsplit(value).scheme.casefold() in {"http", "https"}
            ):
                cell.hyperlink = value
        worksheet.row_dimensions[row_index].height = None

    def _emit(self, phase: str, processed: int, total: int) -> None:
        if self.progress is None:
            return
        try:
            self.progress(phase, max(0, int(processed)), max(0, int(total)))
        except Exception:
            logger.exception("excel_export_progress_callback_failed phase=%s", phase)


__all__ = [
    "CustomerDetailExcelExporter",
    "DETAIL_EXPORT_COLUMNS",
    "DETAIL_TEMPLATE_WIDTHS",
]
