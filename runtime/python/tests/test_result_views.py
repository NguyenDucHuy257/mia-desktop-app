import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

import mia_runtime
from app.job_engine.models import JobRecord
from app.repositories.invoice_detail_repository import InvoiceDetailRepository
from app.repositories.invoice_overview_repository import InvoiceOverviewRepository
from app.services.overview_downloader import _find_header_row
from mia_backend import ProductionBackend
from mia_source_results import _DETAIL_EXPORT_COLUMNS, _available_path


class SourceJobIntentTests(unittest.TestCase):
    def test_refresh_options_are_forwarded_to_source_create_job_body(self):
        backend = object.__new__(ProductionBackend)
        source_job = SimpleNamespace(job_id="job-source")
        backend.service = Mock()
        backend.service.create_job.return_value = source_job
        with patch.object(
            ProductionBackend,
            "public_job",
            return_value={"job_id": "job-source", "status": "queued"},
        ):
            result = backend.start({
                "idempotency_key": "source-policy-test",
                "intent": {
                    "connection_id": "conn_123456",
                    "date_from": "2026-07-15",
                    "date_to": "2026-08-10",
                    "directions": ["purchase", "sold"],
                    "query_types": ["query", "sco-query"],
                    "scopes": ["overview", "detail"],
                    "data_types": ["invoice"],
                    "force_refresh": False,
                    "refresh_latest_month": True,
                },
            })
        self.assertEqual(result["job_id"], "job-source")
        body = backend.service.create_job.call_args.args[0]
        self.assertEqual(body.connection_id, "conn_123456")
        self.assertFalse(body.force_refresh)
        self.assertTrue(body.refresh_latest_month)
        self.assertEqual(body.result_scope, "detail")
        self.assertFalse(body.include_xml)
        backend.service.create_job.assert_called_once()

    def test_force_refresh_is_not_rewritten_by_desktop_policy(self):
        backend = object.__new__(ProductionBackend)
        backend.service = Mock(return_value=None)
        backend.service.create_job.return_value = SimpleNamespace(job_id="job-source")
        with patch.object(
            ProductionBackend,
            "public_job",
            return_value={"job_id": "job-source", "status": "queued"},
        ):
            backend.start({
                "idempotency_key": "source-force-refresh-test",
                "intent": {
                    "connection_id": "conn_123456",
                    "date_from": "2023-01-01",
                    "date_to": "2026-08-31",
                    "directions": ["purchase"],
                    "query_types": ["query"],
                    "scopes": ["overview"],
                    "data_types": ["invoice"],
                    "force_refresh": True,
                    "refresh_latest_month": False,
                },
            })
        body = backend.service.create_job.call_args.args[0]
        self.assertTrue(body.force_refresh)
        self.assertFalse(body.refresh_latest_month)


class ResultViewTests(unittest.TestCase):
    def test_result_filename_collision_starts_with_copy_two(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            (destination / "result.xlsx").write_bytes(b"existing")
            self.assertEqual(
                _available_path(destination, "result.xlsx").name,
                "result (2).xlsx",
            )

    @staticmethod
    def source_job() -> JobRecord:
        return JobRecord(
            job_id="job-latest",
            account_key="conn_account_1",
            company_tax_code="0100000000",
            job_type="invoice_crawl",
            queue_order=1,
            parameters={
                "connection_id": "conn_account_1",
                "date_from": "2026-01-01",
                "date_to": "2026-01-31",
                "directions": ["purchase", "sold"],
                "query_types": ["query", "sco-query"],
            },
            status="completed",
            current_stage="finalize",
            worker_id=None,
            lease_token=None,
            lease_generation=0,
            lease_expires_at=None,
            available_at="2026-08-20T10:00:00+00:00",
            cancel_requested_at=None,
            warning_count=0,
            created_at="2026-08-20T10:00:00+00:00",
            updated_at="2026-08-20T10:00:00+00:00",
            started_at="2026-08-20T10:00:00+00:00",
            finished_at="2026-08-20T10:05:00+00:00",
            last_error_code=None,
            last_error_message=None,
            owner_id="mia-desktop-local",
            pipeline_version=2,
            progress_state={
                "version": 3,
                "modules": {
                    name: {
                        "status": "completed",
                        "months": [{
                            "from_date": "2025-01-01",
                            "to_date": "2026-12-31",
                            "status": "completed",
                        }],
                    }
                    for name in ("overview", "detail")
                },
            },
        )

    def backend_with_job(self, *, data_root: Path | None = None):
        backend = object.__new__(ProductionBackend)
        backend.data_root = data_root or Path("source-data")
        job = self.source_job()
        backend.repository = SimpleNamespace(
            list_jobs_for_reconciliation=lambda: [job]
        )
        return backend, job

    def test_result_range_and_source_filters_override_latest_job_view(self):
        backend, _ = self.backend_with_job()
        reader = Mock()
        reader.overview_page.return_value = {
            "items": [{
                "id": 7,
                "nbmst": "0300000000",
                "khhdon": "AA/26E",
                "shdon": "12",
                "khmshdon": "1",
            }],
            "total_count": 1,
            "pagination": {"has_more": False, "next_cursor": None},
        }

        with patch(
            "app.external_api.results.JobResultReader", return_value=reader
        ):
            result = backend.results("overview", {
                "connection_id": "conn_account_1",
                "date_from": "2026-02-01",
                "date_to": "2026-02-28",
                "direction": "sold",
                "query_type": "sco-query",
                "limit": 50,
            })

        read_job = reader.overview_page.call_args.args[0]
        self.assertEqual(read_job.parameters["date_from"], "2026-02-01")
        self.assertEqual(read_job.parameters["date_to"], "2026-02-28")
        self.assertEqual(read_job.parameters["directions"], ["sold"])
        self.assertEqual(read_job.parameters["query_types"], ["sco-query"])
        self.assertEqual(result["items"][0]["direction"], "sold")
        self.assertTrue(result["columns"])
        self.assertTrue(result["column_labels"])
        self.assertEqual(result["total_count"], 1)

    def test_combined_results_dedupe_canonical_invoices_before_global_pagination(self):
        backend, _ = self.backend_with_job()

        class CombinedReader:
            @staticmethod
            def overview_page(job, *, limit, cursor):
                self.assertEqual(job.parameters["query_types"], ["query", "sco-query"])
                return {
                    "items": [
                        {"id": 1, "query_type": "query", "nlap_date": "2026-01-02", "nbmst": "010-001", "khmshdon": "01", "khhdon": " aa/26 ", "shdon": "0002"},
                        {"id": 2, "query_type": "sco-query", "nlap_date": "2026-01-02", "nbmst": "010001", "khmshdon": "1", "khhdon": "AA26", "shdon": "2"},
                        {"id": 3, "query_type": "sco-query", "nlap_date": "2026-01-03", "nbmst": "010002", "khmshdon": "1", "khhdon": "BB26", "shdon": "3"},
                    ],
                    "pagination": {"has_more": False, "next_cursor": None},
                }

        with patch("app.external_api.results.JobResultReader", return_value=CombinedReader()):
            query = {
                "connection_id": "conn_account_1", "date_from": "2026-01-01",
                "date_to": "2026-01-31", "direction": "purchase",
                "query_type": None, "query_types": ["query", "sco-query"], "limit": 1,
            }
            first = backend.results("overview", query)
            second = backend.results("overview", {**query, "cursor": first["pagination"]["next_cursor"]})

        self.assertEqual(first["total_count"], 2)
        self.assertEqual(len(first["items"]), 1)
        self.assertTrue(first["pagination"]["has_more"])
        self.assertEqual(len(second["items"]), 1)
        self.assertFalse(second["pagination"]["has_more"])
        self.assertNotEqual(first["items"][0]["invoice_key"], second["items"][0]["invoice_key"])

    def test_detail_pages_report_total_display_rows_not_invoice_count(self):
        backend, _ = self.backend_with_job()

        class DetailReader:
            @staticmethod
            def detail_page(_job, *, limit, cursor):
                start = int(cursor or 0)
                rows = [
                    {"id": index + 1, "stt": index + 1, "ten": f"Dòng {index + 1}"}
                    for index in range(start, min(start + limit, 73))
                ]
                next_offset = start + len(rows)
                return {
                    "items": rows,
                    "invoice_count": 12,
                    "row_count": len(rows),
                    "pagination": {
                        "limit": limit,
                        "has_more": next_offset < 73,
                        "next_cursor": str(next_offset) if next_offset < 73 else None,
                    },
                }

        with patch(
            "app.external_api.results.JobResultReader", return_value=DetailReader()
        ):
            query = {
                "connection_id": "conn_account_1",
                "date_from": "2026-03-01",
                "date_to": "2026-03-31",
                "direction": "purchase",
                "query_type": "query",
                "limit": 50,
            }
            first = backend.results("details", query)
            second = backend.results(
                "details", {**query, "cursor": first["pagination"]["next_cursor"]}
            )

        self.assertEqual(len(first["items"]), 50)
        self.assertEqual(first["total_count"], 73)
        self.assertEqual(len(second["items"]), 23)
        self.assertEqual(second["total_count"], 73)

    def test_result_dispatch_initializes_source_backend_after_restart(self):
        previous = (
            mia_runtime.data_directory,
            mia_runtime.logger,
            mia_runtime.production_backend,
        )
        backend = Mock()
        backend.results.return_value = {
            "items": [{"overview_id": 1}],
            "pagination": {"limit": 50, "has_more": False, "next_cursor": None},
        }
        try:
            with tempfile.TemporaryDirectory() as directory, patch(
                "mia_runtime.ProductionBackend", return_value=backend
            ) as constructor:
                mia_runtime.data_directory = Path(directory)
                mia_runtime.logger = None
                mia_runtime.production_backend = None
                query = {
                    "connection_id": "conn_account_1",
                    "date_from": "2026-02-01",
                    "date_to": "2026-02-28",
                }
                result, should_stop = mia_runtime.dispatch("results.overview", query)
                self.assertFalse(should_stop)
                self.assertEqual(result["items"][0]["overview_id"], 1)
                constructor.assert_called_once_with(Path(directory), None)
                backend.results.assert_called_once_with("overview", query)
        finally:
            (
                mia_runtime.data_directory,
                mia_runtime.logger,
                mia_runtime.production_backend,
            ) = previous

    def test_reconciliation_dispatch_uses_read_only_source_backend_adapter(self):
        previous = (
            mia_runtime.data_directory,
            mia_runtime.logger,
            mia_runtime.production_backend,
        )
        backend = Mock()
        backend.reconciliation.return_value = {
            "items": [],
            "reconciliation": {"issue_count": 0},
            "pagination": {"limit": 50, "has_more": False, "next_cursor": None},
        }
        try:
            with tempfile.TemporaryDirectory() as directory, patch(
                "mia_runtime.ProductionBackend", return_value=backend
            ):
                mia_runtime.data_directory = Path(directory)
                mia_runtime.logger = None
                mia_runtime.production_backend = None
                query = {"connection_id": "conn_account_1", "limit": 50}
                result, should_stop = mia_runtime.dispatch("results.reconciliation", query)
                self.assertFalse(should_stop)
                self.assertEqual(result["reconciliation"]["issue_count"], 0)
                backend.reconciliation.assert_called_once_with(query)
        finally:
            (
                mia_runtime.data_directory,
                mia_runtime.logger,
                mia_runtime.production_backend,
            ) = previous

    def test_result_export_groups_source_sheets_by_direction_and_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            backend, _ = self.backend_with_job(data_root=root / "source-data")
            overview_row = {
                "khmshdon": "1",
                "khhdon": "AA/26E",
                "shdon": "1",
                "tdlap": "2026-02-01T00:00:00+07:00",
            }
            detail_record = {"raw_detail_path": "unused.json"}

            def write_overview(_rows, **kwargs):
                kwargs["target"].write_bytes(b"source-overview")

            def combine(staged_jobs, target):
                self.assertEqual(len(staged_jobs), 1)
                target.write_bytes(b"combined-source-workbook")

            detail_repository = Mock()
            detail_repository.get_detail_records_for_export.return_value = [detail_record]
            detail_exporter = Mock()
            detail_exporter.export.side_effect = (
                lambda **kwargs: Path(kwargs["output_path"]).write_bytes(b"source-detail")
            )

            with patch(
                "mia_source_results._all_overview_fields",
                return_value=[overview_row],
            ) as overview_rows, patch(
                "mia_source_results._write_overview_excel_from_source_template",
                side_effect=write_overview,
            ) as overview_writer, patch(
                "app.repositories.invoice_detail_query_repository.InvoiceDetailQueryRepository",
                return_value=detail_repository,
            ), patch(
                "app.exporters.invoice_detail_excel_exporter.InvoiceDetailExcelExporter",
                return_value=detail_exporter,
            ), patch(
                "mia_source_results._combine_source_workbooks_atomically",
                side_effect=combine,
            ), patch(
                "mia_source_results._normalize_detail_excel_atomically",
                return_value=None,
            ):
                result = backend.export_results({
                    "destination": directory,
                    "connection_ids": ["conn_account_1"],
                    "result_scopes": ["overview", "details"],
                    "date_from": "2026-02-01",
                    "date_to": "2026-02-28",
                    "direction": "purchase",
                    "query_type": "query",
                    "search": "",
                })

            self.assertEqual(result["count"], 2)
            self.assertEqual(len(result["files"]), 2)
            self.assertTrue(all(Path(path).is_file() for path in result["files"]))
            self.assertEqual(
                {Path(path).name for path in result["files"]},
                {
                    "0100000000 - Mua vào - Tổng quan - 2026-02-01_2026-02-28.xlsx",
                    "0100000000 - Mua vào - Chi tiết - 2026-02-01_2026-02-28.xlsx",
                },
            )
            self.assertEqual(
                {Path(path).parent.relative_to(Path(directory)) for path in result["files"]},
                {
                    Path("0100000000") / "Tổng quan 2026-02-01_2026-02-28",
                    Path("0100000000") / "Chi tiết 2026-02-01_2026-02-28",
                },
            )
            overview_rows.assert_called_once()
            overview_writer.assert_called_once()
            detail_repository.get_detail_records_for_export.assert_called_once()
            detail_exporter.export.assert_called_once()

    def test_combined_export_creates_one_sheet_per_direction_and_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            backend, _ = self.backend_with_job(data_root=root / "source-data")

            def write_marker(path: Path, marker: str) -> None:
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.cell(5, 1).value = "STT"
                worksheet.cell(5, 2).value = "Marker"
                worksheet.cell(6, 1).value = 1
                worksheet.cell(6, 2).value = marker
                workbook.save(path)
                workbook.close()

            def overview_rows(_context, direction, query_type, _search):
                return [{
                    "marker": f"overview:{direction}:{query_type}",
                    "nbmst": f"tax-{query_type}", "khmshdon": "1",
                    "khhdon": "AA", "shdon": "1", "nlap_date": "2026-08-02",
                }]

            def overview_writer(rows, **kwargs):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.cell(5, 1).value = "STT"
                worksheet.cell(5, 2).value = "Marker"
                for index, row in enumerate(rows, start=6):
                    worksheet.cell(index, 1).value = index - 5
                    worksheet.cell(index, 2).value = row["marker"]
                workbook.save(kwargs["target"])
                workbook.close()

            detail_repository = Mock()
            detail_repository.get_detail_records_for_export.side_effect = (
                lambda _tax_code, direction, query_type, _from, _to: [
                    {
                        "marker": f"details:{direction}:{query_type}",
                        "nbmst": f"tax-{query_type}", "khmshdon": "1",
                        "khhdon": "AA", "shdon": "1", "nlap_date": "2026-08-02",
                    }
                ]
            )

            class MarkerDetailExporter:
                def __init__(self, *_args, **_kwargs):
                    pass

                def export(self, *, detail_records, output_path, **_kwargs):
                    workbook = Workbook()
                    worksheet = workbook.active
                    worksheet.cell(5, 1).value = "STT"
                    worksheet.cell(5, 2).value = "Marker"
                    for index, record in enumerate(detail_records, start=6):
                        worksheet.cell(index, 1).value = index - 5
                        worksheet.cell(index, 2).value = record["marker"]
                    workbook.save(output_path)
                    workbook.close()

            with patch(
                "mia_source_results._all_overview_fields",
                side_effect=overview_rows,
            ), patch(
                "mia_source_results._write_overview_excel_from_source_template",
                side_effect=overview_writer,
            ), patch(
                "app.repositories.invoice_detail_query_repository.InvoiceDetailQueryRepository",
                return_value=detail_repository,
            ), patch(
                "app.exporters.invoice_detail_excel_exporter.InvoiceDetailExcelExporter",
                MarkerDetailExporter,
            ), patch(
                "mia_source_results._normalize_detail_excel_atomically",
                return_value=None,
            ):
                result = backend.export_results({
                    "destination": directory,
                    "connection_ids": ["conn_account_1"],
                    "result_scopes": ["overview", "details"],
                    "date_from": "2026-08-01",
                    "date_to": "2026-08-22",
                    "direction": None,
                    "query_type": None,
                    "query_types": ["query", "sco-query"],
                    "search": "",
                })

            expected = {
                f"0100000000 - {direction_label} - {scope_label} - 2026-08-01_2026-08-22.xlsx"
                for direction_label in ("Mua vào", "Bán ra")
                for scope_label in ("Tổng quan", "Chi tiết")
            }
            self.assertEqual(result["count"], 4)
            self.assertEqual({Path(path).name for path in result["files"]}, expected)

            for path_value in result["files"]:
                path = Path(path_value)
                scope = "overview" if "Tổng quan" in path.name else "details"
                direction = "purchase" if "Mua vào" in path.name else "sold"
                workbook = load_workbook(path, data_only=False)
                try:
                    self.assertEqual(len(workbook.sheetnames), 1)
                    sheet = workbook.active
                    self.assertEqual(sheet.cell(5, 2).value, "Marker")
                    self.assertEqual(
                        {sheet.cell(6, 2).value, sheet.cell(7, 2).value},
                        {
                            f"{scope}:{direction}:query",
                            f"{scope}:{direction}:sco-query",
                        },
                    )
                finally:
                    workbook.close()

    def test_real_source_excel_export_uses_selected_result_range(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data_root = root / "source-data"
            backend, _ = self.backend_with_job(data_root=data_root)
            database_path = data_root / "0100000000" / "db" / "invoices.sqlite3"
            overview_repository = InvoiceOverviewRepository(database_path)
            detail_repository = InvoiceDetailRepository(database_path)
            timestamp = "2026-08-22T03:00:00+00:00"
            raw_overview = root / "overview.json"
            raw_overview.write_text("{}", encoding="utf-8")

            def overview_item(day: str, shdon: str) -> dict:
                public = {
                    "khmshdon": "1",
                    "khhdon": "K25TAN",
                    "shdon": shdon,
                    "tdlap": f"{day}T08:00:00+07:00",
                    "nbmst": "0300555450-001",
                    "nbten": "CHI NHÁNH XĂNG DẦU\x0b SÀI GÒN",
                    "nmmst": "0315394414",
                    "nmten": "CÔNG TY TNHH HỒNG TRÀ NGỌC",
                    "nmdchi": "TP Hồ Chí Minh",
                    "tgtcthue": 100,
                    "tgtthue": 8,
                    "ttcktmai": 0,
                    "tgtphi": 0,
                    "tgtttbso": 108,
                    "dvtte": "VND",
                    "tgia": 1,
                    "tthai": 1,
                    "kqcht": "Hợp lệ",
                }
                return {
                    "nbmst": public["nbmst"],
                    "khhdon": public["khhdon"],
                    "shdon": shdon,
                    "khmshdon": public["khmshdon"],
                    "nlap": public["tdlap"],
                    "nlap_date": day,
                    "_public_fields": public,
                }

            overview_repository.upsert_items(
                company_tax_code="0100000000",
                direction="purchase",
                query_type="query",
                invoice_category="electronic",
                raw_json_path=raw_overview,
                items=[
                    overview_item("2025-05-01", "101"),
                    overview_item("2025-05-03", "303"),
                ],
                timestamp=timestamp,
            )

            def detail_payload(day: str, shdon: str) -> dict:
                return {
                    "detail": {
                        "khmshdon": 1,
                        "khhdon": "K25TAN",
                        "shdon": shdon,
                        "tdlap": f"{day}T08:00:00+07:00",
                        "nky": f"{day}T08:05:00+07:00",
                        "mhdon": f"MCCQT-{shdon}",
                        "dvtte": "VND",
                        "tgia": 1,
                        "nbten": "CHI NHÁNH XĂNG DẦU\x0b SÀI GÒN",
                        "nbmst": "0300555450-001",
                        "nbdchi": "TP Hồ Chí Minh",
                        "nmten": "CÔNG TY TNHH HỒNG TRÀ NGỌC",
                        "nmmst": "0315394414",
                        "nmdchi": "TP Hồ Chí Minh",
                        "ttcktmai": 0,
                        "tgtphi": 0,
                        "tgtttbso": 108,
                        "tgtthue": 8,
                        "tthai": 1,
                        "ttxly": 5,
                        "thtttoan": "CK",
                        "ttkhac": [],
                        "urltracuu": "https" + "://invoice.example/lookup?code=ABC123",
                        "cttkhac": [{"ttruong": "Mã tra cứu", "dlieu": "ABC123"}],
                        "hdhhdvu": [{
                            "ten": "Dịch vụ\x0b kiểm thử",
                            "dvtinh": "Lần",
                            "sluong": 1,
                            "dgia": 100,
                            "stckhau": 0,
                            "tsuat": 8,
                            "thtien": 100,
                            "tthue": 8,
                            "tchat": 1,
                            "ttkhac": [],
                        }],
                    }
                }

            for day, shdon in (("2025-05-01", "101"), ("2025-05-03", "303")):
                raw_detail = root / f"detail-{shdon}.json"
                raw_detail.write_text(
                    json.dumps(detail_payload(day, shdon), ensure_ascii=False),
                    encoding="utf-8",
                )
                detail_repository.upsert_detail_success(
                    company_tax_code="0100000000",
                    direction="purchase",
                    query_type="query",
                    invoice_category="electronic",
                    nbmst="0300555450-001",
                    khhdon="K25TAN",
                    shdon=shdon,
                    khmshdon="1",
                    nlap=f"{day}T08:00:00+07:00",
                    nlap_date=day,
                    raw_detail_path=raw_detail,
                    http_status=200,
                    fetched_at=timestamp,
                )

            reconciliation = backend.reconciliation({
                "connection_id": "conn_account_1",
                "date_from": "2025-05-01",
                "date_to": "2025-05-02",
                "direction": "purchase",
                "query_type": "query",
                "limit": 50,
            })
            self.assertEqual(
                reconciliation["reconciliation"]["overview_invoice_count"], 1
            )
            self.assertEqual(
                reconciliation["reconciliation"]["detail_invoice_count"], 1
            )
            self.assertEqual(reconciliation["reconciliation"]["issue_count"], 0)

            export_dir = root / "export"
            result = backend.export_results({
                "destination": str(export_dir),
                "connection_ids": ["conn_account_1"],
                "result_scopes": ["overview", "details"],
                "date_from": "2025-05-01",
                "date_to": "2025-05-02",
                "direction": "purchase",
                "query_type": "query",
                "search": "",
            })

            self.assertEqual(result["count"], 2)
            overview_path = next(
                Path(path) for path in result["files"]
                if "Tổng quan" in Path(path).name
            )
            detail_path = next(
                Path(path) for path in result["files"]
                if "Chi tiết" in Path(path).name
            )

            overview_workbook = load_workbook(overview_path, data_only=False)
            try:
                self.assertEqual(overview_workbook.sheetnames, ["Hóa đơn điện tử"])
                worksheet = overview_workbook["Hóa đơn điện tử"]
                header_row = _find_header_row(worksheet)
                self.assertEqual(
                    worksheet.cell(4, 1).value,
                    "Từ ngày 01/05/2025 đến ngày 02/05/2025",
                )
                rows = [
                    row for row in worksheet.iter_rows(
                        min_row=header_row + 1,
                        max_row=worksheet.max_row,
                        values_only=True,
                    )
                    if any(value is not None for value in row)
                ]
                self.assertEqual(len(rows), 1)
                self.assertEqual(str(rows[0][3]), "101")
                self.assertNotIn("\x0b", str(rows[0][6]))
            finally:
                overview_workbook.close()

            detail_workbook = load_workbook(detail_path, data_only=False)
            try:
                self.assertEqual(detail_workbook.sheetnames, ["Sheet1"])
                worksheet = detail_workbook["Sheet1"]
                header_row = 1
                headers = [
                    str(worksheet.cell(header_row, column).value or "")
                    for column in range(1, worksheet.max_column + 1)
                ]
                self.assertEqual(headers, [label for _key, label in _DETAIL_EXPORT_COLUMNS])
                self.assertEqual(len(headers), 37)
                self.assertNotIn("STT", headers)
                shdon_column = headers.index("Số hóa đơn") + 1
                url_column = 30
                code_column = 31
                values = [
                    worksheet.cell(row, shdon_column).value
                    for row in range(header_row + 1, worksheet.max_row + 1)
                    if worksheet.cell(row, shdon_column).value is not None
                ]
                self.assertEqual([str(value) for value in values], ["101"])
                self.assertNotIn("\x0b", str(worksheet.cell(header_row + 1, 18).value))
                url_cell = worksheet.cell(header_row + 1, url_column)
                self.assertEqual(url_cell.value, "https" + "://invoice.example/lookup?code=ABC123")
                self.assertIsNotNone(url_cell.hyperlink)
                self.assertEqual(url_cell.hyperlink.target, url_cell.value)
                self.assertEqual(worksheet.cell(header_row + 1, code_column).value, "ABC123")
            finally:
                detail_workbook.close()

    def test_real_combined_overview_export_has_both_sources_in_one_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data_root = root / "source-data"
            backend, _ = self.backend_with_job(data_root=data_root)
            database_path = data_root / "0100000000" / "db" / "invoices.sqlite3"
            repository = InvoiceOverviewRepository(database_path)
            timestamp = "2026-08-22T03:00:00+00:00"
            raw = root / "overview.json"
            raw.write_text("{}", encoding="utf-8")

            for query_type, shdon, partner in (
                ("query", "000101", "0300000001"),
                ("sco-query", "000202", "0300000002"),
            ):
                public = {
                    "khmshdon": "1", "khhdon": "K26T", "shdon": shdon,
                    "tdlap": "2026-01-10T08:00:00+07:00", "nbmst": partner,
                    "nbten": f"Nguồn {query_type}", "nmmst": "0100000000",
                    "nmten": "Công ty kiểm thử", "tgtcthue": 100,
                    "tgtthue": 10, "tgtttbso": 110, "tthai": 1,
                }
                repository.upsert_items(
                    company_tax_code="0100000000", direction="purchase",
                    query_type=query_type,
                    invoice_category="cash_register" if query_type == "sco-query" else "electronic",
                    raw_json_path=raw,
                    items=[{
                        "nbmst": partner, "khhdon": "K26T", "shdon": shdon,
                        "khmshdon": "1", "nlap": public["tdlap"],
                        "nlap_date": "2026-01-10", "_public_fields": public,
                    }],
                    timestamp=timestamp,
                )

            result = backend.export_results({
                "destination": str(root / "export"),
                "connection_ids": ["conn_account_1"], "result_scopes": ["overview"],
                "date_from": "2026-01-01", "date_to": "2026-01-31",
                "direction": "purchase", "query_type": None,
                "query_types": ["query", "sco-query"], "search": "",
            })
            self.assertEqual(result["count"], 1)
            output = Path(result["files"][0])
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(len(workbook.sheetnames), 1)
                worksheet = workbook.active
                header_row = _find_header_row(worksheet)
                headers = [worksheet.cell(header_row, column).value for column in range(1, worksheet.max_column + 1)]
                invoice_column = headers.index("Số hóa đơn") + 1
                values = {
                    str(worksheet.cell(row, invoice_column).value)
                    for row in range(header_row + 1, worksheet.max_row + 1)
                    if worksheet.cell(row, invoice_column).value is not None
                }
                self.assertEqual(values, {"000101", "000202"})
            finally:
                workbook.close()
            fixture_path = os.environ.get("MIA_COMBINED_RESULT_FIXTURE")
            if fixture_path:
                Path(fixture_path).parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(output, fixture_path)

    def test_excel_export_keeps_scope_filters_separate_and_forwards_exclusion(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            backend, _ = self.backend_with_job(data_root=root / "source-data")
            detail_repository = Mock()
            detail_repository.get_detail_records_for_export.return_value = [{
                "direction": "purchase", "query_type": "query", "nbmst": "0101",
                "khhdon": "AA", "shdon": "1", "khmshdon": "1",
                "raw_detail_path": str(root / "detail.json"),
            }]
            (root / "detail.json").write_text("{}", encoding="utf-8")
            row_builder = Mock()
            row_builder.build_rows.return_value = [{"ten": "Dịch vụ"}]
            detail_exporter = Mock()
            detail_exporter.export.side_effect = lambda **kwargs: Path(kwargs["output_path"]).write_bytes(b"detail")

            def combine(_staged, target):
                target.write_bytes(b"combined")

            with patch(
                "mia_source_results._resolve_excluded_keys",
                return_value=frozenset({"purchase|query|0101|AA|999|1"}),
            ) as resolve_exclusion, patch(
                "mia_source_results._all_overview_fields",
                return_value=[{"shdon": "1"}],
            ) as overview_rows, patch(
                "mia_source_results._write_overview_excel_from_source_template",
                side_effect=lambda _rows, **kwargs: kwargs["target"].write_bytes(b"overview"),
            ), patch(
                "app.repositories.invoice_detail_query_repository.InvoiceDetailQueryRepository",
                return_value=detail_repository,
            ), patch(
                "mia_source_results._FilteredExcelSafeDetailRowBuilder",
                return_value=row_builder,
            ) as filtered_builder, patch(
                "app.exporters.invoice_detail_excel_exporter.InvoiceDetailExcelExporter",
                return_value=detail_exporter,
            ), patch(
                "mia_source_results._combine_source_workbooks_atomically",
                side_effect=combine,
            ), patch(
                "mia_source_results._normalize_detail_excel_atomically",
                return_value=None,
            ):
                result = backend.export_results({
                    "destination": directory,
                    "connection_ids": ["conn_account_1"],
                    "result_scopes": ["overview", "details"],
                    "date_from": "2026-02-01", "date_to": "2026-02-28",
                    "direction": "purchase", "query_type": "query", "search": "",
                    "result_filters": {
                        "overview": {"search": "overview-term", "column_filters": {"nbten": {"values": ["A"]}}},
                        "details": {"search": "detail-term", "column_filters": {"ten": {"values": ["Dịch vụ"]}}},
                    },
                    "exclusion": {"keys": ["purchase|query|0101|AA|999|1"], "rules": []},
                })

            self.assertEqual(result["count"], 2)
            self.assertEqual(overview_rows.call_args.args[3], "overview-term")
            self.assertEqual(overview_rows.call_args.args[4], {"nbten": {"values": ["A"]}})
            filtered_builder.assert_any_call(
                search="detail-term", column_filters={"ten": {"values": ["Dịch vụ"]}}
            )
            resolve_exclusion.assert_called_once()


if __name__ == "__main__":
    unittest.main()
