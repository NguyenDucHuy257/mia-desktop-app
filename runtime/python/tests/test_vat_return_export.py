import json
import hashlib
import sqlite3
import tempfile
import unittest
import zipfile
import re
import os
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from app.repositories.invoice_detail_repository import InvoiceDetailRepository
from app.repositories.invoice_overview_repository import InvoiceOverviewRepository
from app.parsers.invoice_detail_excel_row_builder import _to_number, resolve_tax_amount
from mia_vat_return_export import (MONEY_CELLS, MONEY_NUMBER_FORMAT, OBLIGATION_FORMULAS,
                                   PURCHASE_SHEET_NAME, REDUCTION_SHEET_NAME, SOLD_SHEET_NAME, TEMPLATE_NAME, aggregate_vat_return,
                                   build_vat_return_data, export_vat_return, vat_return_filename,
                                   _tax_group, _template_path)


class VatReturnExportTests(unittest.TestCase):
    tax_code = "0101234567"

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name); self.db = self.root / self.tax_code / "db" / "invoices.sqlite3"
        InvoiceOverviewRepository(self.db).init_db(); InvoiceDetailRepository(self.db).init_db()
        self.counter = 0

    def invoice(self, direction, rate, base, tax, *, query_type="query", status=1,
                lines=None, save_detail=True, number=None, khhdon="K", nbmst="PARTNER",
                overview_fields=None, invoice_date="2023-10-15", khmshdon="1"):
        self.counter += 1
        number = str(number or self.counter)
        fields = {"tgtcthue": str(base), "tgtthue": str(tax), "tthai": status}
        fields.update(overview_fields or {})
        with sqlite3.connect(self.db) as c:
            c.execute("""INSERT INTO invoice_overview_items(company_tax_code,direction,query_type,invoice_category,nbmst,khhdon,shdon,khmshdon,nlap,nlap_date,raw_json_path,detail_fetched,detail_path,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (self.tax_code,direction,query_type,"cash_register" if query_type=="sco-query" else "electronic",nbmst,khhdon,"S"+number,khmshdon,invoice_date,invoice_date,"raw",1 if save_detail else 0,"detail" if save_detail else "","now","now"))
            invoice_id=c.execute("select last_insert_rowid()").fetchone()[0]
            c.executemany("insert into invoice_overview_attributes(invoice_item_id,field_name,value_json) values(?,?,?)",[(invoice_id,k,json.dumps(v)) for k,v in fields.items()])
        if not save_detail:
            return
        detail_lines = lines if lines is not None else [(rate, base, tax)]
        normalized_lines=[]
        for index,line in enumerate(detail_lines,1):
            if isinstance(line,dict):
                normalized_lines.append(line)
            elif len(line)==4:
                name,line_rate,line_base,line_tax=line
                normalized_lines.append({"ten":name,"tsuat":line_rate,"thtien":str(line_base),"tthue":str(line_tax)})
            else:
                line_rate,line_base,line_tax=line
                normalized_lines.append({"ten":f"Mặt hàng {index}","tsuat":line_rate,
                                         "thtien":str(line_base),"tthue":str(line_tax)})
        InvoiceDetailRepository(self.db).replace_normalized_detail_success(company_tax_code=self.tax_code,direction=direction,query_type=query_type,invoice_category="cash_register" if query_type=="sco-query" else "electronic",nbmst=nbmst,khhdon=khhdon,shdon="S"+number,khmshdon=khmshdon,nlap=invoice_date,nlap_date=invoice_date,raw_detail_path="detail",http_status=200,fetched_at="now",lines=normalized_lines)

    def fixture(self):
        self.invoice("purchase","0%",100,0); self.invoice("purchase","5%",200,10)
        self.invoice("purchase","10%",300,30); self.invoice("purchase","8%",800,64)
        self.invoice("purchase","10%",50,5,status=6); self.invoice("purchase","10%",-20,-2,status=2)
        self.invoice("sold","KCT",30,0); self.invoice("sold","0.00",40,0)
        self.invoice("sold","5%",150,7.5,lines=[("5%",100,5),("0.05",50,2.5)])
        self.invoice("sold","10%",-100,-10,status=3); self.invoice("sold","KKKNT",20,0)
        self.invoice("sold","8%",900,72); self.invoice("sold","10%",70,7,status=4)
        self.invoice("sold","10%",200,20,query_type="sco-query")
        self.invoice("sold","8%",130,9,lines=[("8%",80,6.4),("5%",50,2.5)])

    def ready_coverage(self):
        direction={"ready":True,"missing":[]}
        return {"accounts":[{"connection_id":"conn", "purchase":direction,"sold":direction}]}

    def export_book(self, destination_name="out"):
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,
                                _load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            return export_vat_return(backend,{"connection_ids":["conn"],
                "destination":str(self.root/destination_name),"date_from":"2023-10-01",
                "date_to":"2023-10-31"})

    @staticmethod
    def raw_cells(workbook_path):
        from xml.etree import ElementTree as ET
        namespace={"m":"http:" "//schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(workbook_path) as archive:
            sheet=ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            styles=ET.fromstring(archive.read("xl/styles.xml"))
        formats={int(item.get("numFmtId")):item.get("formatCode") for item in styles.find("m:numFmts",namespace) or []}
        xfs=list(styles.find("m:cellXfs",namespace))
        output={}
        for cell in sheet.findall(".//m:c",namespace):
            value=cell.find("m:v",namespace); formula=cell.find("m:f",namespace)
            style=xfs[int(cell.get("s","0"))]
            output[cell.get("r")]={"value":None if value is None else value.text,"formula":None if formula is None else formula.text,
                                   "format":formats.get(int(style.get("numFmtId","0")),style.get("numFmtId","0")),"type":cell.get("t")}
        return output

    @staticmethod
    def sheet_signature(sheet):
        def style(cell):
            color=lambda value:(value.type,value.rgb,value.indexed,value.theme) if value else None
            side=lambda value:(value.style,color(value.color))
            return (cell.font.name,cell.font.sz,cell.font.bold,cell.font.italic,color(cell.font.color),cell.fill.fill_type,color(cell.fill.fgColor),
                    side(cell.border.left),side(cell.border.right),side(cell.border.top),side(cell.border.bottom),
                    cell.alignment.horizontal,cell.alignment.vertical,cell.alignment.wrap_text,cell.number_format,cell.protection.locked)
        return {
            "title": sheet.title, "cells": [(cell.coordinate, cell.value, style(cell)) for row in sheet.iter_rows() for cell in row if cell.value is not None],
            "merged": [str(item) for item in sheet.merged_cells.ranges],
            "columns": {key: (value.width, value.hidden) for key, value in sheet.column_dimensions.items()},
            "rows": {key: (value.height, value.hidden) for key, value in sheet.row_dimensions.items()},
            "print_area": str(sheet.print_area), "freeze": str(sheet.freeze_panes),
            "orientation": sheet.page_setup.orientation, "paper": sheet.page_setup.paperSize,
        }

    def test_aggregation_includes_8_percent_and_classifies_mixed_invoice_by_line(self):
        self.fixture(); totals=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(totals["purchase_base"],1380); self.assertEqual(totals["purchase_tax"],102)
        self.assertEqual((totals["kct"],totals["0"],totals["5_base"],totals["5_tax"]),(30,40,200,11))
        self.assertEqual((totals["10_base"],totals["10_tax"],totals["kkknt"]),(1080,88,20))
        self.assertEqual(totals["kct"]+totals["0"]+totals["5_base"]+totals["10_base"]+totals["kkknt"],1370)

    def test_sqlite_trace_totals_match_direct_excel_cells_for_8_percent(self):
        self.fixture()
        trace=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        destination=self.root/"reconciliation"
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,_load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            result=export_vat_return(backend,{"connection_ids":["conn"],"destination":str(destination),"date_from":"2023-10-01","date_to":"2023-10-31"})
        raw=self.raw_cells(result["files"][0])
        mapping={"F14":"purchase_base","H14":"purchase_tax","F19":"kct","F21":"0",
                 "F22":"5_base","H22":"5_tax","F23":"10_base","H23":"10_tax","F24":"kkknt"}
        for cell,total_name in mapping.items():
            self.assertEqual(Decimal(raw[cell]["value"]),trace[total_name],f"SQLite trace != Excel {cell}")
        self.assertEqual(Decimal(raw["F14"]["value"]),Decimal("1380"))
        self.assertEqual(Decimal(raw["H14"]["value"]),Decimal("102"))
        self.assertEqual(Decimal(raw["F23"]["value"]),Decimal("1080"))
        self.assertEqual(Decimal(raw["H23"]["value"]),Decimal("88"))

    def test_exports_vat_sheets_and_preserves_remaining_unimplemented_template_sheet(self):
        self.fixture(); destination=self.root/"out"
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,_load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        template_hash=hashlib.sha256(_template_path().read_bytes()).hexdigest()
        template=load_workbook(_template_path(),data_only=False)
        signatures=[self.sheet_signature(s) for s in template.worksheets]
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            result=export_vat_return(backend,{"connection_ids":["conn"],"destination":str(destination),"date_from":"2023-10-01","date_to":"2023-10-31"})
        self.assertEqual(result["count"],1); self.assertEqual(len(list(destination.iterdir())),1)
        self.assertEqual(Path(result["files"][0]).name,"To_khai_thue_GTGT_0101234567_01-10-2023_31-10-2023.xlsx")
        self.assertEqual(
            Path(result["files"][0]).parent.relative_to(destination),
            Path("0101234567") / "Mua vào & Bán ra",
        )
        book=load_workbook(result["files"][0],data_only=False); sheet=book.worksheets[0]
        with zipfile.ZipFile(_template_path()) as original_zip, zipfile.ZipFile(result["files"][0]) as output_zip:
            for entry in ("xl/worksheets/sheet5.xml",):
                self.assertEqual(output_zip.read(entry),original_zip.read(entry))
        self.assertEqual(book.sheetnames,template.sheetnames); self.assertEqual(sheet["A1"].value,"CÔNG TY KIỂM THỬ"); self.assertEqual(sheet["A2"].value,"Mã số thuế: 0101234567")
        self.assertIsNone(sheet["F8"].value); self.assertIsNone(sheet["H8"].value); self.assertIsNone(sheet["H11"].value); self.assertEqual((sheet["F14"].value,sheet["H14"].value),(1380,102)); self.assertEqual((sheet["F16"].value,sheet["H16"].value),(0,0))
        expected={"H17":"=H14","F20":"=F21+F22+F23+F24","H20":"=H22+H23","F25":"=F19+F20","H25":"=H20","H26":"=H25-H17","H32":"=MAX(H26-H11+H28-H29-H30,0)","H34":"=H32-H33","H35":"=MAX(-(H26-H11+H28-H29-H30),0)","H38":"=H35-H36"}
        for cell,formula in expected.items(): self.assertEqual(sheet[cell].value,formula)
        for cell in ("H28","H29","H30","H33","H36"): self.assertIsNone(sheet[cell].value)
        after=[self.sheet_signature(s) for s in book.worksheets]
        self.assertEqual(after[4],signatures[4])
        for key in ("title","merged","columns","rows","print_area","freeze","orientation","paper"): self.assertEqual(after[0][key],signatures[0][key])
        formulas=[cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value,str) and cell.value.startswith("=")]
        self.assertFalse(any(error in formula for formula in formulas for error in ("#REF!","#VALUE!","#NAME?")))
        self.assertTrue(book.calculation.fullCalcOnLoad); self.assertTrue(book.calculation.forceFullCalc); self.assertEqual(book.calculation.calcMode,"auto")
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            repeated=export_vat_return(backend,{"connection_ids":["conn"],"destination":str(destination),"date_from":"2023-10-01","date_to":"2023-10-31"})
        self.assertEqual(repeated["files"],result["files"]); self.assertEqual(len(list(destination.iterdir())),1)
        self.assertEqual(hashlib.sha256(_template_path().read_bytes()).hexdigest(),template_hash)

    def test_obligation_formulas_cover_positive_negative_zero_and_reopen_as_formulas(self):
        expected_formulas={
            "H32":"=MAX(H26-H11+H28-H29-H30,0)",
            "H35":"=MAX(-(H26-H11+H28-H29-H30),0)",
            "H38":"=H35-H36",
        }
        self.assertEqual(OBLIGATION_FORMULAS,expected_formulas)
        for s,expected_40a,expected_41 in (
            (Decimal("100"),Decimal("100"),Decimal("0")),
            (Decimal("-100"),Decimal("0"),Decimal("100")),
            (Decimal("0"),Decimal("0"),Decimal("0")),
            (Decimal("116376919"),Decimal("116376919"),Decimal("0")),
        ):
            actual_40a=max(s,Decimal(0)); actual_41=max(-s,Decimal(0))
            self.assertEqual((actual_40a,actual_41),(expected_40a,expected_41))
            self.assertFalse(actual_40a > 0 and actual_41 > 0)
            self.assertFalse(actual_40a.is_signed() and actual_40a == 0)
            self.assertFalse(actual_41.is_signed() and actual_41 == 0)
        self.assertEqual(Decimal("100")-Decimal(0),Decimal("100"))
        self.assertEqual(Decimal("100")-Decimal("30"),Decimal("70"))

        self.invoice("purchase","10%",100,100)
        result=self.export_book("obligation-formulas")
        raw=self.raw_cells(result["files"][0])
        for address,formula in expected_formulas.items():
            self.assertEqual("="+str(raw[address]["formula"]),formula)
            self.assertEqual(raw[address]["format"],"#,##0")
            self.assertIsNone(raw[address]["type"])
        reopened=load_workbook(result["files"][0],data_only=False,read_only=True)
        sheet=reopened.worksheets[0]
        self.assertEqual(sheet["H35"].value,expected_formulas["H35"])
        self.assertEqual(sheet["H38"].value,expected_formulas["H38"])
        self.assertIsNone(sheet["H36"].value)
        self.assertTrue(reopened.calculation.fullCalcOnLoad)
        self.assertTrue(reopened.calculation.forceFullCalc)
        self.assertEqual(reopened.calculation.calcMode,"auto")
        reopened.close()
        cached=load_workbook(result["files"][0],data_only=True,read_only=True)
        self.assertEqual(Decimal(str(cached.worksheets[0]["H35"].value)),Decimal("100"))
        self.assertEqual(Decimal(str(cached.worksheets[0]["H38"].value)),Decimal("100"))
        self.assertIsNone(cached.worksheets[0]["H36"].value)
        cached.close()

    def test_positive_obligation_caches_visible_zero_for_41_and_43(self):
        self.invoice("sold","10%",1000,100)
        result=self.export_book("positive-obligation")
        formula_book=load_workbook(result["files"][0],data_only=False,read_only=True)
        value_book=load_workbook(result["files"][0],data_only=True,read_only=True)
        formula_sheet=formula_book.worksheets[0]; value_sheet=value_book.worksheets[0]
        self.assertEqual(formula_sheet["H35"].value,"=MAX(-(H26-H11+H28-H29-H30),0)")
        self.assertEqual(formula_sheet["H38"].value,"=H35-H36")
        self.assertEqual(value_sheet["H35"].value,0)
        self.assertEqual(value_sheet["H38"].value,0)
        self.assertIsNone(value_sheet["H36"].value)
        self.assertEqual((formula_sheet["H35"].number_format,formula_sheet["H38"].number_format),
                         ("#,##0","#,##0"))
        formula_book.close(); value_book.close()
        raw=self.raw_cells(result["files"][0])
        self.assertEqual((raw["H35"]["value"],raw["H38"]["value"]),("0","0"))
        self.assertEqual((raw["H35"]["formula"],raw["H38"]["formula"]),
                         ("MAX(-(H26-H11+H28-H29-H30),0)","H35-H36"))
        with zipfile.ZipFile(result["files"][0]) as archive:
            namespace={"m":"http:" "//schemas.openxmlformats.org/spreadsheetml/2006/main"}
            workbook_xml=ET.fromstring(archive.read("xl/workbook.xml"))
            calculation=workbook_xml.find("m:calcPr",namespace)
            self.assertEqual(calculation.get("calcMode"),"auto")
            self.assertEqual(calculation.get("fullCalcOnLoad"),"1")
            self.assertEqual(calculation.get("forceFullCalc"),"1")
            self.assertNotIn("xl/calcChain.xml",archive.namelist())

    def test_locked_destination_is_preserved_and_does_not_create_a_suffix(self):
        self.invoice("sold","10%",1000,100)
        first=self.export_book("locked-target")
        original=Path(first["files"][0])
        original_bytes=original.read_bytes()
        real_replace=os.replace

        def replace_with_locked_target(source, destination):
            if Path(destination) == original:
                raise PermissionError(13,"destination is open",str(destination))
            return real_replace(source,destination)

        with patch("mia_vat_return_export.os.replace",side_effect=replace_with_locked_target):
            with self.assertRaisesRegex(ValueError,"vat_return_destination_file_locked"):
                self.export_book("locked-target")
        self.assertEqual(original.read_bytes(),original_bytes)
        files=list(original.parent.iterdir())
        self.assertEqual(files,[original])
        self.assertFalse(any(path.name.endswith("_2.xlsx") or path.name.endswith("_5.xlsx") for path in files))
        self.assertFalse(any(path.name.startswith(f".{original.stem}-") for path in files))

    def test_filename_uses_hyphenated_dates_and_no_collision_suffix(self):
        first=vat_return_filename("0109591907","2023-10-01","2023-10-31")
        self.assertEqual(
            first,
            "To_khai_thue_GTGT_0109591907_01-10-2023_31-10-2023.xlsx",
        )
        self.assertNotEqual(first,vat_return_filename("0101234567","2023-10-01","2023-10-31"))
        self.assertNotEqual(first,vat_return_filename("0109591907","2023-11-01","2023-11-30"))
        self.assertNotRegex(first,r"_\d+\.xlsx$")
        with self.assertRaisesRegex(ValueError,"invalid_vat_return_range"):
            vat_return_filename("0109591907","01-10-2023","2023-10-31")

    def test_second_export_atomically_overwrites_the_same_file_with_new_data(self):
        self.invoice("sold","10%",1000,100)
        first=self.export_book("overwrite")
        target=Path(first["files"][0]); before=target.read_bytes()
        self.invoice("sold","10%",2500,250,number="NEW")
        second=self.export_book("overwrite")
        self.assertEqual(second["files"],first["files"])
        self.assertNotEqual(target.read_bytes(),before)
        self.assertEqual([path for path in target.parent.iterdir() if path.suffix==".xlsx"],[target])
        self.assertFalse(any(path.name.startswith(f".{target.stem}-") for path in target.parent.iterdir()))
        raw=self.raw_cells(target)
        self.assertEqual(raw["F23"]["value"],"3500")
        self.assertEqual(raw["H23"]["value"],"350")

    def test_replace_permission_failure_is_not_mislabeled_as_an_open_file(self):
        self.invoice("sold","10%",1000,100)
        first=self.export_book("denied-target")
        target=Path(first["files"][0]); original_bytes=target.read_bytes()
        with patch("mia_vat_return_export.os.replace",side_effect=PermissionError(13,"denied",str(target))), \
             patch("mia_vat_return_export.os.access",return_value=False):
            with self.assertRaisesRegex(ValueError,"^vat_return_destination_not_writable$"):
                self.export_book("denied-target")
        self.assertEqual(target.read_bytes(),original_bytes)
        self.assertEqual(list(target.parent.iterdir()),[target])

    def test_unwritable_destination_cleans_up_without_reporting_success(self):
        self.invoice("sold","10%",1000,100)
        with patch("mia_vat_return_export.tempfile.mkstemp",side_effect=PermissionError(13,"denied")):
            with self.assertRaisesRegex(ValueError,"^vat_return_destination_not_writable$"):
                self.export_book("unwritable")
        destination=self.root/"unwritable"
        self.assertTrue(destination.is_dir())
        self.assertFalse(list(destination.rglob("*.xlsx")))

    def test_money_values_are_integer_vnd_with_thousands_format(self):
        self.invoice("purchase","10%","123456789","17543211")
        self.invoice("sold","KCT","123456789","0")
        self.invoice("sold","0%","17543211","0")
        self.invoice("sold","5%","10000025","500012")
        self.invoice("sold","10%","2345675","234568")
        self.invoice("sold","KKKNT","-1234568","0")
        totals=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        expected={
            "purchase_base":Decimal("123456789"),"purchase_tax":Decimal("17543211"),
            "kct":Decimal("123456789"),"0":Decimal("17543211"),
            "5_base":Decimal("10000025"),"5_tax":Decimal("500012"),
            "10_base":Decimal("2345675"),"10_tax":Decimal("234568"),
            "kkknt":Decimal("-1234568"),
        }
        self.assertEqual(totals,expected)
        destination=self.root/"out"
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,_load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            result=export_vat_return(backend,{"connection_ids":["conn"],"destination":str(destination),"date_from":"2023-10-01","date_to":"2023-10-31"})
        raw=self.raw_cells(result["files"][0])
        cell_values={"F14":"123456789","H14":"17543211","F19":"123456789","F21":"17543211",
                     "F22":"10000025","H22":"500012","F23":"2345675",
                     "H23":"234568","F24":"-1234568"}
        for address, expected_text in cell_values.items():
            self.assertIsNone(raw[address]["type"])
            self.assertEqual(Decimal(raw[address]["value"]),Decimal(expected_text))
            self.assertEqual(Decimal(raw[address]["value"]),Decimal(raw[address]["value"]).to_integral())
            self.assertEqual(raw[address]["format"],MONEY_NUMBER_FORMAT)
        for address in MONEY_CELLS:
            self.assertEqual(raw[address]["format"],MONEY_NUMBER_FORMAT)
            self.assertNotRegex(str(raw[address]["format"]),r"(?i)(million|triệu)")
            self.assertFalse(str(raw[address]["format"]).rstrip().endswith((",",",\"")))
        formulas={address:raw[address]["formula"] for address in ("H17","F20","H20","F25","H25","H26","H32","H34","H35","H38")}
        self.assertFalse(any(re.search(r"(?i)\b(ROUND|ROUNDUP|ROUNDDOWN|MROUND|INT|TRUNC)\s*\(",formula or "") for formula in formulas.values()))
        f20=expected["0"]+expected["5_base"]+expected["10_base"]+expected["kkknt"]
        h20=expected["5_tax"]+expected["10_tax"]
        self.assertEqual(f20,Decimal("28654343")); self.assertEqual(h20,Decimal("734580"))
        self.assertEqual(expected["kct"]+f20,Decimal("152111132"))
        h26=h20-expected["purchase_tax"]
        self.assertEqual(h26,Decimal("-16808631"))
        self.assertEqual(max(h26,Decimal(0)),Decimal(0))
        self.assertEqual(max(-h26,Decimal(0)),Decimal("16808631"))
        reopened=load_workbook(result["files"][0],data_only=False,read_only=True)
        self.assertEqual(Decimal(str(reopened.worksheets[0]["F22"].value)),Decimal("10000025"))
        self.assertEqual(reopened.worksheets[0]["F22"].number_format,MONEY_NUMBER_FORMAT)
        reopened.close()

    def test_detail_normalizer_and_repository_keep_decimal_exact(self):
        self.assertEqual(_to_number("123,456,789"),Decimal("123456789"))
        self.assertEqual(_to_number("10000000.25"),Decimal("10000000.25"))
        self.assertEqual(_to_number("-1234567.50"),Decimal("-1234567.50"))
        calculated=resolve_tax_amount({"thtien":"2345678.75","tsuat":"10"},{},0,1,Decimal(0))
        self.assertEqual(calculated,Decimal("234567.875"))
        self.invoice("sold","10%","2345678.75","234567.875",
                     lines=[("10%",Decimal("2345678.75"),Decimal("234567.875"))])
        with sqlite3.connect(self.db) as connection:
            stored=connection.execute("SELECT CAST(thtien AS TEXT),CAST(tthue AS TEXT) FROM invoice_detail_lines").fetchone()
        self.assertEqual(stored,("2345678.75","234567.875"))

    def test_decimal_parser_rejects_float_money_inside_vat_aggregation(self):
        from mia_vat_return_export import _decimal
        with self.assertRaisesRegex(ValueError,"vat_return_float_money:tgtcthue"):
            _decimal(123456789.25,"tgtcthue")

    def test_money_rounds_once_to_nearest_dong_with_half_up(self):
        from mia_vat_return_export import _decimal
        cases={
            "123456789.49":Decimal("123456789"),
            "123456789.50":Decimal("123456790"),
            "-1234567.49":Decimal("-1234567"),
            "-1234567.50":Decimal("-1234568"),
            "-0.49":Decimal("0"),
        }
        for source, expected in cases.items():
            actual=_decimal(source,"fixture")
            self.assertEqual(actual,expected)
            self.assertEqual(actual,actual.to_integral())
            self.assertFalse(actual.is_signed() and actual == 0)

    def test_fractional_sources_are_written_as_integer_cells(self):
        self.invoice("purchase","10%","123456789.49","123456789.50")
        self.invoice("sold","5%","123456789.50","-1234567.50")
        destination=self.root/"rounded"
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,_load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=self.ready_coverage()):
            result=export_vat_return(backend,{"connection_ids":["conn"],"destination":str(destination),"date_from":"2023-10-01","date_to":"2023-10-31"})
        raw=self.raw_cells(result["files"][0])
        expected={"F14":"123456789","H14":"123456790","F22":"123456790","H22":"-1234568"}
        for address, integer_text in expected.items():
            self.assertEqual(raw[address]["value"],integer_text)
            self.assertEqual(raw[address]["format"],"#,##0")
            self.assertNotIn(".",raw[address]["value"])

    def test_unknown_rate_stops_export(self):
        self.invoice("sold","KHAC",100,10)
        with self.assertRaisesRegex(ValueError,"vat_return_unknown_tax_rate:1"): aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")

    def test_missing_purchase_detail_now_blocks_full_workbook_for_reduction_sheet(self):
        self.invoice("purchase", "10%", 100, 10, save_detail=False)
        backend=SimpleNamespace(data_root=self.root,connection_tax_code=lambda _id:self.tax_code,
                                _load_company_names=lambda:{"conn":"CÔNG TY KIỂM THỬ"})
        blocked={"accounts":[{"connection_id":"conn","purchase":{"ready":False,"missing":[{"scope":"details","date_from":"2023-10-01","date_to":"2023-10-31"}]},"sold":{"ready":True,"missing":[]}}]}
        with patch("mia_vat_return_export.ArtifactInspector.vat_return_coverage",return_value=blocked):
            with self.assertRaisesRegex(ValueError,"vat_return_coverage_missing"):
                export_vat_return(backend,{"connection_ids":["conn"],"destination":str(self.root/"out"),"date_from":"2023-10-01","date_to":"2023-10-31"})

    def test_missing_valid_sold_detail_blocks_with_actionable_count(self):
        self.invoice("sold", "10%", 100, 10, save_detail=False)
        with self.assertRaisesRegex(ValueError, r'vat_return_detail_missing:.*"direction":"sold".*"count":1'):
            aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")

    def test_cancelled_sold_invoice_without_detail_does_not_block(self):
        self.invoice("sold", "10%", 100, 10, status=6, save_detail=False)
        totals=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(totals["10_base"],0)

    def test_cross_query_duplicate_uses_one_available_detail(self):
        self.invoice("sold","10%",100,10,query_type="query",number="77",save_detail=False)
        self.invoice("sold","10%",100,10,query_type="sco-query",number="77")
        totals=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(totals["10_base"],100)

    def test_sold_schedule_aggregates_invoice_plus_group_and_reconciles_tab_one(self):
        self.fixture()
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual([item["shdon"] for item in report["sold_groups"]["5"]],["S9","S15"])
        mixed_5=next(item for item in report["sold_groups"]["5"] if item["shdon"]=="S15")
        mixed_10=next(item for item in report["sold_groups"]["10"] if item["shdon"]=="S15")
        self.assertEqual((mixed_5["base"],mixed_5["tax"]),(Decimal("50"),Decimal("3")))
        self.assertEqual((mixed_10["base"],mixed_10["tax"]),(Decimal("80"),Decimal("6")))
        self.assertTrue(mixed_10["has_reduction"]); self.assertEqual(mixed_10["reduction"],2)
        result=self.export_book("schedule")
        book=load_workbook(result["files"][0],data_only=False); sheet=book[SOLD_SHEET_NAME]
        headers={str(sheet.cell(row,1).value).strip():row for row in range(1,sheet.max_row+1)
                 if str(sheet.cell(row,1).value or "").strip().startswith(("1.","2.","3.","4."))}
        ranges=[]
        for label,row in headers.items():
            total=next(r for r in range(row+1,sheet.max_row+1) if sheet.cell(r,1).value=="Tổng")
            ranges.append((label,row,total))
        expected_keys=("kct","0","5","10")
        for (label,start,total_row),key in zip(ranges,expected_keys):
            data_rows=list(range(start+1,total_row))
            self.assertEqual(sum(Decimal(str(sheet.cell(r,11).value or 0)) for r in data_rows),
                             sum((item["base"] for item in report["sold_groups"][key]),Decimal(0)))
            self.assertEqual(Decimal(str(sheet.cell(total_row,11).value)),
                             sum((item["base"] for item in report["sold_groups"][key]),Decimal(0)))
            self.assertEqual(sheet.cell(total_row,11).number_format,"#,##0")
        sequence=[sheet.cell(r,1).value for _,start,total in ranges for r in range(start+1,total)]
        self.assertEqual(sequence,list(range(1,len(sequence)+1)))
        self.assertEqual(sheet.print_area,f"'{SOLD_SHEET_NAME}'!$A$1:$P${sheet.max_row}")
        self.assertEqual(sheet.sheet_properties.pageSetUpPr.fitToPage,True)
        book.close()

    def test_sold_schedule_omits_zero_value_note_lines_but_keeps_real_zero_rate_sales(self):
        self.invoice(
            "sold", "10%", 100, 10, number="1",
            lines=[
                {"ten": "Dá»‹ch vá»¥ thá»±c", "tsuat": "10%", "thtien": "100", "tthue": "10"},
                {"ten": "Ghi chÃº há»£p Ä‘á»“ng", "tsuat": "0%", "thtien": "0", "tthue": "0"},
                {"ten": "Äiá»u chá»‰nh lÃ m trÃ²n", "tsuat": "0%", "thtien": "0", "tthue": "-0.280"},
            ],
        )
        self.invoice("sold", "0%", 50, 0, number="2")

        report = build_vat_return_data(
            self.db, self.tax_code, "2023-10-01", "2023-10-31",
        )

        self.assertEqual(len(report["sold_groups"]["0"]), 1)
        self.assertEqual(report["sold_groups"]["0"][0]["shdon"], "S2")
        self.assertEqual(report["sold_groups"]["0"][0]["base"], Decimal("50"))
        self.assertEqual(report["totals"]["0"], Decimal("50"))
        excluded = [
            item for item in report["trace"]
            if item["exclusion_reason"] == "zero_value_detail_line"
        ]
        self.assertEqual(len(excluded), 2)
        self.assertEqual(
            {(item["base"], item["actual_tax"]) for item in excluded},
            {("0", "0"), ("0", "-0.280")},
        )

        result = self.export_book("zero-note-lines")
        book = load_workbook(result["files"][0], data_only=False)
        sheet = book[SOLD_SHEET_NAME]
        exported_numbers = [sheet.cell(row, 5).value for row in range(1, sheet.max_row + 1)]
        self.assertNotIn("S1", [
            sheet.cell(row, 5).value for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 11).value == 0 and sheet.cell(row, 12).value == 0
        ])
        self.assertIn("S2", exported_numbers)
        book.close()

    def test_reduction_is_only_explicit_8_or_declared_10_matching_8_percent(self):
        self.invoice("sold","10%",1000,100)
        self.invoice("sold","8%",1000,80)
        self.invoice("sold","10%",1000,80)
        self.invoice("sold","10%",1000,92)
        self.invoice("sold","10%",1000,73)
        self.invoice("sold","5%",1000,40)
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        rows={item["shdon"]:item for item in report["sold_groups"]["10"]}
        self.assertFalse(rows["S1"]["has_reduction"])
        self.assertEqual((rows["S2"]["has_reduction"],rows["S2"]["reduction"]),(True,20))
        self.assertEqual((rows["S3"]["has_reduction"],rows["S3"]["reduction"]),(True,20))
        self.assertFalse(rows["S4"]["has_reduction"])
        self.assertFalse(rows["S5"]["has_reduction"])
        five=next(item for item in report["sold_groups"]["5"] if item["shdon"]=="S6")
        self.assertFalse(five["has_reduction"])
        anomalies={item["line_identity"]:item for item in report["reduction_anomalies"]}
        self.assertEqual(len(anomalies),3)
        self.assertEqual({item["actual_rate_percent"] for item in anomalies.values()},
                         {"9.2","7.3","4"})
        self.assertTrue(all(not item["reduced"] for item in report["trace"]
                            if item["reduction_anomaly"]))

    def test_declared_10_matching_8_percent_accepts_five_dong_tolerance(self):
        self.invoice("sold","10%",1000,85)
        self.invoice("sold","10%",1000,75)
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        rows=report["sold_groups"]["10"]
        self.assertTrue(all(item["has_reduction"] for item in rows))
        self.assertTrue(all(item["reduction"] == Decimal("20") for item in rows))
        self.assertFalse(report["reduction_anomalies"])

    def test_mixed_8_and_full_10_rounds_reduction_once_on_only_reduced_base(self):
        self.invoice("sold","10%",75,7,
                     lines=[("8%",25,2),("8%",25,2),("10%",25,3)])
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        item=report["sold_groups"]["10"][0]
        self.assertEqual((item["base"],item["tax"]),(Decimal("75"),Decimal("7")))
        self.assertTrue(item["has_reduction"])
        self.assertEqual(item["reduction"],Decimal("1"))
        reduced_lines=[line for line in report["trace"] if line["reduced"]]
        self.assertEqual(sum(Decimal(line["reduction_base"]) for line in reduced_lines),Decimal("50"))
        self.assertTrue(all(line["source_tax_rate"] == "8%" for line in reduced_lines))
        result=self.export_book("mixed-reduction")
        book=load_workbook(result["files"][0],data_only=False)
        row=next(r for r in range(7,book[SOLD_SHEET_NAME].max_row+1)
                 if book[SOLD_SHEET_NAME].cell(r,5).value=="S1")
        self.assertEqual((book[SOLD_SHEET_NAME].cell(row,13).value,
                          book[SOLD_SHEET_NAME].cell(row,14).value),("X",1))
        book.close()

    def test_reduction_uses_half_up_to_one_dong_and_export_returns_masked_audit(self):
        self.invoice("sold","8%",75,6)
        self.invoice("sold","10%",1000,92)
        result_report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        reduced=next(item for item in result_report["sold_groups"]["10"] if item["shdon"]=="S1")
        self.assertEqual(reduced["reduction"],Decimal("2"))
        result=self.export_book("audit")
        self.assertEqual(result["audit"]["reduction_anomaly_count"],1)
        anomaly=result["audit"]["reduction_anomalies"][0]
        self.assertNotIn(self.tax_code,anomaly["canonical_invoice_identity"])
        self.assertEqual(anomaly["actual_rate_percent"],"9.2")
        book=load_workbook(result["files"][0],data_only=False); sheet=book[SOLD_SHEET_NAME]
        rows={sheet.cell(r,5).value:r for r in range(7,sheet.max_row+1)}
        self.assertEqual((sheet.cell(rows["S1"],13).value,sheet.cell(rows["S1"],14).value),("X",2))
        self.assertIsNone(sheet.cell(rows["S2"],13).value)
        self.assertIsNone(sheet.cell(rows["S2"],14).value)
        book.close()

    def test_sold_sheet_identifiers_are_text_dates_are_dates_and_empty_reduction_is_none(self):
        self.invoice("sold","10%",100,10,number="00000002")
        result=self.export_book("typed")
        book=load_workbook(result["files"][0],data_only=False); sheet=book[SOLD_SHEET_NAME]
        data_row=next(row for row in range(7,sheet.max_row) if sheet.cell(row,5).value=="S00000002")
        self.assertEqual(sheet.cell(data_row,2).data_type,"s")
        self.assertEqual(sheet.cell(data_row,5).data_type,"s")
        self.assertEqual(sheet.cell(data_row,10).data_type,"s")
        self.assertEqual(sheet.cell(data_row,6).data_type,"d")
        self.assertIsNone(sheet.cell(data_row,13).value)
        self.assertIsNone(sheet.cell(data_row,14).value)
        book.close()

    def test_sold_sheet_maps_overview_buyer_fields_and_preserves_leading_zeroes(self):
        self.invoice("sold","10%",100,10,number="00000002",khhdon="01AB/26E",
                     overview_fields={"nmten":"NGƯỜI MUA A","nmmst":"0012345678",
                                      "ghichu":"Ghi chú nguồn","dgiai":"Diễn giải nguồn"})
        result=self.export_book("mapped")
        book=load_workbook(result["files"][0],data_only=False); sheet=book[SOLD_SHEET_NAME]
        row=next(r for r in range(7,sheet.max_row) if sheet.cell(r,5).value=="S00000002")
        self.assertEqual((sheet.cell(row,4).value,sheet.cell(row,7).value,
                          sheet.cell(row,10).value),("01AB/26E","NGƯỜI MUA A","0012345678"))
        self.assertEqual((sheet.cell(row,15).value,sheet.cell(row,16).value),
                         ("Ghi chú nguồn","Diễn giải nguồn"))
        book.close()

    def test_workbook_reduction_examples_round_half_up_to_dong(self):
        examples=(("67876000","5430080",Decimal("1357520")),
                  ("32159280","2572742",Decimal("643186")),
                  ("9275000","742000",Decimal("185500")),
                  ("7230605","578448",Decimal("144612")))
        for base,tax,_ in examples:
            self.invoice("sold","8%",base,tax)
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        for item,(_,_,expected) in zip(report["sold_groups"]["10"],examples):
            self.assertEqual(item["reduction"],expected)

    def test_kkknt_remains_on_tab_one_but_is_not_forced_into_missing_sheet_group(self):
        self.invoice("sold","KKKNT",123,0)
        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(report["totals"]["kkknt"],123)
        self.assertTrue(any(item["exclusion_reason"]=="sheet_has_no_kkknt_group"
                            for item in report["trace"]))
        self.assertTrue(all(not rows for rows in report["sold_groups"].values()))

    def test_identity_normalizes_case_whitespace_and_separators(self):
        self.invoice("sold","10%",100,10,number="88",khhdon=" ab-26/e ",
                     nbmst=" 0200 000 001 ",save_detail=False)
        InvoiceDetailRepository(self.db).replace_normalized_detail_success(
            company_tax_code=self.tax_code,direction="sold",query_type="query",
            invoice_category="electronic",nbmst="0200000001",khhdon="AB26E",
            shdon="s88",khmshdon="01",nlap="2023-10-15",nlap_date="2023-10-15",
            raw_detail_path="",http_status=200,fetched_at="now",
            lines=[{"tsuat":"10%","thtien":"100","tthue":"10"}],
        )
        totals=aggregate_vat_return(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(totals["10_tax"],10)

    def test_purchase_sheet_is_one_row_per_canonical_invoice_and_reconciles_tab_one(self):
        self.invoice("purchase","0%",100,0,number="000269512",nbmst="0012345678",
                     overview_fields={"nbten":"Công ty cổ phần MISA"})
        self.invoice("purchase","8%",200,16,number="2",khhdon="AA/26E",nbmst="0000000002",
                     overview_fields={"nbten":"Người bán 8%"})
        self.invoice("purchase","8%",200,16,number="2",khhdon="AA/26E",nbmst="0000000002",
                     query_type="sco-query",overview_fields={"nbten":"Người bán 8%"})
        self.invoice("purchase","10%",300,30,number="10",lines=[("5%",100,5),("10%",200,20)],
                     overview_fields={"nbten":"Người bán nhiều thuế suất"})
        self.invoice("purchase","10%",999,99,number="11",save_detail=False,
                     overview_fields={"tthai":"Hóa đơn đã bị hủy","nbten":"Không được xuất"})
        self.invoice("purchase","10%",50,5,number="12",
                     overview_fields={"tthai":"Hóa đơn thay thế","nbten":"Hóa đơn hợp lệ"})
        self.invoice("purchase","10%",-20,-2,number="13",
                     overview_fields={"tthai":"Hóa đơn điều chỉnh","nbten":"Điều chỉnh âm"})
        self.invoice("purchase","10%",500,50,number="14",invoice_date="2023-11-01",
                     save_detail=False,overview_fields={"nbten":"Ngoài kỳ"})

        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        self.assertEqual(len(report["purchase_items"]),5)
        self.assertEqual((report["totals"]["purchase_base"],report["totals"]["purchase_tax"]),
                         (Decimal("630"),Decimal("49")))
        duplicate=next(item for item in report["purchase_items"] if item["shdon"]=="S2")
        self.assertEqual(duplicate["query_type"],"query,sco-query")
        self.assertEqual(duplicate["deductible_tax"],duplicate["tax"])
        self.assertNotIn("S14",[item["shdon"] for item in report["purchase_items"]])
        self.assertTrue(any(not item["exported"] and item["reason"]=="duplicate_canonical"
                            for item in report["purchase_audit"]))
        self.assertTrue(any(not item["exported"] and item["reason"]=="excluded_status"
                            for item in report["purchase_audit"]))

        result=self.export_book("purchase-sheet")
        book=load_workbook(result["files"][0],data_only=False)
        sheet=book[PURCHASE_SHEET_NAME]; tab_one=book.worksheets[0]
        data_rows=[row for row in range(8,sheet.max_row+1)
                   if isinstance(sheet.cell(row,1).value,int)]
        self.assertEqual(len(data_rows),5)
        self.assertEqual([sheet.cell(row,1).value for row in data_rows],[1,2,3,4,5])
        first=next(row for row in data_rows if sheet.cell(row,4).value=="S000269512")
        self.assertEqual(sheet.cell(first,2).data_type,"s")
        self.assertEqual(sheet.cell(first,4).value,"S000269512")
        self.assertEqual(sheet.cell(first,8).value,"0012345678")
        self.assertEqual(sheet.cell(first,8).data_type,"s")
        self.assertEqual(sheet.cell(first,5).data_type,"d")
        self.assertEqual(sheet.cell(first,5).number_format,"dd/mm/yyyy")
        self.assertEqual(sheet.cell(first,13).value,
                         "Mua hàng của Công ty cổ phần MISA theo hóa đơn số S000269512")
        for row in data_rows:
            self.assertEqual(sheet.cell(row,10).value,sheet.cell(row,11).value)
            self.assertIsNone(sheet.cell(row,14).value)
            self.assertEqual((sheet.cell(row,9).number_format,sheet.cell(row,10).number_format,
                              sheet.cell(row,11).number_format),("#,##0","#,##0","#,##0"))
        total_row=next(row for row in range(8,sheet.max_row+1)
                       if sheet.cell(row,1).value=="Tổng")
        self.assertEqual((sheet.cell(total_row,9).value,sheet.cell(total_row,10).value,
                          sheet.cell(total_row,11).value),(630,49,49))
        self.assertEqual((tab_one["F14"].value,tab_one["H14"].value),(630,49))
        self.assertEqual(tab_one["H17"].value,"=H14")
        self.assertEqual(str(sheet.print_area),f"'{PURCHASE_SHEET_NAME}'!$A$1:$O${sheet.max_row}")
        template=load_workbook(_template_path(),read_only=False)
        template_sheet=template[PURCHASE_SHEET_NAME]
        self.assertEqual(sheet.column_dimensions["M"].width,
                         template_sheet.column_dimensions["M"].width)
        self.assertEqual(sheet.page_setup.orientation,
                         template_sheet.page_setup.orientation)
        self.assertEqual(sheet.cell(data_rows[0],1)._style,template_sheet["A8"]._style)
        self.assertEqual(sheet.cell(data_rows[-1],1)._style,template_sheet["A30"]._style)
        self.assertEqual(sheet.cell(total_row,1)._style,template_sheet["A31"]._style)
        self.assertEqual(sheet.row_dimensions[data_rows[0]].height,template_sheet.row_dimensions[8].height)
        for row in data_rows:
            self.assertIn(f"F{row}:G{row}",{str(item) for item in sheet.merged_cells.ranges})
        self.assertNotIn("CÔNG TY TNHH", " ".join(str(cell.value or "") for row in sheet for cell in row))
        template.close(); book.close()
        self.assertEqual(result["audit"]["purchase_invoice_count"],5)
        self.assertEqual(result["audit"]["purchase_base"],"630")
        self.assertEqual(result["audit"]["purchase_tax"],"49")

    def test_purchase_sheet_empty_period_keeps_three_groups_and_zero_totals(self):
        result=self.export_book("purchase-empty")
        book=load_workbook(result["files"][0],data_only=False)
        sheet=book[PURCHASE_SHEET_NAME]
        group_rows=[row for row in range(1,sheet.max_row+1)
                    if str(sheet.cell(row,1).value or "").strip().startswith(("1. ","2. ","3. "))]
        total_rows=[row for row in range(1,sheet.max_row+1) if sheet.cell(row,1).value=="Tổng"]
        self.assertEqual((len(group_rows),len(total_rows)),(3,3))
        for row in total_rows:
            self.assertEqual((sheet.cell(row,9).value,sheet.cell(row,10).value,
                              sheet.cell(row,11).value),(0,0,0))
            self.assertEqual(sheet.cell(row,9).number_format,"#,##0")
        self.assertFalse(any(isinstance(sheet.cell(row,1).value,int)
                             for row in range(7,sheet.max_row+1)))
        book.close()

    def test_purchase_missing_required_overview_money_blocks_export(self):
        self.invoice("purchase","10%",100,10,save_detail=False,
                     overview_fields={"tgtthue":None,"nbten":"Người bán"})
        with self.assertRaisesRegex(ValueError,"vat_return_purchase_invalid"):
            self.export_book("purchase-invalid")
        self.assertFalse((self.root/"purchase-invalid").exists())

    def test_purchase_null_overview_totals_use_only_balanced_complete_detail(self):
        self.invoice(
            "purchase", "", 5_700_000, 0, query_type="sco-query",
            overview_fields={
                "tgtcthue": None, "tgtthue": None,
                "tgtttbso": "5700000", "ttcktmai": 0, "tgtphi": None,
                "nbten": "Người bán máy tính tiền",
            },
            lines=[
                ("Kẹp tôn 5T", "", 3_700_000, 0),
                ("Kẹp tôn 6T", "", 2_000_000, 0),
            ],
        )
        report = build_vat_return_data(
            self.db, self.tax_code, "2023-10-01", "2023-10-31",
        )
        self.assertEqual(len(report["purchase_items"]), 1)
        item = report["purchase_items"][0]
        self.assertEqual((item["base"], item["tax"]),
                         (Decimal("5700000"), Decimal("0")))
        self.assertEqual(item["totals_source"], "detail_verified_fallback")
        self.assertEqual(report["purchase_audit"][0]["reason"],
                         "detail_verified_fallback")
        result = self.export_book("purchase-detail-fallback")
        self.assertEqual(result["audit"]["purchase_base"], "5700000")
        self.assertEqual(result["audit"]["purchase_tax"], "0")

    def test_purchase_null_overview_totals_reject_unbalanced_detail(self):
        self.invoice(
            "purchase", "", 5_700_000, 0, query_type="sco-query",
            overview_fields={
                "tgtcthue": None, "tgtthue": None,
                "tgtttbso": "6000000", "ttcktmai": 0, "tgtphi": None,
                "nbten": "Người bán máy tính tiền",
            },
            lines=[("Không cân với tổng thanh toán", "", 5_700_000, 0)],
        )
        with self.assertRaisesRegex(ValueError, "vat_return_purchase_invalid"):
            build_vat_return_data(
                self.db, self.tax_code, "2023-10-01", "2023-10-31",
            )

    def test_purchase_8_percent_reduction_sheet_uses_one_row_per_canonical_detail_line(self):
        mixed_lines=[
            ("Mặt hàng trùng", "8", 100, 8),
            ("Mặt hàng trùng", "8%", 100, 8),
            ("Mặt hàng 8.0", "8.0", 200, 16),
            ("Mặt hàng 0.08", "0.08", 300, 24),
            ("10% nhưng thuế bằng 8%", "10%", 100, 8),
            ("Mặt hàng 5%", "5%", 100, 5),
            ("Mặt hàng 0%", "0%", 100, 0),
            ("Mặt hàng KCT", "KCT", 100, 0),
            ("Thiếu thuế suất", None, 100, 8),
        ]
        self.invoice("purchase","8%",1100,77,number="1",khhdon="AA/26E",lines=mixed_lines)
        self.invoice("purchase","8%",1100,77,number="1",khhdon="AA/26E",
                     query_type="sco-query",lines=mixed_lines)
        self.invoice("purchase","8.00%",400,32,number="2",
                     lines=[("Mặt hàng 8.00%","8.00%",400,32)])
        self.invoice("purchase","8.0%",500,40,number="3",
                     lines=[("Mặt hàng 8.0%","8.0%",500,40)])
        self.invoice("purchase","8%",600,48,number="4",
                     overview_fields={"tthai":"Hóa đơn thay thế"},
                     lines=[("Hàng thay thế hợp lệ","8%",600,48)])
        self.invoice("purchase","8%",-100,-8,number="5",
                     overview_fields={"tthai":"Hóa đơn điều chỉnh"},
                     lines=[("Điều chỉnh âm","8%",-100,-8)])
        self.invoice("purchase","8%",100,1,number="6",
                     lines=[("8% chênh thuế bất thường","8%",100,1)])
        self.invoice("purchase","8%",700,56,number="7",status=6,
                     lines=[("Hóa đơn bị hủy","8%",700,56)])
        self.invoice("purchase","8%",800,64,number="8",
                     overview_fields={"tthai":"Hóa đơn đã bị thay thế"},
                     lines=[("Hóa đơn bị thay thế","8%",800,64)])
        self.invoice("purchase","8%",900,72,number="9",invoice_date="2023-11-01",
                     lines=[("Ngoài kỳ","8%",900,72)])
        with sqlite3.connect(self.db) as connection:
            connection.execute(
                """INSERT INTO invoice_detail_items(
                       company_tax_code,direction,query_type,invoice_category,nbmst,
                       khhdon,shdon,khmshdon,nlap,nlap_date,raw_detail_path,http_status,
                       fetched_at,created_at,updated_at,error_message,normalized_ready,
                       detail_outcome,normalized_line_count
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (self.tax_code,"purchase","query","electronic","ORPHAN","O","S999","1",
                 "2023-10-15","2023-10-15","detail",200,"now","now","now",None,1,"with_lines",1),
            )
            detail_id=connection.execute("SELECT last_insert_rowid()").fetchone()[0]
            connection.execute(
                """INSERT INTO invoice_detail_lines(
                       detail_item_id,line_number,ten,tsuat,thtien,tthue,created_at,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?)""",
                (detail_id,1,"Chi tiết mồ côi","8%","100","8","now","now"),
            )

        report=build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")
        lines=report["purchase_reduction_lines"]
        self.assertEqual(len(lines),9,[(item["name"],item["shdon"],item["line_number"]) for item in lines])
        self.assertEqual(sum((item["base"] for item in lines),Decimal(0)),Decimal("2200"))
        self.assertEqual(sum((item["tax"] for item in lines),Decimal(0)),Decimal("169"))
        self.assertEqual(sum(1 for item in lines if item["name"]=="Mặt hàng trùng"),2)
        self.assertEqual({item["line_number"] for item in lines if item["shdon"]=="S1"},
                         {1,2,3,4})
        self.assertTrue(all(item["parent_identity"] in {row["identity"] for row in report["purchase_items"]}
                            for item in lines))
        audit=report["purchase_reduction_audit"]
        looks_like=next(item for item in audit if item["name"]=="10% nhưng thuế bằng 8%"
                        and item["query_type"]=="query")
        self.assertFalse(looks_like["included"])
        self.assertEqual(looks_like["reason"],"tax_rate_not_8")
        self.assertTrue(looks_like["looks_like_8_percent"])
        self.assertTrue(any(item["reason"]=="duplicate_canonical_line" for item in audit))
        self.assertTrue(any(item["reason"]=="excluded_status" for item in audit))
        self.assertTrue(any(item["reason"]=="outside_period" for item in audit))
        self.assertTrue(any(item["reason"]=="orphan_detail" for item in audit))
        self.assertTrue(any(item["name"]=="8% chênh thuế bất thường"
                            and item["tax_difference_warning"] for item in audit))

        result=self.export_book("purchase-reduction")
        book=load_workbook(result["files"][0],data_only=False)
        sheet=book[REDUCTION_SHEET_NAME]
        data_rows=[row for row in range(8,sheet.max_row+1)
                   if isinstance(sheet.cell(row,1).value,int)]
        self.assertEqual(len(data_rows),9)
        self.assertEqual([sheet.cell(row,1).value for row in data_rows],list(range(1,10)))
        self.assertEqual(sum(1 for row in data_rows if sheet.cell(row,2).value=="Mặt hàng trùng"),2)
        total_row=next(row for row in range(8,sheet.max_row+1)
                       if str(sheet.cell(row,1).value or "").strip()=="Tổng")
        self.assertEqual((sheet.cell(total_row,6).value,sheet.cell(total_row,7).value),(2200,169))
        for row in data_rows:
            self.assertEqual((sheet.cell(row,6).number_format,sheet.cell(row,7).number_format),
                             ("#,##0","#,##0"))
            self.assertEqual(sheet.cell(row,6).value,int(sheet.cell(row,6).value))
            self.assertIn(f"B{row}:E{row}",{str(item) for item in sheet.merged_cells.ranges})
        part_two=next(row for row in range(1,sheet.max_row+1)
                      if str(sheet.cell(row,1).value or "").startswith("II. "))
        self.assertEqual(sheet.cell(part_two+1,1).value,"STT")
        self.assertEqual(sheet.cell(part_two+2,1).value,"(1)")
        self.assertEqual(str(sheet.cell(part_two+3,1).value).strip(),"Tổng")
        self.assertEqual((sheet.cell(part_two+3,6).value,sheet.cell(part_two+3,11).value),(0,0))
        self.assertTrue(str(sheet.cell(part_two+4,1).value).startswith("III. "))
        self.assertEqual(sheet.cell(part_two+5,3).value,
                         f"=K{part_two+3}-G{total_row}")
        self.assertEqual(sheet.max_row,24)
        self.assertEqual(str(sheet.print_area),f"'{REDUCTION_SHEET_NAME}'!$A$1:$L$24")
        template=load_workbook(_template_path(),data_only=False)
        template_sheet=template[REDUCTION_SHEET_NAME]
        self.assertEqual({key:sheet.column_dimensions[key].width for key in "ABCDEFGHIJKL"},
                         {key:template_sheet.column_dimensions[key].width for key in "ABCDEFGHIJKL"})
        self.assertEqual((sheet.page_setup.orientation,sheet.page_setup.paperSize,
                          sheet.page_setup.fitToHeight),
                         (template_sheet.page_setup.orientation,template_sheet.page_setup.paperSize,
                          template_sheet.page_setup.fitToHeight))
        self.assertEqual(sheet.cell(data_rows[0],1)._style,template_sheet["A8"]._style)
        self.assertEqual(sheet.cell(total_row,1)._style,template_sheet["A23"]._style)
        self.assertEqual(sheet.cell(part_two,1)._style,template_sheet["A24"]._style)
        template.close()
        sample_names=("Băng dính khổ 47mm x 6kg","Hộp carton 3 lớp B 1 nâu 180x100x80mm")
        all_text=" ".join(str(cell.value or "") for row in sheet for cell in row)
        self.assertTrue(all(name not in all_text for name in sample_names))
        book.close()
        self.assertEqual(result["audit"]["purchase_reduction_line_count"],9)
        self.assertEqual((result["audit"]["purchase_reduction_base"],
                          result["audit"]["purchase_reduction_tax"]),("2200","169"))

    def test_sold_reduction_sheet_lists_detail_lines_and_uses_template_formulas(self):
        self.invoice("purchase", "8%", 500, 40,
                     lines=[("Đầu vào 8%", "8%", 500, 40)])
        self.invoice("sold", "8%", 300, 24, number="20", lines=[
            ("Dịch vụ vận tải", "8%", 100, 8),
            ("Dịch vụ xuất bến", "8", 200, 16),
        ])
        self.invoice("sold", "10%", 400, 32, number="21", lines=[
            ("Khai thuế suất 10%, thực tế giảm còn 8%", "10%", 400, 32),
        ])
        self.invoice("sold", "10%", 500, 50, number="22", lines=[
            ("Không thuộc diện giảm", "10%", 500, 50),
        ])

        report = build_vat_return_data(
            self.db, self.tax_code, "2023-10-01", "2023-10-31",
        )
        self.assertEqual(
            [item["name"] for item in report["sold_reduction_lines"]],
            ["Dịch vụ vận tải", "Dịch vụ xuất bến",
             "Khai thuế suất 10%, thực tế giảm còn 8%"],
        )
        self.assertEqual(
            sum((item["base"] for item in report["sold_reduction_lines"]), Decimal(0)),
            Decimal("700"),
        )
        self.assertEqual(
            sum((item["reduction"] for item in report["sold_reduction_lines"]), Decimal(0)),
            Decimal("14"),
        )

        result = self.export_book("sold-reduction")
        formula_book = load_workbook(result["files"][0], data_only=False)
        value_book = load_workbook(result["files"][0], data_only=True)
        sheet = formula_book[REDUCTION_SHEET_NAME]
        values = value_book[REDUCTION_SHEET_NAME]
        part_two = next(
            row for row in range(1, sheet.max_row + 1)
            if str(sheet.cell(row, 1).value or "").startswith("II. ")
        )
        data_rows = [part_two + 3, part_two + 4, part_two + 5]
        self.assertEqual(
            [sheet.cell(row, 2).value for row in data_rows],
            ["Dịch vụ vận tải", "Dịch vụ xuất bến",
             "Khai thuế suất 10%, thực tế giảm còn 8%"],
        )
        for row, base, reduction in zip(data_rows, (100, 200, 400), (2, 4, 8)):
            self.assertEqual(sheet.cell(row, 6).value, base)
            self.assertEqual(sheet.cell(row, 7).value, 0.1)
            self.assertEqual(sheet.cell(row, 9).value, f"=G{row}*80%")
            self.assertEqual(sheet.cell(row, 11).value, f"=F{row}*(G{row}-I{row})")
            self.assertEqual(values.cell(row, 9).value, 0.08)
            self.assertEqual(values.cell(row, 11).value, reduction)
            self.assertIn(f"K{row}:L{row}", {str(item) for item in sheet.merged_cells.ranges})
        sold_total = data_rows[-1] + 1
        self.assertEqual((sheet.cell(sold_total, 6).value,
                          sheet.cell(sold_total, 11).value), (700, 14))
        difference_row = sold_total + 2
        self.assertEqual(sheet.cell(difference_row, 3).value,
                         f"=K{sold_total}-G{part_two-1}")
        self.assertEqual(values.cell(difference_row, 3).value, -26)
        self.assertEqual(result["audit"]["sold_reduction_line_count"], 3)
        self.assertEqual(result["audit"]["sold_reduction_base"], "700")
        self.assertEqual(result["audit"]["sold_reduction_tax"], "14")
        formula_book.close()
        value_book.close()

    def test_sold_reduction_sheet_groups_equal_display_columns_and_sums_revenue(self):
        self.invoice("sold", "8%", 100, 8, number="30", lines=[
            ("Dịch vụ vận tải", "8%", 100, 8),
        ])
        self.invoice("sold", "8%", 250, 20, number="31", lines=[
            ("  Dịch vụ   vận tải  ", "8", 250, 20),
        ])
        self.invoice("sold", "8%", 400, 32, number="32", lines=[
            ("Dịch vụ khác", "8%", 400, 32),
        ])

        report = build_vat_return_data(
            self.db, self.tax_code, "2023-10-01", "2023-10-31",
        )
        lines = report["sold_reduction_lines"]
        self.assertEqual(report["sold_reduction_source_line_count"], 3)
        self.assertEqual(len(lines), 2)
        transport = next(item for item in lines if item["name"] == "Dịch vụ vận tải")
        self.assertEqual(transport["base"], Decimal("350"))
        self.assertEqual(transport["statutory_rate"], Decimal("0.10"))
        self.assertEqual(transport["reduced_rate"], Decimal("0.08"))
        self.assertEqual(transport["reduction"], Decimal("7"))
        self.assertEqual(transport["source_line_count"], 2)

        result = self.export_book("sold-reduction-grouped")
        formula_book = load_workbook(result["files"][0], data_only=False)
        value_book = load_workbook(result["files"][0], data_only=True)
        sheet = formula_book[REDUCTION_SHEET_NAME]
        values = value_book[REDUCTION_SHEET_NAME]
        part_two = next(
            row for row in range(1, sheet.max_row + 1)
            if str(sheet.cell(row, 1).value or "").startswith("II. ")
        )
        data_rows = [part_two + 3, part_two + 4]
        self.assertEqual(
            [(sheet.cell(row, 2).value, sheet.cell(row, 6).value) for row in data_rows],
            [("Dịch vụ khác", 400), ("Dịch vụ vận tải", 350)],
        )
        transport_row = data_rows[1]
        self.assertEqual(sheet.cell(transport_row, 7).value, 0.1)
        self.assertEqual(sheet.cell(transport_row, 9).value,
                         f"=G{transport_row}*80%")
        self.assertEqual(sheet.cell(transport_row, 11).value,
                         f"=F{transport_row}*(G{transport_row}-I{transport_row})")
        self.assertEqual(values.cell(transport_row, 11).value, 7)
        self.assertEqual(result["audit"]["sold_reduction_source_line_count"], 3)
        self.assertEqual(result["audit"]["sold_reduction_line_count"], 2)
        formula_book.close()
        value_book.close()

    def test_purchase_reduction_tax_rate_normalizer_accepts_only_explicit_8_variants(self):
        for value in ("8","8%","8.0","8.00","8.0%","8.00%","0.08"):
            self.assertEqual(_tax_group(value),"8",value)
        for value in ("10","10%","0.10","5%","0%","KCT",None,"không xác định"):
            self.assertNotEqual(_tax_group(value),"8",value)

    def test_purchase_reduction_invalid_8_percent_name_or_tax_blocks_export(self):
        self.invoice("purchase","8%",200,16,lines=[
            {"ten":"","tsuat":"8%","thtien":"100","tthue":"8"},
            {"ten":"Thiếu tiền thuế","tsuat":"8%","thtien":"100","tthue":None},
        ])
        with self.assertRaisesRegex(ValueError,r'vat_return_purchase_reduction_invalid:.*"count":2'):
            build_vat_return_data(self.db,self.tax_code,"2023-10-01","2023-10-31")

    def test_purchase_reduction_empty_period_keeps_empty_part_two_frame(self):
        self.invoice("purchase","5%",100,5)
        result=self.export_book("purchase-reduction-empty")
        book=load_workbook(result["files"][0],data_only=False)
        sheet=book[REDUCTION_SHEET_NAME]
        totals=[row for row in range(1,sheet.max_row+1)
                if str(sheet.cell(row,1).value or "").strip()=="Tổng"]
        self.assertEqual(len(totals),2)
        self.assertEqual((sheet.cell(totals[0],6).value,sheet.cell(totals[0],7).value),(0,0))
        self.assertEqual((sheet.cell(totals[1],6).value,sheet.cell(totals[1],11).value),(0,0))
        self.assertEqual(sheet.max_row,15)
        self.assertFalse(any(isinstance(sheet.cell(row,1).value,int)
                             for row in range(8,sheet.max_row+1)))
        book.close()

    def test_template_is_packaged_resource(self):
        self.assertTrue(_template_path().is_file()); self.assertEqual(_template_path().name,TEMPLATE_NAME)
        spec=(Path(__file__).resolve().parents[2]/"mia-runtime.spec").read_text(encoding="utf-8")
        self.assertIn('str(vendor_root / "resources" / "templates")',spec)
        self.assertIn('hiddenimports=["mia_vat_return_export"]',spec)


if __name__ == "__main__": unittest.main()
