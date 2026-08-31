import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import app.exporters.invoice_detail_excel_exporter as source_module
from openpyxl import load_workbook

from mia_detail_excel_format import (
    DETAIL_EXPORT_COLUMNS,
    DETAIL_TEMPLATE_WIDTHS,
    FixedDetailExcelExporter,
    combine_detail_workbooks_atomically,
)
from mia_progressive_excel_exporter import ProgressiveInvoiceDetailExcelExporter
from mia_source_results import (
    _ExcelSafeDetailRowBuilder,
    _ExportProgressReporter,
    _source_template_dir,
    _write_overview_excel_from_source_template,
)


def detail_payload(number: str, products: int = 2) -> dict:
    return {
        "detail": {
            "khmshdon": "1", "khhdon": "K26T", "shdon": number,
            "tdlap": "2026-08-01T08:00:00+07:00", "dvtte": "VND",
            "nbten": "Người bán", "nbmst": "0100000000",
            "nmten": "Người mua", "nmmst": "0200000000",
            "tthai": 1, "ttxly": 5, "ttkhac": [],
            "hdhhdvu": [
                {
                    "ten": f"Hàng hóa {index}", "dvtinh": "Cái",
                    "sluong": 1, "dgia": 100, "tsuat": 8,
                    "thtien": 100, "tthue": 8, "tchat": 1,
                }
                for index in range(products)
            ],
        }
    }


class ExcelExportProgressTests(unittest.TestCase):
    def _records(self, root: Path) -> list[dict]:
        records = []
        for number in ("1", "2"):
            path = root / f"detail-{number}.json"
            path.write_text(json.dumps(detail_payload(number), ensure_ascii=False), encoding="utf-8")
            records.append({
                "raw_detail_path": str(path), "nbmst": "0100000000",
                "khhdon": "K26T", "shdon": number, "khmshdon": "1",
                "nlap": "2026-08-01T08:00:00+07:00",
                "material_codes_json": "[]",
            })
        return records

    def test_progressive_detail_export_normalizes_to_customer_template(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            records = self._records(root)
            first_payload = detail_payload("1")
            first_payload["detail"]["nbten"] = "N" * 200
            Path(records[0]["raw_detail_path"]).write_text(
                json.dumps(first_payload, ensure_ascii=False), encoding="utf-8"
            )
            staged_path = root / "staged.xlsx"
            final_path = root / "detail.xlsx"
            template = _source_template_dir() / "invoice_detail.xlsx"
            events = []
            width_calls = 0

            def measured_width(_value, **_kwargs):
                nonlocal width_calls
                width_calls += 1
                return 999

            def progress(phase, processed, total):
                if phase == "save" and processed == total == 1:
                    self.assertTrue(staged_path.is_file())
                events.append((phase, processed, total))

            with patch.object(source_module, "_display_width", side_effect=measured_width):
                ProgressiveInvoiceDetailExcelExporter(
                    template,
                    _ExcelSafeDetailRowBuilder(),
                    progress=progress,
                ).export(records, staged_path, "2026-08-01", "2026-08-31")
            combine_detail_workbooks_atomically(
                [("purchase", "electronic", staged_path)], final_path
            )

            workbook = load_workbook(final_path, data_only=False)
            try:
                worksheet = workbook.active
                expected_headers = [label for _key, label in DETAIL_EXPORT_COLUMNS]
                self.assertEqual(workbook.sheetnames, ["Sheet1"])
                self.assertEqual(worksheet.max_column, 37)
                self.assertEqual(
                    [worksheet.cell(1, column).value for column in range(1, 38)],
                    expected_headers,
                )
                self.assertNotIn("STT", expected_headers)
                self.assertEqual(worksheet["A2"].value, "1")
                self.assertEqual(worksheet["Q2"].value, "Hàng hóa 0")
                self.assertEqual(worksheet["J2"].value, "N" * 200)
                self.assertIsInstance(worksheet["S2"].value, (int, float))
                self.assertIsInstance(worksheet["W2"].value, (int, float))
                self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "FFFFFF00")
                self.assertEqual(worksheet["A1"].font.sz, 12)
                self.assertEqual(worksheet["A2"].font.sz, 12)
                self.assertEqual(
                    tuple(
                        getattr(worksheet["A1"].border, side).style
                        for side in ("left", "right", "top", "bottom")
                    ),
                    ("thin", "thin", "thin", "thin"),
                )
                self.assertFalse(worksheet.merged_cells.ranges)
                self.assertIsNone(worksheet.freeze_panes)
                self.assertIsNone(worksheet.auto_filter.ref)
                for column, expected_width in enumerate(DETAIL_TEMPLATE_WIDTHS, start=1):
                    letter = worksheet.cell(1, column).column_letter
                    self.assertAlmostEqual(
                        worksheet.column_dimensions[letter].width,
                        expected_width,
                        places=4,
                    )
                for row in range(1, worksheet.max_row + 1):
                    self.assertIsNone(worksheet.row_dimensions[row].height)
            finally:
                workbook.close()

            self.assertEqual(width_calls, 0)
            self.assertIn(("build_rows", 2, 2), events)
            self.assertIn(("write_rows", 4, 4), events)
            self.assertIn(("format", 38, 38), events)
            self.assertIn(("save", 1, 1), events)

    def test_detail_combiner_appends_categories_under_one_header(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            records = self._records(root)
            template = _source_template_dir() / "invoice_detail.xlsx"
            electronic = root / "electronic.xlsx"
            cash = root / "cash.xlsx"
            combined = root / "combined.xlsx"
            for target in (electronic, cash):
                FixedDetailExcelExporter(
                    template, _ExcelSafeDetailRowBuilder()
                ).export(records, target, "2026-08-01", "2026-08-31")

            combine_detail_workbooks_atomically(
                [
                    ("purchase", "electronic", electronic),
                    ("purchase", "cash_register", cash),
                ],
                combined,
            )

            combined_book = load_workbook(combined, data_only=False)
            try:
                worksheet = combined_book.active
                self.assertEqual(combined_book.sheetnames, ["Sheet1"])
                self.assertEqual(worksheet.max_row, 9)
                self.assertEqual(worksheet.max_column, 37)
                self.assertEqual(worksheet["A1"].value, "Mẫu số HD")
                self.assertEqual(worksheet["AK1"].value, "Hạn dùng")
                self.assertEqual(
                    sum(worksheet.cell(row, 3).value == "1" for row in range(2, 10)),
                    4,
                )
            finally:
                combined_book.close()

    def test_reporter_events_are_monotonic_and_terminal(self):
        events = []
        reporter = _ExportProgressReporter(events.append)
        reporter.planning(1, 2)
        reporter.planning(2, 2)
        reporter.set_units(2)
        reporter.start_unit("overview")
        reporter.unit("write_rows", 2, 10)
        reporter.unit("save", 1, 1)
        reporter.complete_unit()
        reporter.start_unit("details")
        reporter.unit("build_rows", 1, 2)
        reporter.unit("format", 38, 38)
        reporter.unit("save", 1, 1)
        reporter.complete_unit()
        reporter.complete()
        percents = [event["percent"] for event in events]
        self.assertEqual(percents[-1], 100)
        self.assertTrue(all(0 <= value <= 100 for value in percents))
        self.assertEqual(percents, sorted(percents))
        self.assertEqual(events[-1]["status"], "completed")

        failed = []
        failure_reporter = _ExportProgressReporter(failed.append)
        failure_reporter.planning(1, 2)
        failure_reporter.failed()
        self.assertEqual(failed[-1]["status"], "failed")
        self.assertLess(failed[-1]["percent"], 100)

    def test_overview_emits_actual_rows_and_finishes_after_atomic_replace(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "overview.xlsx"
            events = []

            def progress(phase, processed, total):
                if phase == "save" and processed == total == 1:
                    self.assertTrue(output.is_file())
                events.append((phase, processed, total))

            _write_overview_excel_from_source_template(
                [{
                    "khmshdon": "1", "khhdon": "K26T", "shdon": "1",
                    "tdlap": "2026-08-01T08:00:00+07:00", "tthai": 1,
                }],
                direction="purchase",
                category="electronic",
                date_from="2026-08-01",
                date_to="2026-08-31",
                target=output,
                progress=progress,
            )
            self.assertIn(("load_template", 1, 1), events)
            self.assertIn(("write_rows", 1, 1), events)
            self.assertEqual(events[-1], ("save", 1, 1))


if __name__ == "__main__":
    unittest.main()
