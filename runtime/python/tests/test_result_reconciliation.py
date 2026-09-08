import unittest
import tempfile
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from app.job_engine.models import JobRecord

from mia_backend import ProductionBackend
from mia_source_results import _RECONCILIATION_CACHE, _compare_money_values


SCHEMA = (
    ("khmshdon", "Ký hiệu mẫu số"), ("khhdon", "Ký hiệu hóa đơn"),
    ("shdon", "Số hóa đơn"), ("tdlap", "Ngày lập"),
    ("ntao", "Ngày tạo"), ("nbmst", "MST người bán"),
    ("nbten", "Tên người bán"), ("nmmst", "MST người mua"),
    ("nmten", "Tên người mua"), ("tgtcthue", "Tiền trước thuế"),
    ("tgtthue", "Tiền thuế"), ("thtien", "Thành tiền"),
    ("tthue", "Tiền thuế dòng"), ("ttcktmai", "Chiết khấu"),
    ("tgtphi", "Phí"), ("tgtttbso", "Tổng thanh toán"),
)


class FakeResultReader:
    def __init__(self, overview, details):
        self.overview = overview
        self.details = details

    @staticmethod
    def _page(rows, limit, cursor):
        start = int(cursor or 0)
        visible = rows[start:start + limit]
        next_offset = start + len(visible)
        return {
            "items": visible,
            "pagination": {
                "limit": limit,
                "has_more": next_offset < len(rows),
                "next_cursor": str(next_offset) if next_offset < len(rows) else None,
            },
        }

    def overview_page(self, _job, *, limit, cursor):
        return self._page(self.overview, limit, cursor)

    def detail_page(self, _job, *, limit, cursor):
        return self._page(self.details, limit, cursor)


class ResultReconciliationTests(unittest.TestCase):
    def setUp(self):
        _RECONCILIATION_CACHE.clear()

    @staticmethod
    def job(
        *, status="completed", coverage_from="2026-08-01",
        overview_to="2026-08-31", detail_to="2026-08-31",
    ):
        def module(to_date):
            return {
                "status": "completed",
                "months": [{
                    "from_date": coverage_from, "to_date": to_date,
                    "status": "completed",
                }],
            }
        return JobRecord(
            job_id="job-reconciliation", account_key="conn_account_1",
            company_tax_code="0100000000", job_type="invoice_crawl", queue_order=1,
            parameters={
                "connection_id": "conn_account_1", "date_from": "2026-01-01",
                "date_to": "2026-12-31", "directions": ["purchase"],
                "query_types": ["query"],
            },
            status=status, current_stage="finalize", worker_id=None,
            lease_token=None, lease_generation=0, lease_expires_at=None,
            available_at="2026-08-20T10:00:00+00:00", cancel_requested_at=None,
            warning_count=0, created_at="2026-08-20T10:00:00+00:00",
            updated_at="2026-08-20T10:00:00+00:00",
            started_at="2026-08-20T10:00:00+00:00",
            finished_at="2026-08-20T10:05:00+00:00",
            last_error_code=None, last_error_message=None,
            owner_id="mia-desktop-local", pipeline_version=2,
            progress_state={
                "version": 3,
                "modules": {
                    "overview": module(overview_to), "detail": module(detail_to),
                },
            },
        )

    def backend(self, jobs=None):
        backend = object.__new__(ProductionBackend)
        backend.data_root = Path("source-data")
        job = self.job()
        backend.repository = SimpleNamespace(
            list_jobs_for_reconciliation=lambda: list(jobs or [job])
        )
        return backend

    @staticmethod
    def overview(shdon, *, total=550):
        return {
            "id": f"o-{shdon}", "nbmst": " 0101000000 ", "khmshdon": 1,
            "khhdon": " AA/26E ", "shdon": shdon, "tdlap": "2026-08-01",
            "nbten": "Bên bán", "nmmst": "0202000000", "nmten": "Bên mua",
            "tgtcthue": 500, "tgtthue": 50, "ttcktmai": 5,
            "tgtphi": 2, "tgtttbso": total,
        }

    @staticmethod
    def detail_lines(shdon, count=5, *, total=550):
        return [{
            "id": f"d-{shdon}-{line}", "nbmst": "0101000000", "khmshdon": "1",
            "khhdon": "AA/26E", "shdon": str(shdon), "ntao": "2026-08-01",
            "nbten": "Bên bán", "nmmst": "0202000000", "nmten": "Bên mua",
            "thtien": 100, "tthue": 10, "ttcktmai": 5,
            "tgtphi": 2, "tgtttbso": total,
        } for line in range(count)]

    def reconcile(self, overview, details, *, schema=SCHEMA, jobs=None, **overrides):
        reader = FakeResultReader(overview, details)
        query = {
            "connection_id": "conn_account_1", "date_from": "2026-08-01",
            "date_to": "2026-08-31", "direction": "purchase",
            "query_type": "query", "limit": 50, **overrides,
        }
        with patch("app.external_api.results.JobResultReader", return_value=reader), patch(
            "mia_source_results._result_schema", return_value=schema
        ):
            return self.backend(jobs).reconciliation(query)

    def test_detail_lines_are_deduplicated_to_unique_invoices(self):
        overview = [self.overview(index) for index in range(1, 11)]
        details = [line for index in range(1, 11) for line in self.detail_lines(index)]
        result = self.reconcile(overview, details)
        self.assertEqual(result["reconciliation"]["overview_invoice_count"], 10)
        self.assertEqual(result["reconciliation"]["detail_invoice_count"], 10)
        self.assertEqual(result["reconciliation"]["issue_count"], 0)
        self.assertEqual(result["items"], [])

    def test_missing_both_directions_and_equal_counts_with_different_keys(self):
        result = self.reconcile(
            [self.overview(1), self.overview(2)],
            self.detail_lines(1) + self.detail_lines(3),
        )
        statuses = {item["fields"]["reconciliation_status"] for item in result["items"]}
        self.assertEqual(statuses, {"Thiếu chi tiết", "Thiếu tổng quan"})
        self.assertEqual(result["reconciliation"]["difference"], 0)
        self.assertEqual(result["reconciliation"]["issue_count"], 2)

    def test_money_mismatch_uses_line_sums_and_one_invoice_level_total(self):
        result = self.reconcile([self.overview(1, total=551)], self.detail_lines(1))
        self.assertEqual(result["reconciliation"]["money_mismatch_count"], 1)
        fields = result["items"][0]["fields"]
        self.assertEqual(
            fields["reconciliation_status"],
            "Chênh lệch tiền",
        )
        self.assertEqual(fields["mismatch_fields"], "Tổng thanh toán")
        self.assertEqual(fields["detail_thtien"], 500)
        self.assertEqual(fields["detail_ttcktmai"], 5)
        self.assertEqual(fields["detail_tgtttbso"], 550)
        self.assertEqual(fields["difference_tgtttbso"], 1)
        self.assertEqual(result["aggregate"]["totals"]["difference_tgtthue"], 0)

    def test_signed_zero_is_canonical_and_never_a_money_mismatch(self):
        overview = self.overview(1)
        overview["tgtcthue"] = Decimal("-0.00")
        overview["tgtthue"] = Decimal("+0.00")
        details = self.detail_lines(1, count=1, total=550)
        details[0]["thtien"] = Decimal("0.00")
        details[0]["tthue"] = Decimal("-0.00")
        result = self.reconcile([overview], details)
        self.assertEqual(result["reconciliation"]["money_mismatch_count"], 0)
        self.assertEqual(result["items"], [])

    def test_money_comparison_quantizes_vnd_before_deciding_mismatch(self):
        equal_pairs = (
            (20154017, 20154017),
            (9612466, Decimal("9612466.0000000002")),
            ("15.190.623", 15190623),
            (Decimal("5385547.00"), 5385547),
            (483008372, Decimal("483008371.9999999998")),
        )
        for overview, detail in equal_pairs:
            with self.subTest(overview=overview, detail=detail):
                comparison = _compare_money_values(overview, detail)
                self.assertEqual(comparison["difference"], Decimal("0"))
                self.assertFalse(comparison["difference"].is_signed())
                self.assertFalse(comparison["is_mismatch"])

        mismatch = _compare_money_values(6500000, 0)
        self.assertEqual(mismatch["difference"], Decimal("6500000"))
        self.assertTrue(mismatch["is_mismatch"])

    def test_microscopic_detail_residue_does_not_create_an_invoice_issue(self):
        overview = self.overview(1)
        details = self.detail_lines(1, count=1)
        overview.update({"tgtcthue": 100, "tgtthue": 10})
        details[0].update({
            "thtien": Decimal("100.0000000002"),
            "tthue": Decimal("9.9999999998"),
            "ttcktmai": Decimal("5.0000000002"),
            "tgtphi": Decimal("1.9999999998"),
            "tgtttbso": Decimal("550.0000000002"),
        })
        result = self.reconcile([overview], details)
        self.assertEqual(result["reconciliation"]["money_mismatch_count"], 0)
        self.assertEqual(result["reconciliation"]["issue_count"], 0)
        self.assertEqual(result["items"], [])

    def test_canonical_detail_decimal_text_is_not_scaled_as_thousands(self):
        overview = self.overview(1)
        overview.update({"tgtcthue": 1000000, "tgtthue": 10760000})
        details = self.detail_lines(1, count=2)
        details[0].update({"thtien": "400000.000", "tthue": "76800.000"})
        details[1].update({"thtien": "600000.000", "tthue": "10683200.000"})
        result = self.reconcile([overview], details)
        self.assertEqual(result["reconciliation"]["money_mismatch_count"], 0)
        self.assertEqual(result["items"], [])

        _RECONCILIATION_CACHE.clear()
        mismatch = self.reconcile(
            [{**overview, "tgtthue": 10760001}], details,
        )
        fields = mismatch["items"][0]["fields"]
        self.assertEqual(fields["detail_tthue"], 10760000)
        self.assertEqual(fields["difference_tgtthue"], 1)

    def test_all_cursor_pages_are_reconciled_before_ui_pagination(self):
        overview = [self.overview(index) for index in range(1, 206)]
        first = self.reconcile(overview, [])
        second = self.reconcile(overview, [], cursor=first["pagination"]["next_cursor"])
        self.assertEqual(first["total_count"], 205)
        self.assertEqual(first["reconciliation"]["issue_count"], 205)
        self.assertEqual(len(first["items"]), 50)
        self.assertEqual(len(second["items"]), 50)
        self.assertEqual(second["items"][0]["fields"]["stt"], 51)

    def test_search_filters_issue_page_without_changing_scope_summary(self):
        first = self.overview(1)
        first["nbten"] = "Alpha"
        second = self.overview(2)
        second["nbten"] = "Beta"
        result = self.reconcile([first, second], [], search="alpha")
        self.assertEqual(result["total_count"], 1)
        self.assertEqual(result["reconciliation"]["issue_count"], 2)

    def test_money_field_omitted_by_overview_template_is_not_a_false_mismatch(self):
        schema_without_fee = tuple(item for item in SCHEMA if item[0] != "tgtphi")
        result = self.reconcile(
            [self.overview(1)], self.detail_lines(1), schema=schema_without_fee
        )
        self.assertEqual(result["reconciliation"]["issue_count"], 0)

    def test_coverage_intersection_excludes_dates_not_completed_by_detail(self):
        overview = [self.overview(1), {**self.overview(2), "tdlap": "2026-08-20"}]
        details = self.detail_lines(1)
        result = self.reconcile(
            overview, details,
            jobs=[self.job(overview_to="2026-08-31", detail_to="2026-08-15")],
        )
        self.assertEqual(result["reconciliation"]["overview_invoice_count"], 1)
        self.assertEqual(result["reconciliation"]["issue_count"], 0)
        self.assertEqual(result["reconciliation"]["coverage_ranges"], [{
            "date_from": "2026-08-01", "date_to": "2026-08-15",
        }])

    def test_july_is_not_missing_when_detail_coverage_stops_in_june(self):
        january = {**self.overview(1), "tdlap": "2026-01-10"}
        july = {**self.overview(2), "tdlap": "2026-07-10"}
        january_details = [
            {**line, "ntao": "2026-01-10"} for line in self.detail_lines(1)
        ]
        result = self.reconcile(
            [january, july], january_details,
            jobs=[self.job(
                coverage_from="2026-01-01",
                overview_to="2026-07-31", detail_to="2026-06-30",
            )],
            date_from="2026-01-01", date_to="2026-07-31",
        )
        self.assertEqual(result["reconciliation"]["overview_invoice_count"], 1)
        self.assertEqual(result["reconciliation"]["detail_invoice_count"], 1)
        self.assertEqual(result["reconciliation"]["issue_count"], 0)
        self.assertEqual(result["reconciliation"]["coverage_ranges"], [{
            "date_from": "2026-01-01", "date_to": "2026-06-30",
        }])

    def test_selected_counts_are_separate_from_partial_comparable_counts(self):
        comparable_overview = [
            {**self.overview(index), "tdlap": "2026-06-15", "tgtcthue": 100, "tgtthue": 10}
            for index in range(1, 1768)
        ]
        uncovered_overview = [
            {**self.overview(index), "tdlap": "2026-07-15", "tgtcthue": 100, "tgtthue": 10}
            for index in range(1768, 1961)
        ]
        details = [
            {**self.detail_lines(index, count=1)[0], "ntao": "2026-06-15"}
            for index in range(1, 1768)
        ]
        result = self.reconcile(
            comparable_overview + uncovered_overview, details,
            jobs=[self.job(
                coverage_from="2026-01-01",
                overview_to="2026-07-31", detail_to="2026-06-30",
            )],
            date_from="2026-01-01", date_to="2026-07-31",
        )
        summary = result["reconciliation"]
        self.assertEqual(summary["selected_overview_invoice_count"], 1960)
        self.assertEqual(summary["selected_detail_invoice_count"], 1767)
        self.assertEqual(summary["selected_difference"], 193)
        self.assertEqual(summary["overview_invoice_count"], 1767)
        self.assertEqual(summary["detail_invoice_count"], 1767)
        self.assertEqual(summary["difference"], 0)
        self.assertEqual(summary["missing_detail_count"], 0)
        self.assertEqual(summary["uncovered_ranges"], [{
            "date_from": "2026-07-01", "date_to": "2026-07-31",
        }])

    def test_full_coverage_reports_all_missing_detail_invoice_keys(self):
        overview = [
            {**self.overview(index), "tdlap": "2026-06-15", "tgtcthue": 100, "tgtthue": 10}
            for index in range(1, 1961)
        ]
        details = [
            {**self.detail_lines(index, count=1)[0], "ntao": "2026-06-15"}
            for index in range(1, 1768)
        ]
        result = self.reconcile(
            overview, details,
            jobs=[self.job(
                coverage_from="2026-01-01",
                overview_to="2026-07-31", detail_to="2026-07-31",
            )],
            date_from="2026-01-01", date_to="2026-07-31",
        )
        summary = result["reconciliation"]
        self.assertEqual(summary["selected_overview_invoice_count"], 1960)
        self.assertEqual(summary["selected_detail_invoice_count"], 1767)
        self.assertEqual(summary["overview_invoice_count"], 1960)
        self.assertEqual(summary["detail_invoice_count"], 1767)
        self.assertEqual(summary["difference"], 193)
        self.assertEqual(summary["missing_detail_count"], 193)
        self.assertEqual(summary["uncovered_ranges"], [])

    def test_active_job_range_is_not_reconciled(self):
        completed = self.job()
        active = self.job(status="running")
        reader = FakeResultReader([self.overview(1)], [])
        query = {
            "connection_id": "conn_account_1", "date_from": "2026-08-01",
            "date_to": "2026-08-31", "direction": "purchase",
            "query_type": "query", "limit": 50,
        }
        with patch("app.external_api.results.JobResultReader", return_value=reader), patch(
            "mia_source_results._result_schema", return_value=SCHEMA
        ):
            result = self.backend([completed, active]).reconciliation(query)
        self.assertEqual(result["reconciliation"]["coverage_ranges"], [])
        self.assertEqual(result["reconciliation"]["issue_count"], 0)

    def test_missing_side_money_values_are_blank_not_zero(self):
        result = self.reconcile([self.overview(1)], [])
        fields = result["items"][0]["fields"]
        self.assertEqual(fields["reconciliation_status"], "Thiếu chi tiết")
        self.assertIsNone(fields["detail_thtien"])
        self.assertIsNone(fields["difference_tgtcthue"])

    def test_reconciliation_export_writes_fixed_schema_and_blank_missing_values(self):
        reader = FakeResultReader(
            [self.overview(1), self.overview(2, total=551)],
            self.detail_lines(2),
        )
        backend = self.backend()
        progress_events = []
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.external_api.results.JobResultReader", return_value=reader
        ), patch("mia_source_results._result_schema", return_value=SCHEMA):
            result = backend.export_results({
                "destination": directory,
                "connection_ids": ["conn_account_1"],
                "result_scopes": ["reconciliation"],
                "date_from": "2026-08-01", "date_to": "2026-08-31",
                "direction": "purchase", "query_type": "query",
                "result_filters": {"reconciliation": {
                    "search": "", "column_filters": {},
                }},
            }, progress_callback=progress_events.append)
            self.assertEqual(result["count"], 1)
            self.assertTrue(any(
                event.get("scope") == "reconciliation"
                and event.get("phase") == "write_rows"
                for event in progress_events
            ))
            self.assertEqual(progress_events[-1]["percent"], 100.0)
            self.assertEqual(progress_events[-1]["status"], "completed")
            workbook = load_workbook(result["files"][0], data_only=False)
            try:
                worksheet = workbook["Bao cao doi chieu"]
                header_row = next(
                    row for row in range(1, worksheet.max_row + 1)
                    if worksheet.cell(row, 1).value == "STT"
                )
                headers = [cell.value for cell in worksheet[header_row]]
                self.assertIn("Trạng thái đối chiếu", headers)
                self.assertIn("Lý do chênh lệch", headers)
                detail_total_column = headers.index("Tổng tiền trước thuế - Chi tiết") + 1
                difference_column = headers.index("Tổng tiền trước thuế - Chênh lệch") + 1
                self.assertIsNone(worksheet.cell(header_row + 1, detail_total_column).value)
                self.assertIsNone(worksheet.cell(header_row + 1, difference_column).value)
                reason_column = headers.index("Lý do chênh lệch") + 1
                self.assertIn("không tìm thấy dữ liệu Chi tiết", worksheet.cell(header_row + 1, reason_column).value)
                self.assertIn("lớn hơn Chi tiết 1 đồng", worksheet.cell(header_row + 2, reason_column).value)
                self.assertFalse(any(
                    cell.value in {"-0", "+0", "-0.00", "+0.00"}
                    for row in worksheet.iter_rows() for cell in row
                ))
                self.assertEqual(workbook.sheetnames, ["Bao cao doi chieu"])
                self.assertEqual(worksheet["A1"].value,
                                 "BÁO CÁO ĐỐI CHIẾU TỔNG QUAN & CHI TIẾT")
                summary_labels = {
                    worksheet.cell(row, 1).value
                    for row in range(2, header_row)
                }
                self.assertIn("Dữ liệu hiện có - Tổng quan", summary_labels)
                self.assertIn("Phạm vi đủ điều kiện đối chiếu", summary_labels)
                self.assertIn("Thiếu Chi tiết", summary_labels)
                self.assertEqual(worksheet.freeze_panes, f"A{header_row + 1}")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
