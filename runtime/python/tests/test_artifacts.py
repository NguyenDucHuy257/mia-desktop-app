import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from mia_artifacts import ArtifactExporter, safe_excel_value
from mia_storage import Storage


class ArtifactExporterTests(unittest.TestCase):
    def test_excel_preserves_unicode_blocks_formula_injection_and_suffixes_duplicates(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = Storage(root / "mia.sqlite3")
            storage.initialize()
            storage.create_account({
                "account_id": "account-1",
                "tax_code": "0101234567",
                "encrypted_password": "cipher",
                "timestamp": "now",
            })
            storage.import_overviews({
                "connection_id": "account-1",
                "timestamp": "now",
                "items": [{
                    "direction": "purchase",
                    "business_key": "=MỞ-CALC",
                    "payload": {"company": "Công ty Việt Nam"},
                }],
            })
            destination = root / "output"
            exporter = ArtifactExporter(storage, root)
            first = exporter.export({
                "destination": str(destination),
                "connection_ids": ["account-1"],
                "kinds": ["excel"],
            })
            second = exporter.export({
                "destination": str(destination),
                "connection_ids": ["account-1"],
                "kinds": ["excel"],
            })
            self.assertEqual(first["count"], 1)
            self.assertIn(" (1).xlsx", second["files"][0])
            workbook = load_workbook(first["files"][0], read_only=True)
            values = list(workbook["Tong quan"].values)
            self.assertEqual(values[1][2], "'=MỞ-CALC")
            self.assertIn("Công ty Việt Nam", values[1][3])
            workbook.close()

    def test_rejects_relative_destination_invalid_kind_and_duplicate_accounts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = Storage(root / "mia.sqlite3")
            storage.initialize()
            exporter = ArtifactExporter(storage, root)
            with self.assertRaisesRegex(ValueError, "invalid_artifact_directory"):
                exporter.export({
                    "destination": "relative",
                    "connection_ids": ["a"],
                    "kinds": ["xml"],
                })
            with self.assertRaisesRegex(ValueError, "invalid_artifact_kind"):
                exporter.export({
                    "destination": str(root),
                    "connection_ids": ["a"],
                    "kinds": ["exe"],
                })
            with self.assertRaisesRegex(ValueError, "invalid_artifact_accounts"):
                exporter.export({
                    "destination": str(root),
                    "connection_ids": ["a", "a"],
                    "kinds": ["xml"],
                })

    def test_excel_export_prefers_production_pipeline_workbooks(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = Storage(root / "mia.sqlite3")
            storage.initialize()
            storage.create_account({
                "account_id": "account-1",
                "tax_code": "0101234567",
                "encrypted_password": "cipher",
                "timestamp": "now",
            })
            source = (
                root
                / "source-data"
                / "0101234567"
                / "exports"
                / "overview"
                / "job-id"
                / "invoice.xlsx"
            )
            source.parent.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(["production pipeline"])
            workbook.save(source)
            destination = root / "output"

            result = ArtifactExporter(storage, root).export({
                "destination": str(destination),
                "connection_ids": ["account-1"],
                "kinds": ["excel"],
            })

            self.assertEqual(result["count"], 1)
            copied = load_workbook(result["files"][0], read_only=True)
            self.assertEqual(
                copied.active.cell(1, 1).value, "production pipeline"
            )
            copied.close()

    def test_source_connection_lists_and_exports_package_files_without_legacy_account(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = Storage(root / "mia.sqlite3")
            storage.initialize()
            connection_id = "conn_source_123456"
            tax_code = "0101234567"

            with closing(sqlite3.connect(root / "source-control.sqlite3")) as control:
                control.executescript("""
                    CREATE TABLE account_connections(
                        connection_id TEXT PRIMARY KEY,
                        username TEXT NOT NULL,
                        status TEXT NOT NULL
                    );
                    CREATE TABLE crawl_jobs(
                        job_id TEXT PRIMARY KEY,
                        account_key TEXT NOT NULL,
                        created_at TEXT NOT NULL
                    );
                """)
                control.execute(
                    "INSERT INTO account_connections VALUES (?,?,?)",
                    (connection_id, tax_code, "ready"),
                )
                control.execute(
                    "INSERT INTO crawl_jobs VALUES (?,?,?)",
                    ("job-source", connection_id, "2026-08-21T08:00:00+00:00"),
                )
                control.commit()

            xml_path = (
                root
                / "source-data"
                / tax_code
                / "exports"
                / "invoice_packages"
                / "purchase"
                / "query"
                / "2026-08-01_2026-08-21"
                / "xml"
                / "hoa-don-source.xml"
            )
            xml_path.parent.mkdir(parents=True)
            xml_path.write_text("<invoice/>", encoding="utf-8")

            source_db = root / "source-data" / tax_code / "db" / "invoices.sqlite3"
            source_db.parent.mkdir(parents=True)
            with closing(sqlite3.connect(source_db)) as invoices:
                invoices.execute("""
                    CREATE TABLE invoice_package_items(
                        id INTEGER PRIMARY KEY,
                        company_tax_code TEXT NOT NULL,
                        direction TEXT NOT NULL,
                        nlap_date TEXT,
                        updated_at TEXT,
                        xml_path TEXT,
                        html_path TEXT,
                        xml_fetched INTEGER NOT NULL DEFAULT 0,
                        html_fetched INTEGER NOT NULL DEFAULT 0,
                        unavailable INTEGER NOT NULL DEFAULT 0
                    )
                """)
                invoices.execute(
                    """INSERT INTO invoice_package_items(
                        id,company_tax_code,direction,nlap_date,updated_at,
                        xml_path,xml_fetched,html_fetched,unavailable
                    ) VALUES (?,?,?,?,?,?,?,?,?)""",
                    (
                        1,
                        tax_code,
                        "purchase",
                        "2026-08-15",
                        "2026-08-21T08:00:00+00:00",
                        str(xml_path),
                        1,
                        0,
                        0,
                    ),
                )
                invoices.commit()

            exporter = ArtifactExporter(storage, root)
            page = exporter.list({
                "connection_ids": [connection_id],
                "kind": "xml",
                "direction": "purchase",
                "date_from": "2026-08-01",
                "date_to": "2026-08-21",
                "limit": 50,
            })
            self.assertEqual(len(page["items"]), 1)
            self.assertEqual(page["items"][0]["job_id"], "job-source")
            self.assertEqual(page["items"][0]["filename"], xml_path.name)
            self.assertEqual(page["items"][0]["connection_id"], connection_id)

            outside_range = exporter.list({
                "connection_ids": [connection_id],
                "kind": "xml",
                "direction": "purchase",
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "limit": 50,
            })
            self.assertEqual(outside_range["items"], [])

            destination = root / "output"
            exported = exporter.export({
                "destination": str(destination),
                "connection_ids": [connection_id],
                "kinds": ["xml"],
            })
            self.assertEqual(exported["count"], 1)
            copied = Path(exported["files"][0])
            self.assertTrue(copied.is_file())
            self.assertEqual(copied.read_text(encoding="utf-8"), "<invoice/>")

    def test_safe_excel_value_covers_all_formula_prefixes(self):
        for prefix in ("=", "+", "-", "@"):
            self.assertTrue(safe_excel_value(prefix + "payload").startswith("'"))

    def test_lists_only_selected_local_job_artifacts_with_cursor(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            storage = Storage(root / "mia.sqlite3")
            storage.initialize()
            storage.create_account({
                "account_id": "account-1",
                "tax_code": "0101234567",
                "encrypted_password": "cipher",
                "timestamp": "now",
            })
            storage.create_job({
                "job_id": "job_1",
                "connection_id": "account-1",
                "idempotency_key": "key",
                "intent": {},
                "timestamp": "now",
            })
            export_root = (
                root
                / "crawler-data"
                / "0101234567"
                / "exports"
                / "invoice_packages"
                / "purchase"
            )
            export_root.mkdir(parents=True)
            (export_root / "hóa-đơn-1.xml").write_text(
                "<xml/>", encoding="utf-8"
            )
            (export_root / "hóa-đơn-2.xml").write_text(
                "<xml/>", encoding="utf-8"
            )
            exporter = ArtifactExporter(storage, root)
            first = exporter.list({
                "connection_ids": ["account-1"],
                "kind": "xml",
                "direction": "purchase",
                "limit": 1,
            })
            second = exporter.list({
                "connection_ids": ["account-1"],
                "kind": "xml",
                "direction": "purchase",
                "limit": 1,
                "cursor": first["pagination"]["next_cursor"],
            })
            self.assertEqual(len(first["items"]), 1)
            self.assertEqual(len(second["items"]), 1)
            self.assertNotEqual(
                first["items"][0]["artifact_id"],
                second["items"][0]["artifact_id"],
            )
            self.assertFalse(second["pagination"]["has_more"])
            with self.assertRaisesRegex(ValueError, "invalid_artifact_cursor"):
                exporter.list({
                    "connection_ids": ["account-1"],
                    "kind": "xml",
                    "cursor": "bad",
                    "limit": 1,
                })


if __name__ == "__main__":
    unittest.main()
